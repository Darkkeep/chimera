"""Prepare only the isolated authored-scene documentation fixture."""
import json
import os
from pathlib import Path
import subprocess
import sys

ROOT = Path(__file__).resolve().parents[3]
DEMO = Path(os.environ.get("CHIMERA_GUIDE_DIR", ROOT / "workspace/user-guide-authored-20261027")).resolve()
if not (DEMO / "fixture.json").is_file():
    subprocess.run([sys.executable, str(Path(__file__).with_name("seed.py"))], env={**os.environ, "CHIMERA_GUIDE_DIR": str(DEMO)}, check=True)
os.environ["CHIMERA_WORKSPACE"] = str(DEMO / "data")
sys.path.insert(0, str(ROOT / "backend"))
from app import storage
from app.authoring import native_layout
from app.models import SceneElement, SceneLayout
from app.render import render_slides
from PIL import Image, ImageDraw
from pptx import Presentation

fixture = json.loads((DEMO / "fixture.json").read_text(encoding="utf-8"))
project = storage.read_project(fixture["projectId"])
if project.title != "Chimera Guide Demonstration" or project.storyboard.scenes[0].id != fixture["sceneId"]:
    raise SystemExit("Refusing to modify an unexpected documentation fixture.")
deck = Presentation()
slide = deck.slides.add_slide(deck.slide_layouts[1])
slide.shapes.title.text = "Storyboard basics - source"
slide.placeholders[1].text = "Draft objective: identify the purpose of a storyboard.\nDraft objective: review before publishing."
source = storage.project_path(project.id) / "original/source-deck.pptx"
deck.save(source)
renders, warning = render_slides(source, storage.project_path(project.id))
if 1 not in renders:
    raise SystemExit(warning or "The synthetic source slide could not be rendered.")
first = project.storyboard.scenes[0]
image = next(asset for asset in project.storyboard.assets if asset.kind == "image")
layout = SceneLayout(enabled=True, backgroundMode="inherit", backgroundColor="#F1F5F9", elements=[
    SceneElement(id="guide-title", kind="text", text=first.title, x=72, y=72, width=1100, height=100,
                 fontSize=56, fontFamily="Aptos", color="#17324D", bold=True),
    *[SceneElement(id=f"guide-objective-{index}", kind="text", text=value, x=72, y=220 + index * 120,
                  width=830, height=100, fontSize=32, fontFamily="Aptos", color="#17324D")
      for index, value in enumerate(first.extractedText)],
    SceneElement(id="guide-illustration", kind="image", assetId=image.id, x=970, y=210, width=230, height=250),
])
scenes = [first.model_copy(update={"visualLayout": layout, "previewImage": renders[1], "sourceSlide": 1, "renderWarning": None})]
scenes.extend(scene if scene.visualLayout else scene.model_copy(update={"visualLayout": native_layout(scene)})
              for scene in project.storyboard.scenes[1:])
storage.write_storyboard(storage.project_path(project.id), project.storyboard.model_copy(update={"scenes": scenes}), project.storyboard.revision)
background = Image.new("RGB", (1280, 720))
draw = ImageDraw.Draw(background)
for y in range(720):
    tone = int(y / 720 * 12)
    draw.line((0, y, 1280, y), fill=(240 - tone, 248 - tone, 254 - tone))
for x in range(930, 1450, 90):
    draw.line((x, 0, x - 220, 720), fill=(198, 221, 240), width=16)
background.save(DEMO / "project-backdrop.png")
fixture.update({"backgroundFile": str(DEMO / "project-backdrop.png"), "source": "Synthetic PowerPoint rendered by local LibreOffice; no user course or provider."})
(DEMO / "fixture.json").write_text(json.dumps(fixture, indent=2), encoding="utf-8")
print("Prepared isolated authored-scene guide fixture:", DEMO)
