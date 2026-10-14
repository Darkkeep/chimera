"""Create disposable guide data. Never reads or changes the user's projects."""
import io
import hashlib
import json
import os
from pathlib import Path
import shutil
import sys
import wave

ROOT = Path(__file__).resolve().parents[3]
DEMO = Path(os.environ.get("CHIMERA_GUIDE_DIR", ROOT / "workspace" / "user-guide-restored"))
if (DEMO / "data" / "projects").exists():
    raise SystemExit("Guide fixture already exists. Use it, or explicitly choose a fresh fixture directory.")
os.environ["CHIMERA_WORKSPACE"] = str(DEMO / "data")
sys.path.insert(0, str(ROOT / "backend"))
from app import storage as store
from app.models import (
    AssessmentChoice, AssessmentSettings, EnhancementProposalCreate,
    ProjectCreate, SceneCreate, SceneUpdate, SingleChoiceAssessment,
    GeneratedAssetCandidate,
)
from pptx import Presentation

project = store.create_project(ProjectCreate(title="Chimera Guide Demonstration"))
project = store.create_scene(project.id, SceneCreate(
    title="Welcome and objectives", layoutMode="native",
    extractedText=["Identify the purpose of a storyboard.", "Review a scene before publishing."],
    actionDescription="Introduce the two objectives.", notes="Author note: introduce the demonstration.",
))
first_id = project.storyboard.scenes[0].id
for title in ["Knowledge check", "Final review"]:
    project = store.create_scene(project.id, SceneCreate(title=title, sceneType="assessment", layoutMode="native"))
    scene_id = project.storyboard.scenes[-1].id
    project = store.update_scene(project.id, scene_id, SceneUpdate(
        assessmentItem=SingleChoiceAssessment(
            identifier="Q" + str(len(project.storyboard.scenes)),
            prompt="What should you do before publishing?",
            choices=[AssessmentChoice(identifier="A", text="Preview and validate the course"),
                     AssessmentChoice(identifier="B", text="Skip the review")],
            correctResponse="A",
        ),
        assessmentSettings=AssessmentSettings(remediationSceneId=first_id),
    ))
# Reuse a repository graphic as fixture content, without modifying it.
image_file = ROOT / "frontend/src/assets/cutout-7-398bc27260c0.png"
project = store.add_project_asset(project.id, "demonstration-image.png", image_file.read_bytes(), "image")
image_id = project.storyboard.assets[-1].id
audio = io.BytesIO()
with wave.open(audio, "wb") as wav:
    wav.setnchannels(1)
    wav.setsampwidth(2)
    wav.setframerate(8000)
    wav.writeframes(b"\0\0" * 8000)
project = store.add_project_asset(project.id, "demonstration-audio.wav", audio.getvalue(), "audio")
project = store.update_scene(project.id, first_id, SceneUpdate(
    imageAssets=[image_id], audioAssets=[project.storyboard.assets[-1].id],
    imageAltText={image_id: "Demonstration graphic for the user guide."},
))
render = store.project_path(project.id) / "renders" / "slide-001.png"
render.parent.mkdir(exist_ok=True)
shutil.copyfile(image_file, render)
scenes = list(project.storyboard.scenes)
scenes[0] = scenes[0].model_copy(update={"sourceSlide": 1, "previewImage": "renders/slide-001.png"})
project = store.replace_storyboard(project.id, project.storyboard.model_copy(update={"scenes": scenes}))
deck = Presentation()
slide = deck.slides.add_slide(deck.slide_layouts[1])
slide.shapes.title.text = "Demonstration template"
slide.placeholders[1].text = "Replace with your course content."
source = io.BytesIO()
deck.save(source)
store.create_template("Demonstration template", "A sample template for guide screenshots.", "demonstration-template.pptx", source.getvalue())
store.create_enhancement_proposal(project.id, EnhancementProposalCreate(
    targetSceneId=first_id,
    proposedValue="Introduce each objective, then ask learners to identify what they will practice.",
    rationale="Demonstration proposal: make the opening activity more specific.", origin="author",
))
(DEMO / "fixture.json").write_text(json.dumps({"projectId": project.id, "sceneId": first_id}), encoding="utf-8")
(DEMO / "study-guide.txt").write_text("Study guide\nReview each scene for clarity. Preview and validate the course before publishing.", encoding="utf-8")
content = image_file.read_bytes()
store.write_generated_candidate(GeneratedAssetCandidate(
    id="candidate-0123456789ab", projectId=project.id, jobId="media-generation-0123456789ab",
    sourceRevision=project.storyboard.revision, sceneId=first_id, kind="image",
    filename="demonstration-candidate.png", contentType="image/png", sizeBytes=len(content),
    sha256=hashlib.sha256(content).hexdigest(), path="generated/candidates/demonstration-candidate.png",
    purpose="Demonstrate candidate review controls.", prompt="Documentation fixture; no AI request was made.",
    altText="Demonstration graphic.", provider="Documentation fixture", createdAt=store.now_iso(),
), content)
print("Created isolated documentation fixture:", DEMO)
