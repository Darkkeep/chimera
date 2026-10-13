import io
import json
import sys
from copy import deepcopy
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Inches, Pt

from backend.app import storage
from backend.app.ingest import extract_presentation
from backend.app.models import ProjectCreate
from backend.app.scorm import build_scorm_12, build_scorm_2004
from zipfile import ZipFile


def fixture() -> None:
    project = storage.create_project(ProjectCreate(title="Visual qualification"))
    root = storage.project_path(project.id)
    presentation = Presentation()
    picture = io.BytesIO()
    Image.new("RGB", (120, 80), "#19a887").save(picture, format="PNG")
    background = io.BytesIO()
    Image.new("RGB", (120, 80), "#edf4f8").save(background, format="PNG")
    for index in range(24):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        title = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(1))
        title.text = f"Qualification slide {index + 1}"
        title.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
        body = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(1))
        body.text = "Editable scene text"
        body.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
        slide.shapes.add_picture(io.BytesIO(picture.getvalue()), Inches(1), Inches(4), Inches(2))
        if index == 0:
            master = slide.slide_layout.slide_master
            master_picture = deepcopy(slide.shapes[-1]._element)
            blip = master_picture.find(".//" + qn("a:blip"))
            part = slide.part.rels[blip.get(qn("r:embed"))].target_part
            blip.set(qn("r:embed"), master.part.relate_to(part, RT.IMAGE))
            master_picture.find(".//" + qn("a:off")).set("x", "0")
            master_picture.find(".//" + qn("a:off")).set("y", "0")
            master.shapes._spTree.append(master_picture)
            temporary_picture = slide.shapes.add_picture(io.BytesIO(background.getvalue()), 0, 0, Inches(1))
            image_part = slide.part.rels[temporary_picture._element.find(".//" + qn("a:blip")).get(qn("r:embed"))].target_part
            relationship_id = master.part.relate_to(image_part, RT.IMAGE)
            temporary_picture._element.getparent().remove(temporary_picture._element)
            background_node = OxmlElement("p:bg")
            properties = OxmlElement("p:bgPr")
            fill = OxmlElement("a:blipFill")
            background_blip = OxmlElement("a:blip")
            background_blip.set(qn("r:embed"), relationship_id)
            fill.append(background_blip)
            stretch = OxmlElement("a:stretch")
            stretch.append(OxmlElement("a:fillRect"))
            fill.append(stretch)
            properties.append(fill)
            background_node.append(properties)
            master._element.find(qn("p:cSld")).insert(0, background_node)
            parallelogram = slide.shapes.add_shape(MSO_SHAPE.PARALLELOGRAM, Inches(8), Inches(1), Inches(3), Inches(1.5))
            parallelogram.fill.solid()
            parallelogram.fill.fore_color.rgb = RGBColor(30, 120, 180)
            slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(8), Inches(3), Inches(11), Inches(3))
    source = root / "original" / "source-deck.pptx"
    presentation.save(source)
    scenes, assets = extract_presentation(source, root)
    for asset in assets:
        asset.attributionText = "Generated solid-color test artwork for qualification."
        asset.attributionStatus = "complete"
    for scene in scenes:
        scene.imageAltText = {asset_id: "Qualification artwork" for asset_id in scene.imageAssets}
    storage.update_from_import(project.id, scenes, assets)
    legacy = storage.create_project(ProjectCreate(title="Legacy visual qualification"))
    legacy_root = storage.project_path(legacy.id)
    (legacy_root / "original" / "source-deck.pptx").write_bytes(source.read_bytes())
    legacy_scenes, legacy_assets = extract_presentation(source, legacy_root)
    for scene in legacy_scenes:
        scene.visualLayout = None
    legacy_scenes[0].notes = "Preserved legacy notes"
    storage.update_from_import(legacy.id, legacy_scenes, legacy_assets)
    print(json.dumps({"projectId": project.id, "legacyProjectId": legacy.id}))


def exports(project_id: str, destination: str) -> None:
    project = storage.read_project(project_id)
    for profile, builder in [("scorm12", build_scorm_12), ("scorm2004", build_scorm_2004)]:
        package, report = builder(project)
        target = Path(destination) / profile
        target.mkdir(parents=True, exist_ok=True)
        with ZipFile(io.BytesIO(package)) as archive:
            archive.extractall(target)
        (target / "report.json").write_text(json.dumps(report, indent=2), encoding="utf-8")


if len(sys.argv) > 1:
    exports(sys.argv[1], sys.argv[2])
else:
    fixture()