"""Small inert presentation for onboarding import/template checks."""
import io
import sys
from pptx import Presentation

deck = Presentation()
slide = deck.slides.add_slide(deck.slide_layouts[1])
slide.shapes.title.text = "Onboarding fixture"
slide.placeholders[1].text = "Demonstration content only."
output = io.BytesIO()
deck.save(output)
sys.stdout.buffer.write(output.getvalue())
