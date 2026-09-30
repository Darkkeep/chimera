import io
from pathlib import Path

import pymupdf
from docx import Document
from pptx import Presentation

from .models import ReferencePassage


SUPPORTED_REFERENCE_EXTENSIONS = {".pdf", ".docx", ".pptx", ".txt", ".md"}


def extract_reference_passages(filename: str, content: bytes) -> tuple[list[ReferencePassage], str, str | None]:
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_REFERENCE_EXTENSIONS:
        raise ValueError("Reference files must be PDF, DOCX, PPTX, TXT, or Markdown")
    passages: list[ReferencePassage] = []
    try:
        if extension == ".pdf":
            document = pymupdf.open(stream=content, filetype="pdf")
            for page_number, page in enumerate(document, start=1):
                text = page.get_text("text").strip()
                if text:
                    passages.append(ReferencePassage(locator=f"page:{page_number}", text=text[:50000]))
            document.close()
        elif extension == ".docx":
            document = Document(io.BytesIO(content))
            for paragraph_number, paragraph in enumerate(document.paragraphs, start=1):
                text = paragraph.text.strip()
                if text:
                    passages.append(ReferencePassage(locator=f"paragraph:{paragraph_number}", text=text[:50000]))
            for table_index, table in enumerate(document.tables, start=1):
                for row_index, row in enumerate(table.rows, start=1):
                    text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if text:
                        passages.append(ReferencePassage(locator=f"table:{table_index}:row:{row_index}", text=text[:50000]))
        elif extension == ".pptx":
            presentation = Presentation(io.BytesIO(content))
            for slide_number, slide in enumerate(presentation.slides, start=1):
                lines = [shape.text.strip() for shape in slide.shapes if getattr(shape, "has_text_frame", False) and shape.text.strip()]
                if lines:
                    passages.append(ReferencePassage(locator=f"slide:{slide_number}", text="\n".join(lines)[:50000]))
        else:
            text = content.decode("utf-8-sig").strip()
            for paragraph_number, paragraph in enumerate((part.strip() for part in text.splitlines()), start=1):
                if paragraph:
                    passages.append(ReferencePassage(locator=f"line:{paragraph_number}", text=paragraph[:50000]))
    except Exception as error:
        return [], "failed", f"Could not extract readable text: {error}"
    if not passages:
        return [], "failed", "No readable text was extracted. Scanned PDFs require OCR before they can ground suggestions."
    return passages, "complete", None
