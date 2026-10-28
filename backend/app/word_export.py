import io
import tempfile
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.section import Section
from docx.shared import Emu, Inches, Pt, RGBColor
from docx.table import _Cell
from docx.text.paragraph import Paragraph
from PIL import Image, UnidentifiedImageError
from pptx import Presentation
from pptx.oxml.ns import qn as pptx_qn

from .models import AssessmentItem, ExtractedAsset, ProjectDetail, Scene
from .render import render_slides
from .storage import project_path, list_reference_resources, list_enhancement_proposals, list_review_comments
from .media import resolved_project
from .scene_print import scene_image, with_scene_images


DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
TEMPLATE_ROOT = Path(__file__).resolve().parents[2]
SME_TEMPLATE_PATH = TEMPLATE_ROOT / "Visual_Slide_Storyboard_Template.pptx"
FULL_TEMPLATE_PATH = TEMPLATE_ROOT / "Multimedia_Instructional_Storyboard_Template.docx"


class TemplateArtworkError(RuntimeError):
    pass


@dataclass(frozen=True)
class SMEVisualStyle:
    font_name: str
    title_color: str
    metadata_color: str
    accent_color: str
    body_color: str
    panel_fill: str
    panel_border: str


def _clean_xml_value(value: object) -> object:
    if isinstance(value, str):
        return "".join(
            character
            for character in value
            if character in "\t\n\r" or 0x20 <= ord(character) <= 0xD7FF or 0xE000 <= ord(character) <= 0xFFFD or 0x10000 <= ord(character) <= 0x10FFFF
        )
    if isinstance(value, list):
        return [_clean_xml_value(item) for item in value]
    if isinstance(value, dict):
        return {_clean_xml_value(key): _clean_xml_value(item) for key, item in value.items()}
    return value


def _asset_path(project: ProjectDetail, relative_path: str, folder: str) -> Path | None:
    project_root = project_path(project.id).resolve()
    allowed_root = (project_root / folder).resolve()
    candidate = (project_root / relative_path).resolve()
    if candidate.is_relative_to(allowed_root) and candidate.is_file():
        return candidate
    return None


def _attribution_label(asset: ExtractedAsset) -> str:
    if asset.kind == "image" and (asset.attributionStatus != "complete" or not asset.attributionText.strip()):
        return f"Attribution Pending: {asset.attributionText}" if asset.attributionText.strip() else "Attribution Pending"
    return f"{asset.attributionText} ({asset.attributionStatus})" if asset.attributionText else asset.attributionStatus


def _add_image(document: DocumentType, path: Path, label: str) -> bool:
    try:
        with Image.open(path) as source:
            image_buffer = io.BytesIO()
            source.convert("RGB").save(image_buffer, format="PNG")
        image_buffer.seek(0)
        document.add_paragraph(label, style="Caption")
        document.add_picture(image_buffer, width=Inches(5.8))
        return True
    except (OSError, UnidentifiedImageError, ValueError):
        return False


def _add_media_thumbnail(cell: _Cell, path: Path | None) -> bool:
    if path is None:
        return False
    try:
        with Image.open(path) as source:
            aspect_ratio = source.width / source.height
            max_width = 0.55
            max_height = 0.42
            if aspect_ratio >= max_width / max_height:
                width = max_width
                height = max_width / aspect_ratio
            else:
                height = max_height
                width = max_height * aspect_ratio
            image_buffer = io.BytesIO()
            source.convert("RGB").save(image_buffer, format="PNG")
        image_buffer.seek(0)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(image_buffer, width=Inches(width), height=Inches(height))
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        return True
    except (OSError, UnidentifiedImageError, ValueError):
        return False


def _assessment_choices(item: AssessmentItem) -> list[tuple[str, str]]:
    if item.type in {"single-choice", "multiple-response", "ordering"}:
        return [(choice.identifier, choice.text) for choice in item.choices]
    if item.type == "true-false":
        return [("true", "True"), ("false", "False")]
    if item.type == "drag-drop":
        return [(choice.identifier, f"Item: {choice.text}") for choice in item.draggables] + [
            (choice.identifier, f"Target: {choice.text}") for choice in item.targets
        ]
    if item.type == "matching":
        return [(choice.identifier, f"Left: {choice.text}") for choice in item.leftItems] + [
            (choice.identifier, f"Right: {choice.text}") for choice in item.rightItems
        ]
    return []


def _answer_key(item: AssessmentItem) -> str:
    choices = dict(_assessment_choices(item))
    if item.type in {"single-choice", "true-false"}:
        return choices.get(item.correctResponse, item.correctResponse)
    if item.type == "multiple-response":
        return "; ".join(choices.get(identifier, identifier) for identifier in item.correctResponses)
    if item.type == "ordering":
        return " -> ".join(choices.get(identifier, identifier) for identifier in item.correctOrder)
    if item.type in {"drag-drop", "matching"}:
        return "; ".join(
            f"{choices.get(pair.sourceIdentifier, pair.sourceIdentifier)} -> {choices.get(pair.targetIdentifier, pair.targetIdentifier)}"
            for pair in item.correctMatches
        )
    if item.type == "short-answer":
        return "; ".join(item.acceptedAnswers)
    return ""


def _add_assessment(document: DocumentType, scene: Scene, include_answer_key: bool) -> None:
    item = scene.assessmentItem
    if item is None:
        return
    document.add_heading("Assessment specification", level=2)
    interaction_name = "Fill-in-the-blank" if item.type == "short-answer" else item.type.replace("-", " ").title()
    document.add_paragraph(f"Interaction: {interaction_name}")
    document.add_paragraph(f"Identifier: {item.identifier}")
    document.add_paragraph(f"Prompt: {item.prompt}")
    choices = _assessment_choices(item)
    if choices:
        table = document.add_table(rows=1, cols=2)
        table.style = "Light Shading Accent 1"
        table.rows[0].cells[0].text = "Identifier"
        table.rows[0].cells[1].text = "Response / interaction element"
        for identifier, text in choices:
            cells = table.add_row().cells
            cells[0].text = identifier
            cells[1].text = text
    elif item.type == "short-answer":
        document.add_paragraph("Accepted response: learner fills in one blank with an exact-match word or short phrase.")

    if item.type == "ordering":
        document.add_paragraph("Expected interaction: order the listed elements.")
    elif item.type in {"drag-drop", "matching"}:
        document.add_paragraph("Expected interaction: match each source element with a target.")
    elif item.type == "multiple-response":
        document.add_paragraph("Expected interaction: select every applicable response.")
    elif item.type in {"single-choice", "true-false"}:
        document.add_paragraph("Expected interaction: select one response.")

    document.add_paragraph(f"Correct feedback: {item.correctFeedback}")
    document.add_paragraph(f"Incorrect feedback: {item.incorrectFeedback}")
    if scene.assessmentSettings:
        settings = scene.assessmentSettings
        document.add_paragraph(
            f"Scoring: {settings.maxAttempts} attempts; {settings.passingScore}% question threshold; "
            f"weight {settings.scoreWeight:g}."
        )
    if include_answer_key:
        paragraph = document.add_paragraph()
        paragraph.add_run("Answer key (authoring material): ").bold = True
        paragraph.add_run(_answer_key(item))


def _add_scene(document: DocumentType, project: ProjectDetail, scene: Scene, index: int, include_answer_key: bool) -> None:
    document.add_page_break()
    document.add_heading(f"Scene {index}: {scene.title or 'Untitled scene'}", level=1)
    metadata = document.add_table(rows=0, cols=2)
    metadata.style = "Light Shading Accent 1"
    metadata_rows = [
        ("Scene ID", scene.id),
        ("Type", scene.sceneType),
        ("Required", "Yes" if scene.required else "No (optional)"),
    ]
    if scene.sourceSlide is not None:
        metadata_rows.append(("Source slide", str(scene.sourceSlide)))
    if scene.templateId:
        metadata_rows.append(("Template ID", scene.templateId))
    for label, value in metadata_rows:
        cells = metadata.add_row().cells
        cells[0].text = label
        cells[1].text = value

    document.add_heading("Learner-facing content", level=2)
    if scene.extractedText:
        for text in scene.extractedText:
            document.add_paragraph(text)
    else:
        document.add_paragraph("No learner-facing text recorded.")

    if scene.visualLayout or scene.previewImage:
        preview = scene_image(scene.id)
        if preview is None and not scene.visualLayout and scene.previewImage:
            preview = _asset_path(project, scene.previewImage, "renders")
        if preview is None or not _add_image(document, preview, "Authored scene"):
            raise RuntimeError(f"Authored scene visual is unavailable for '{scene.title}'.")

    assets_by_id = {asset.id: asset for asset in project.storyboard.assets}
    attached_images = [assets_by_id[asset_id] for asset_id in scene.imageAssets if asset_id in assets_by_id]
    attached_audio = [assets_by_id[asset_id] for asset_id in scene.audioAssets if asset_id in assets_by_id]
    attached_audio += [assets_by_id[asset_id] for asset_id in scene.videoAssets if asset_id in assets_by_id]
    if attached_images or attached_audio:
        document.add_heading("Media", level=2)
        media_table = document.add_table(rows=1, cols=5)
        media_table.style = "Light Shading Accent 1"
        # Long opaque asset IDs must not consume the accessibility/credit columns.
        media_table.autofit = False
        section = document.sections[-1]
        available_width = section.page_width - section.left_margin - section.right_margin
        width_ratios = (0.65, 0.55, 1.7, 2.5, 1.6)
        ratio_total = sum(width_ratios)
        column_widths = [Emu(round(available_width * ratio / ratio_total)) for ratio in width_ratios]
        column_widths[-1] = Emu(available_width - sum(column_widths[:-1]))
        for column, width in zip(media_table.columns, column_widths):
            column.width = width
        for cell, width in zip(media_table.rows[0].cells, column_widths):
            cell.width = width
        media_table.rows[0].cells[0].text = "Role"
        media_table.rows[0].cells[1].text = "Preview"
        media_table.rows[0].cells[2].text = "File / reference"
        media_table.rows[0].cells[3].text = "Accessibility / attribution"
        media_table.rows[0].cells[4].text = "Credit"
        for asset in attached_images:
            cells = media_table.add_row().cells
            cells[0].text = "Image"
            image_path = _asset_path(project, asset.path, "assets")
            _add_media_thumbnail(cells[1], image_path)
            cells[2].text = f"{asset.filename} ({asset.id})"
            alt_text = scene.imageAltText.get(asset.id, "")
            cells[3].text = "Decorative image" if asset.id in scene.decorativeImages else alt_text or "Alternative text not recorded"
            cells[4].text = _attribution_label(asset)
        for asset in attached_audio:
            cells = media_table.add_row().cells
            cells[0].text = asset.kind.title()
            cells[2].text = f"{asset.filename} ({asset.id})"
            cells[3].text = "Required for completion" if scene.audioRequired else "Optional playback"
            cells[4].text = _attribution_label(asset)
            if asset.transcript:
                document.add_paragraph(f"Transcript — {asset.filename}: {asset.transcript}")
            if asset.captionsVtt:
                document.add_paragraph(f"Captions ({asset.language}) — {asset.filename}:\n{asset.captionsVtt}")

    if scene.actionDescription:
        document.add_heading("Production directions (authoring material)", level=2)
        document.add_paragraph(scene.actionDescription)
    if scene.notes:
        document.add_heading("Speaker notes (authoring material)", level=2)
        document.add_paragraph(scene.notes)
    _add_assessment(document, scene, include_answer_key)


def _set_widescreen_section(section: Section, slide_width: int, slide_height: int) -> None:
    width = Inches(11)
    height = Inches(11 * slide_height / slide_width)
    section.page_width = width
    section.page_height = height
    section.left_margin = Inches(0.4)
    section.right_margin = Inches(0.4)
    section.top_margin = Inches(0.35)
    section.bottom_margin = Inches(0.35)


def _paragraph_label(paragraph: Paragraph, label: str, style: SMEVisualStyle) -> None:
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(label.upper())
    run.bold = True
    run.font.size = Pt(8)
    run.font.name = style.font_name
    run.font.color.rgb = RGBColor.from_string(style.accent_color)


def _cell_label(cell: _Cell, label: str, style: SMEVisualStyle) -> None:
    _paragraph_label(cell.paragraphs[0], label, style)


def _set_cell_border(cell: _Cell, color: str = "D7DEE7") -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    borders = cell_properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        cell_properties.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def _set_cell_fill(cell: _Cell, color: str) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    shading = cell_properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        cell_properties.append(shading)
    shading.set(qn("w:fill"), color)


def _add_sme_scene(document: DocumentType, project: ProjectDetail, scene: Scene, index: int, include_answer_key: bool, style: SMEVisualStyle) -> None:
    if index > 1:
        document.add_page_break()

    header = document.add_table(rows=1, cols=2)
    header.autofit = False
    header.columns[0].width = Inches(7.0)
    header.columns[1].width = Inches(3.2)
    header.cell(0, 0).width = Inches(7.0)
    header.cell(0, 1).width = Inches(3.2)
    title = header.cell(0, 0).paragraphs[0]
    title.paragraph_format.space_after = Pt(0)
    title_run = title.add_run(scene.title or "Untitled scene")
    title_run.bold = True
    title_run.font.size = Pt(20)
    title_run.font.name = style.font_name
    title_run.font.color.rgb = RGBColor.from_string(style.title_color)
    header.cell(0, 1).text = f"SME Master - {project.title}\nSCENE {scene.id} | {index}/{len(project.storyboard.scenes)} | REV {project.storyboard.revision} | PROJECT {project.productionStage}"
    for run in header.cell(0, 1).paragraphs[0].runs:
        run.font.size = Pt(8)
        run.font.name = style.font_name
        run.font.color.rgb = RGBColor.from_string(style.metadata_color)

    metadata = document.add_table(rows=1, cols=4)
    labels = ("OBJECTIVE", "VERSION", "STATUS", "SCENE TYPE")
    values = ("Not linked", str(project.storyboard.revision), "Not tracked", scene.sceneType)
    for cell, label, value in zip(metadata.rows[0].cells, labels, values):
        cell.text = ""
        _cell_label(cell, label, style)
        cell.add_paragraph(value)

    purpose = document.add_table(rows=1, cols=1)
    _cell_label(purpose.cell(0, 0), "Learning purpose", style)
    purpose.cell(0, 0).add_paragraph("Not recorded in Chimera.")

    body = document.add_table(rows=1, cols=2)
    body.autofit = False
    body.columns[0].width = Inches(6.5)
    body.columns[1].width = Inches(3.7)
    visual = body.cell(0, 0)
    content = body.cell(0, 1)
    _cell_label(visual, "What the learner sees", style)
    visual_area = visual.add_table(rows=1, cols=1)
    visual_area.autofit = False
    # Leave room for editable video access details and production notes.
    visual_area.rows[0].height = Inches(2.0 if scene.videoAssets else 2.35)
    visual_area.rows[0].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    visual_area_cell = visual_area.cell(0, 0)
    visual_area_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _set_cell_border(visual_area_cell, style.panel_border)
    _set_cell_fill(visual_area_cell, style.panel_fill)
    preview = scene_image(scene.id)
    if preview is None and not scene.visualLayout and scene.previewImage:
        preview = _asset_path(project, scene.previewImage, "renders")
    if preview is not None:
        try:
            with Image.open(preview) as source:
                aspect_ratio = source.width / source.height
                image_width = min(6.0, 2.15 * aspect_ratio)
                image_buffer = io.BytesIO()
                source.convert("RGB").save(image_buffer, format="PNG")
            image_buffer.seek(0)
            visual_area_cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            visual_area_cell.paragraphs[0].add_run().add_picture(image_buffer, width=Inches(image_width))
        except (OSError, UnidentifiedImageError, ValueError):
            visual_area_cell.text = "Preview image could not be embedded."
    else:
        visual_area_cell.text = "Add an editable sketch or screenshot."
        visual_area_cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    _paragraph_label(visual.add_paragraph(), "Exact on-screen wording", style)
    if scene.extractedText:
        for text in scene.extractedText:
            visual.add_paragraph(text)
    else:
        visual.add_paragraph("No learner-facing text is recorded.")

    _cell_label(content, "What they hear", style)
    content.add_paragraph(scene.narration or "No narration script is recorded.")
    _paragraph_label(content.add_paragraph(), "What they do", style)
    item = scene.assessmentItem
    if item is not None:
        content.add_paragraph(item.prompt)
        for identifier, choice_text in _assessment_choices(item):
            content.add_paragraph(f"{identifier}. {choice_text}")
    else:
        content.add_paragraph("Learner action is not separately recorded.")
    if scene.actionDescription:
        content.add_paragraph(f"Production notes: {scene.actionDescription}")
    _paragraph_label(content.add_paragraph(), "What happens next", style)
    if item is not None:
        content.add_paragraph(f"Correct: {item.correctFeedback}\nIncorrect: {item.incorrectFeedback}")
        if scene.assessmentSettings:
            settings = scene.assessmentSettings
            content.add_paragraph(f"Attempts: {settings.maxAttempts} | Passing score: {settings.passingScore}%")
        if include_answer_key:
            content.add_paragraph(f"Answer key (authoring material): {_answer_key(item)}")
    else:
        content.add_paragraph("No feedback or next-screen rule is recorded.")
    if scene.notes:
        content.add_paragraph(f"Speaker notes: {scene.notes}")

    footer = document.add_table(rows=1, cols=2)
    footer.autofit = False
    footer.columns[0].width = Inches(6.7)
    footer.columns[1].width = Inches(3.5)
    _cell_label(footer.cell(0, 0), "Media and access", style)
    assets_by_id = {asset.id: asset for asset in project.storyboard.assets}
    scene_assets = [assets_by_id[asset_id] for asset_id in scene.imageAssets + scene.audioAssets + scene.videoAssets if asset_id in assets_by_id]
    if scene_assets:
        media_table = footer.cell(0, 0).add_table(rows=1, cols=4)
        media_table.autofit = False
        media_widths = (0.55, 2.0, 2.5, 1.6)
        for column, width in zip(media_table.columns, media_widths):
            column.width = Inches(width)
        for cell, width, label in zip(media_table.rows[0].cells, media_widths, ("Media", "File / reference", "Access", "Credit")):
            cell.width = Inches(width)
            cell.text = label
        for asset in scene_assets:
            cells = media_table.add_row().cells
            for cell, width in zip(cells, media_widths):
                cell.width = Inches(width)
            cells[1].text = asset.filename
            if asset.kind == "image":
                image_path = _asset_path(project, asset.path, "assets")
                _add_media_thumbnail(cells[0], image_path)
                cells[2].text = "Decorative image" if asset.id in scene.decorativeImages else scene.imageAltText.get(asset.id, "") or "Alternative text not recorded"
            else:
                cells[2].text = "Video with captions" if asset.kind == "video" else "Required audio" if scene.audioRequired else "Optional audio"
                if asset.transcript:
                    cells[2].text += "\nTranscript in media appendix."
            cells[3].text = _attribution_label(asset)
    else:
        footer.cell(0, 0).add_paragraph("No attached media is recorded.")
    _cell_label(footer.cell(0, 1), "Review note", style)
    footer.cell(0, 1).add_paragraph("No reviewer note is recorded.")

    for table in (header, metadata, purpose, body, visual_area, footer):
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(3)
                    for run in paragraph.runs:
                        run.font.name = style.font_name
                        if run.font.color.rgb is None:
                            run.font.color.rgb = RGBColor.from_string(style.body_color)
                        if run.font.size is None:
                            run.font.size = Pt(10)


def _normalized_project(project: ProjectDetail) -> ProjectDetail:
    return ProjectDetail.model_validate(_clean_xml_value(resolved_project(project).model_dump(mode="python")))


def _template_slide_dimensions() -> tuple[int, int]:
    if not SME_TEMPLATE_PATH.is_file():
        raise FileNotFoundError(f"SME storyboard template was not found: {SME_TEMPLATE_PATH.name}")
    presentation = Presentation(str(SME_TEMPLATE_PATH))
    slide_width = presentation.slide_width
    slide_height = presentation.slide_height
    if slide_width is None or slide_height is None:
        raise ValueError("SME storyboard template has no slide dimensions")
    return int(slide_width), int(slide_height)


def _pptx_rgb(color, fallback: str) -> str:
    try:
        value = color.rgb
    except (AttributeError, TypeError, ValueError):
        return fallback
    return str(value) if value is not None else fallback


def _shape_font_and_color(shape, fallback_color: str) -> tuple[str | None, str]:
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.text.strip():
                return run.font.name, _pptx_rgb(run.font.color, fallback_color)
    return None, fallback_color


def _extract_sme_visual_style(presentation: Presentation) -> SMEVisualStyle:
    if not presentation.slides:
        return SMEVisualStyle("Aptos", "17283D", "627386", "365B83", "26374B", "FFFFFF", "CCD4DD")
    slide = presentation.slides[min(1, len(presentation.slides) - 1)]
    text_shapes = [shape for shape in slide.shapes if getattr(shape, "text", "").strip()]
    title = next((shape for shape in text_shapes if shape.top < presentation.slide_height / 8), None)
    label = next((shape for shape in text_shapes if "WHAT THE LEARNER SEES" in shape.text.upper()), None)
    body = next((shape for shape in text_shapes if shape.left > presentation.slide_width * 0.6 and shape.top > presentation.slide_height / 3 and not shape.text.isupper()), None)
    panel = next((shape for shape in slide.shapes if not getattr(shape, "text", "").strip() and shape.shape_type == 1), None)
    title_font, title_color = _shape_font_and_color(title, "17283D") if title else (None, "17283D")
    body_font, body_color = _shape_font_and_color(body, "26374B") if body else (None, "26374B")
    _, metadata_color = _shape_font_and_color(text_shapes[1], "627386") if len(text_shapes) > 1 else (None, "627386")
    _, accent_color = _shape_font_and_color(label, "365B83") if label else (None, "365B83")
    panel_fill = _pptx_rgb(panel.fill.fore_color, "FFFFFF") if panel else "FFFFFF"
    panel_border = _pptx_rgb(panel.line.color, "CCD4DD") if panel else "CCD4DD"
    return SMEVisualStyle(title_font or body_font or "Aptos", title_color, metadata_color, accent_color, body_color, panel_fill, panel_border)


def _has_explicit_background(background) -> bool:
    if background is None:
        return False
    background_properties = background.find(pptx_qn("p:bgPr"))
    if background_properties is None:
        return False
    return any(child.tag in {pptx_qn("a:solidFill"), pptx_qn("a:gradFill"), pptx_qn("a:blipFill"), pptx_qn("a:pattFill")} for child in background_properties)


def _template_has_master_artwork(slide) -> bool:
    master = slide.slide_layout.slide_master
    layout = slide.slide_layout
    return bool(
        len(master.shapes)
        or len(layout.shapes)
        or _has_explicit_background(master._element.cSld.bg)
        or _has_explicit_background(layout._element.cSld.bg)
    )


def _extract_master_artwork(template_path: Path, slide_index: int, output_root: Path) -> Path | None:
    presentation = Presentation(str(template_path))
    if slide_index >= len(presentation.slides) or not _template_has_master_artwork(presentation.slides[slide_index]):
        return None
    output_root.mkdir(parents=True, exist_ok=True)
    for slide in presentation.slides:
        for shape in list(slide.shapes):
            shape._element.getparent().remove(shape._element)
    sanitized_path = output_root / "master-artwork.pptx"
    presentation.save(sanitized_path)
    rendered, warning = render_slides(sanitized_path, output_root)
    if warning:
        raise TemplateArtworkError(warning)
    image_path = output_root / rendered.get(slide_index + 1, "")
    if not image_path.is_file():
        raise TemplateArtworkError("PPTX master artwork could not be extracted")
    return image_path


def _set_section_artwork(section: Section, image_path: Path) -> None:
    section.header.is_linked_to_previous = False
    section.header_distance = Inches(0)
    run = section.header.paragraphs[0].add_run()
    inline = run.add_picture(str(image_path), width=section.page_width, height=section.page_height)._inline
    anchor = OxmlElement("wp:anchor")
    for name, value in (("distT", "0"), ("distB", "0"), ("distL", "0"), ("distR", "0"), ("simplePos", "0"), ("relativeHeight", "0"), ("behindDoc", "1"), ("locked", "0"), ("layoutInCell", "0"), ("allowOverlap", "1")):
        anchor.set(name, value)
    simple_position = OxmlElement("wp:simplePos")
    simple_position.set("x", "0")
    simple_position.set("y", "0")
    anchor.append(simple_position)
    for axis in ("H", "V"):
        position = OxmlElement(f"wp:position{axis}")
        position.set("relativeFrom", "page")
        offset = OxmlElement("wp:posOffset")
        offset.text = "0"
        position.append(offset)
        anchor.append(position)
    anchor.append(inline.find(qn("wp:extent")))
    effect = OxmlElement("wp:effectExtent")
    for edge in ("l", "t", "r", "b"):
        effect.set(edge, "0")
    anchor.append(effect)
    anchor.append(OxmlElement("wp:wrapNone"))
    for name in ("wp:docPr", "wp:cNvGraphicFramePr", "a:graphic"):
        anchor.append(inline.find(qn(name)))
    inline.getparent().replace(inline, anchor)


def _append_sme_storyboard(document: DocumentType, project: ProjectDetail, include_answer_key: bool, dimensions: tuple[int, int], style: SMEVisualStyle, artwork: Path | None) -> None:
    section = document.add_section(WD_SECTION.NEW_PAGE)
    _set_widescreen_section(section, *dimensions)
    if artwork:
        _set_section_artwork(section, artwork)
    for index, scene in enumerate(project.storyboard.scenes, start=1):
        _add_sme_scene(document, project, scene, index, include_answer_key, style)


def _add_authoring_appendix(document: DocumentType, project: ProjectDetail, include_review_comments: bool) -> None:
    media_details = [a for a in project.storyboard.assets if a.kind != "image" and (a.transcript or a.captionsVtt)]
    if media_details:
        document.add_page_break()
        document.add_heading("Media transcripts and captions", level=1)
        for asset in media_details:
            document.add_heading(asset.filename, level=2)
            document.add_paragraph(f"Media ID: {asset.id}")
            for scene in project.storyboard.scenes:
                placement = next((p for p in scene.mediaPlacements if p.assetId == asset.id), None)
                if placement:
                    document.add_paragraph(f"Scene {scene.title}: {placement.role}, {placement.position} content, order {placement.order}; loop {'on' if placement.loop else 'off'}, start muted {'yes' if placement.muted else 'no'}.")
            if asset.transcript:
                document.add_paragraph(asset.transcript)
            if asset.captionsVtt:
                document.add_paragraph(f"Captions ({asset.language}):\n{asset.captionsVtt}")
    resources = list_reference_resources(project.id)
    proposals = [p for p in list_enhancement_proposals(project.id) if p.status == "accepted" and p.referenceCitations]
    if proposals:
        document.add_page_break()
        document.add_heading("Source citations (accepted enhancement provenance)", level=1)
        document.add_paragraph("These citations identify the exact source versions behind accepted suggestions. Later author edits may require renewed source review.")
        for proposal in proposals:
            document.add_paragraph(f"Proposal {proposal.id} | source revision {proposal.sourceRevision}")
            for citation in proposal.referenceCitations:
                resource = next((r for r in resources if r.id == citation.resourceId), None)
                version = next((v for v in resource.versions if v.version == citation.version), None) if resource else None
                label = f"{resource.title if resource else citation.resourceId} | version {citation.version} | {citation.locator}"
                label += f" | SHA-256 {version.sha256}" if version else " | source version unavailable; verify before handoff"
                document.add_paragraph(str(_clean_xml_value(label)))
                if citation.excerpt:
                    document.add_paragraph(str(_clean_xml_value(citation.excerpt)))
    if include_review_comments:
        document.add_page_break()
        document.add_heading("Review discussion (authoring material)", level=1)
        for comment in list_review_comments(project.id):
            status = "resolved" if comment.resolved else "unresolved"
            context = " | outdated/deleted scene context" if comment.outdated or comment.sceneDeleted else ""
            document.add_heading(str(_clean_xml_value(f"{comment.sceneTitle} | revision {comment.revision} | {status}{context}")), level=2)
            document.add_paragraph(str(_clean_xml_value(f"{comment.authorLabel}: {comment.body}")))
            for reply in comment.replies:
                document.add_paragraph(str(_clean_xml_value(f"Reply — {reply.authorLabel}: {reply.body}")))


@with_scene_images
def build_sme_storyboard_docx(project: ProjectDetail, include_answer_key: bool = False, include_review_comments: bool = False) -> bytes:
    project = _normalized_project(project)
    dimensions = _template_slide_dimensions()
    template = Presentation(str(SME_TEMPLATE_PATH))
    style = _extract_sme_visual_style(template)
    document = Document()
    document.core_properties.title = f"{project.title} - SME Master"
    document.core_properties.subject = f"Editable SME storyboard revision {project.storyboard.revision}"
    section = document.sections[0]
    _set_widescreen_section(section, *dimensions)
    with tempfile.TemporaryDirectory(prefix="chimera-master-") as temporary:
        artwork = _extract_master_artwork(SME_TEMPLATE_PATH, 1, Path(temporary))
        if artwork:
            _set_section_artwork(section, artwork)
        for index, scene in enumerate(project.storyboard.scenes, start=1):
            _add_sme_scene(document, project, scene, index, include_answer_key, style)
        _add_authoring_appendix(document, project, include_review_comments)
        output = io.BytesIO()
        document.save(output)
        return output.getvalue()


@with_scene_images
def build_full_storyboard_docx(project: ProjectDetail, include_answer_key: bool = False, include_review_comments: bool = False) -> bytes:
    project = _normalized_project(project)
    if not FULL_TEMPLATE_PATH.is_file():
        raise FileNotFoundError(f"Full storyboard template was not found: {FULL_TEMPLATE_PATH.name}")
    dimensions = _template_slide_dimensions()
    template = Presentation(str(SME_TEMPLATE_PATH))
    style = _extract_sme_visual_style(template)
    document = Document(str(FULL_TEMPLATE_PATH))
    document.core_properties.title = f"{project.title} - Full Storyboard"
    document.core_properties.subject = f"Multimedia instructional storyboard revision {project.storyboard.revision}"
    template_section = document.sections[-1]
    portrait_settings = (
        template_section.page_width,
        template_section.page_height,
        template_section.left_margin,
        template_section.right_margin,
        template_section.top_margin,
        template_section.bottom_margin,
    )
    with tempfile.TemporaryDirectory(prefix="chimera-master-") as temporary:
        artwork = _extract_master_artwork(SME_TEMPLATE_PATH, 1, Path(temporary))
        _append_sme_storyboard(document, project, include_answer_key, dimensions, style, artwork)
    details_section = document.add_section(WD_SECTION.NEW_PAGE)
    (
        details_section.page_width,
        details_section.page_height,
        details_section.left_margin,
        details_section.right_margin,
        details_section.top_margin,
        details_section.bottom_margin,
    ) = portrait_settings
    details_section.header.is_linked_to_previous = False
    document.add_heading("Detailed Chimera scene records", level=0)
    document.add_paragraph(f"{project.title} | Storyboard revision {project.storyboard.revision}")
    if include_answer_key:
        notice = document.add_paragraph()
        notice.add_run("Contains answer keys for authoring use only.").bold = True
    for index, scene in enumerate(project.storyboard.scenes, start=1):
        _add_scene(document, project, scene, index, include_answer_key)
    _add_authoring_appendix(document, project, include_review_comments)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


@with_scene_images
def build_storyboard_docx(project: ProjectDetail, include_answer_key: bool = False, include_review_comments: bool = False) -> bytes:
    project = _normalized_project(project)
    document = Document()
    document.core_properties.title = f"{project.title} - Storyboard"
    document.core_properties.subject = f"Editable Chimera storyboard revision {project.storyboard.revision}"
    document.add_heading(project.title, level=0)
    document.add_paragraph(f"Storyboard revision: {project.storyboard.revision}")
    document.add_paragraph(f"Scenes: {len(project.storyboard.scenes)}")
    if project.storyboard.description:
        document.add_paragraph(project.storyboard.description)
    if include_answer_key:
        notice = document.add_paragraph()
        notice.add_run("Contains answer keys for authoring use only.").bold = True
    for index, scene in enumerate(project.storyboard.scenes, start=1):
        _add_scene(document, project, scene, index, include_answer_key)
    _add_authoring_appendix(document, project, include_review_comments)

    output = io.BytesIO()
    document.save(output)
    return output.getvalue()
