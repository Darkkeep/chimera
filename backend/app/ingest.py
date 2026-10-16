import hashlib
from pathlib import Path
from typing import Any, Literal
from xml.etree import ElementTree

from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

from .models import ExtractedAsset, Scene
from .pptx_backgrounds import render_master_background
from .pptx_layout import extract_master_layout, extract_slide_layout


IMAGE_EXTENSIONS = {".bmp", ".emf", ".gif", ".jpeg", ".jpg", ".png", ".svg", ".tif", ".tiff", ".wmf"}
AUDIO_EXTENSIONS = {".aac", ".m4a", ".mid", ".midi", ".mp3", ".oga", ".ogg", ".wav", ".wma"}
ASSET_DIRECTORIES = {"image": "images", "audio": "audio"}
MIME_EXTENSIONS = {
    "image/jpeg": ".jpg",
    "image/png": ".png",
    "image/gif": ".gif",
    "image/bmp": ".bmp",
    "audio/aac": ".aac",
    "audio/mp4": ".m4a",
    "audio/midi": ".mid",
    "audio/mpeg": ".mp3",
    "audio/ogg": ".ogg",
    "audio/wav": ".wav",
    "audio/x-wav": ".wav",
    "audio/x-ms-wma": ".wma",
}


def _media_kind(relationship: Any) -> Literal["image", "audio"] | None:
    if relationship.is_external:
        return None
    part = relationship.target_part
    content_type = part.content_type.lower()
    extension = Path(str(part.partname)).suffix.lower()
    if relationship.reltype == RT.IMAGE or content_type.startswith("image/"):
        return "image"
    if relationship.reltype == RT.AUDIO or content_type.startswith("audio/"):
        return "audio"
    if relationship.reltype == RT.MEDIA and extension in AUDIO_EXTENSIONS:
        return "audio"
    return None


def _asset_from_relationship(relationship: Any, kind: Literal["image", "audio"], output_root: Path) -> ExtractedAsset:
    part = relationship.target_part
    content = part.blob
    digest = hashlib.sha256(content).hexdigest()
    content_type = part.content_type.lower()
    extension = Path(str(part.partname)).suffix.lower()
    valid_extensions = IMAGE_EXTENSIONS if kind == "image" else AUDIO_EXTENSIONS
    if extension not in valid_extensions:
        extension = MIME_EXTENSIONS.get(content_type, ".bin")
    asset_id = f"{kind}-{digest}"
    filename = f"{asset_id}{extension}"
    relative_path = Path("assets") / ASSET_DIRECTORIES[kind] / filename
    destination = output_root / relative_path
    destination.parent.mkdir(parents=True, exist_ok=True)
    if not destination.exists():
        destination.write_bytes(content)
    return ExtractedAsset(
        id=asset_id,
        kind=kind,
        filename=filename,
        contentType=content_type,
        sizeBytes=len(content),
        sha256=digest,
        path=relative_path.as_posix(),
    )


def extract_presentation(source: Path, output_root: Path) -> tuple[list[Scene], list[ExtractedAsset]]:
    presentation = Presentation(str(source))
    scenes = []
    assets_by_id: dict[str, ExtractedAsset] = {}
    for index, slide in enumerate(presentation.slides, start=1):
        texts = []
        for shape in slide.shapes:
            text = getattr(shape, "text", "").strip()
            if text:
                texts.append(text)
        title = texts[0][:120] if texts else f"Slide {index}"
        notes_frame = slide.notes_slide.notes_text_frame
        image_assets = []
        audio_assets = []
        for relationship in slide.part.rels.values():
            kind = _media_kind(relationship)
            if kind is None:
                continue
            asset = _asset_from_relationship(relationship, kind, output_root)
            assets_by_id[asset.id] = asset
            asset_refs = image_assets if kind == "image" else audio_assets
            if asset.id not in asset_refs:
                asset_refs.append(asset.id)
        def register_image(relationship: Any) -> str:
            asset = _asset_from_relationship(relationship, "image", output_root)
            assets_by_id[asset.id] = asset
            if asset.id not in image_assets:
                image_assets.append(asset.id)
            return asset.id

        visual_layout = extract_slide_layout(presentation, slide, register_image)
        slide_xml = ElementTree.fromstring(slide._element.xml)
        timing = slide_xml.find(qn("p:timing"))
        transition = slide_xml.find(qn("p:transition"))
        scenes.append(
            Scene(
                id=f"scene-{index:03d}",
                sourceSlide=index,
                title=title,
                extractedText=texts[1:],
                backgroundAsset=visual_layout.backgroundAsset,
                visualLayout=visual_layout,
                imageAssets=image_assets,
                audioAssets=audio_assets,
                notes=notes_frame.text.strip() if notes_frame else "",
                sourceAnimationXml=ElementTree.tostring(timing, encoding="unicode") if timing is not None else None,
                sourceTransitionXml=ElementTree.tostring(transition, encoding="unicode") if transition is not None else None,
            )
        )
    def register_master_image(relationship: Any) -> str:
        asset = _asset_from_relationship(relationship, "image", output_root)
        assets_by_id[asset.id] = asset
        return asset.id

    for index, master in enumerate(presentation.slide_masters, start=1):
        master_layout = extract_master_layout(presentation, master, register_master_image)
        content = render_master_background(master_layout, assets_by_id, output_root)
        digest = hashlib.sha256(content).hexdigest()
        asset_id = f"master-background-{digest}"
        if asset_id in assets_by_id:
            continue
        relative_path = Path("assets") / "images" / f"{asset_id}.svg"
        destination = output_root / relative_path
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_bytes(content)
        assets_by_id[asset_id] = ExtractedAsset(
            id=asset_id, kind="image", role="background", filename=f"Master background {index}.svg",
            contentType="image/svg+xml", sizeBytes=len(content), sha256=digest, path=relative_path.as_posix(),
        )
    return scenes, list(assets_by_id.values())
