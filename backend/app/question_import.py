"""Bounded, offline question interchange.

CSV text beginning with a spreadsheet formula prefix (or an apostrophe) is
escaped with one apostrophe. Import reverses that escape; XLSX uses literal
string cells instead. Neither format includes scene-level assessment settings.
"""

import base64
import csv
import json
import posixpath
import re
from io import BytesIO, StringIO
from urllib.parse import unquote, urlsplit
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from defusedxml import ElementTree as SafeET
from openpyxl import Workbook, load_workbook
from pydantic import TypeAdapter

from .models import AssessmentItem, AssessmentSettings, QuestionMetadata, Scene


MAX_UPLOAD = 10 * 1024 * 1024
MAX_UNCOMPRESSED = 30 * 1024 * 1024
MAX_QUESTIONS = 500
MAX_ZIP_ENTRIES = 2000
QTI = "http://www.imsglobal.org/xsd/imsqti_v2p1"
CP = "http://www.imsglobal.org/xsd/imscp_v1p1"
XML_LANG = "{http://www.w3.org/XML/1998/namespace}lang"
COLUMNS = [
    "identifier", "type", "title", "prompt", "choices", "secondaryChoices",
    "answers", "correctFeedback", "incorrectFeedback", "caseSensitive",
    "description", "tags", "author", "source", "language", "difficulty",
    "learningObjectiveRefs",
]
REQUIRED = {"identifier", "type", "title", "prompt", "answers"}
ITEM_ADAPTER = TypeAdapter(AssessmentItem)
SETTINGS_WARNING = "Scene assessment settings are not included; default settings will be used."
TEST_WARNING = (
    "Package test-level settings, navigation, sequencing, selection, timing, "
    "and aggregate scoring are not preserved."
)


def _row(source: str, *, title: str = "", item: dict | None = None,
         errors: list[str] | None = None, warnings: list[str] | None = None) -> dict:
    return {"index": 0, "source": source, "title": title, "item": item,
            "errors": errors or [], "warnings": warnings or []}


def _validate_item(item: dict) -> dict:
    return ITEM_ADAPTER.validate_python(item).model_dump(mode="json")


def _json(value: str, label: str, expected: type = list):
    try:
        result = json.loads(value)
    except (ValueError, TypeError) as exc:
        raise ValueError(f"{label} must contain valid JSON") from exc
    if not isinstance(result, expected):
        raise ValueError(f"{label} must be a JSON {expected.__name__}")
    return result


def _formula_prefix(value: str) -> bool:
    return value.startswith(("'", "\t", "\r", "\n")) or value.lstrip().startswith(("=", "+", "-", "@"))


def _csv_escape(value: str) -> str:
    return "'" + value if _formula_prefix(value) else value


def _csv_unescape(value: str) -> str:
    return value[1:] if value.startswith("'") and _formula_prefix(value[1:]) else value


def _spreadsheet_item(values: dict[str, str]) -> dict:
    for key in REQUIRED - {"title"}:
        if not values.get(key, "").strip():
            raise ValueError(f"{key} must not be empty")
    item = {key: values[key] for key in ("identifier", "type", "prompt")}
    metadata = {key: values.get(key, "") for key in (
        "title", "description", "author", "source", "difficulty",
    )}
    metadata["language"] = values.get("language") or "en"
    for key in ("tags", "learningObjectiveRefs"):
        metadata[key] = _json(values[key], key) if values.get(key) else []
        if any(not isinstance(value, str) for value in metadata[key]):
            raise ValueError(f"{key} must be an array of strings")
    item["metadata"] = QuestionMetadata.model_validate(metadata).model_dump(mode="json")
    for key in ("correctFeedback", "incorrectFeedback"):
        if key in values:
            item[key] = values[key]
    answers = _json(values["answers"], "answers")
    kind = item["type"]
    if kind in {"single-choice", "multiple-response", "ordering", "drag-drop", "matching"}:
        choices = _json(values.get("choices", ""), "choices")
        for choice in choices:
            if not isinstance(choice, dict) or set(choice) != {"identifier", "text"}:
                raise ValueError("choices must contain only {identifier,text} objects")
        item[{"drag-drop": "draggables", "matching": "leftItems"}.get(kind, "choices")] = choices
    elif values.get("choices"):
        raise ValueError(f"choices is not supported for {kind}")
    if kind in {"drag-drop", "matching"}:
        secondary = _json(values.get("secondaryChoices", ""), "secondaryChoices")
        for choice in secondary:
            if not isinstance(choice, dict) or set(choice) != {"identifier", "text"}:
                raise ValueError("secondaryChoices must contain only {identifier,text} objects")
        item["targets" if kind == "drag-drop" else "rightItems"] = secondary
        if any(not isinstance(pair, dict) or set(pair) != {"sourceIdentifier", "targetIdentifier"} for pair in answers):
            raise ValueError("answers must contain only {sourceIdentifier,targetIdentifier} pairs")
        item["correctMatches"] = answers
    else:
        if values.get("secondaryChoices"):
            raise ValueError(f"secondaryChoices is not supported for {kind}")
        if kind in {"single-choice", "true-false"}:
            if len(answers) != 1:
                raise ValueError("answers must contain exactly one answer")
            answer = answers[0]
            if kind == "true-false" and isinstance(answer, bool):
                answer = str(answer).lower()
            item["correctResponse"] = answer
        elif kind == "multiple-response":
            item["correctResponses"] = answers
        elif kind == "ordering":
            item["correctOrder"] = answers
        elif kind == "short-answer":
            item["acceptedAnswers"] = answers
    case = values.get("caseSensitive", "")
    if case:
        if kind != "short-answer" or case.lower() not in {"true", "false"}:
            raise ValueError("caseSensitive must be true/false and is only supported for short-answer")
        item["caseSensitive"] = case.lower() == "true"
    return _validate_item(item)


def _table_rows(filename: str, records, *, xlsx: bool = False) -> list[dict]:
    try:
        header_cells = next(records)
    except StopIteration:
        raise ValueError("The question table is empty")
    headers = []
    if len(header_cells) > 100:
        raise ValueError("Question tables may contain at most 100 columns")
    for cell in header_cells:
        if xlsx:
            if cell.data_type == "f":
                raise ValueError("Formulas are not allowed in headers")
            cell = cell.value
        if cell is not None and not isinstance(cell, str):
            raise ValueError("Headers must be strings")
        headers.append((cell or "").strip())
    nonempty = [header for header in headers if header]
    if len(nonempty) != len(set(nonempty)):
        raise ValueError("Duplicate column headers are not allowed")
    missing = REQUIRED - set(headers)
    if missing:
        raise ValueError("Missing required headers: " + ", ".join(sorted(missing)))
    rows = []
    for number, cells in enumerate(records, start=2):
        if number > MAX_QUESTIONS + 1:
            raise ValueError(f"At most {MAX_QUESTIONS} question rows are allowed")
        values = []
        errors = []
        for cell in cells:
            if xlsx:
                if cell.data_type == "f":
                    errors.append("Formulas are not allowed, including formulas with cached results")
                cell = cell.value
            if cell is not None and not isinstance(cell, (str, bool, int, float)):
                errors.append("Cells must contain plain text")
            value = "" if cell is None else str(cell)
            values.append(value if xlsx else _csv_unescape(value))
        if not any(values) and not errors:
            continue
        if not xlsx and len(values) != len(headers):
            errors.append("Malformed row: column count differs from the header")
        if len(values) > len(headers) and any(values[len(headers):]):
            errors.append("Values without column headers are not allowed")
        mapped = dict(zip(headers, values))
        for key, value in zip(headers, values):
            if key not in COLUMNS and value:
                errors.append(f"Unsupported nonempty column: {key or '(blank header)'}")
        title = mapped.get("title", "")
        item = None
        if not errors:
            try:
                item = _spreadsheet_item(mapped)
            except (ValueError, TypeError) as exc:
                errors.append(str(exc))
        rows.append(_row(f"{filename}:row {number}", title=title, item=item,
                         errors=errors, warnings=[SETTINGS_WARNING]))
    if not rows:
        raise ValueError("The question table contains no questions")
    return rows


def _checked_zip(content: bytes) -> ZipFile:
    archive = ZipFile(BytesIO(content))
    try:
        entries = archive.infolist()
        if len(entries) > MAX_ZIP_ENTRIES:
            raise ValueError(f"ZIP archives may contain at most {MAX_ZIP_ENTRIES} entries")
        names = set()
        total = 0
        for entry in entries:
            name = entry.orig_filename
            if (not name or name != entry.filename or "\\" in name or "\x00" in name or ":" in name
                    or name.startswith("/") or any(part in {"..", "."} for part in name.rstrip("/").split("/"))
                    or "//" in name):
                raise ValueError(f"Unsafe ZIP member path: {name}")
            if name in names:
                raise ValueError(f"Duplicate ZIP member: {name}")
            names.add(name)
            if entry.flag_bits & 1:
                raise ValueError("Encrypted ZIP members are not supported")
            if (entry.external_attr >> 16) & 0o170000 == 0o120000:
                raise ValueError("ZIP symbolic links are not supported")
            total += entry.file_size
            if total > MAX_UNCOMPRESSED:
                raise ValueError("ZIP uncompressed size exceeds 30 MB")
        return archive
    except Exception:
        archive.close()
        raise


def _xml(content: bytes):
    if len(content) > MAX_UNCOMPRESSED:
        raise ValueError("XML exceeds the uncompressed size limit")
    return SafeET.fromstring(content, forbid_dtd=True, forbid_entities=True, forbid_external=True)


def _q(name: str) -> str:
    return f"{{{QTI}}}{name}"


def _local(node) -> str:
    return node.tag.rsplit("}", 1)[-1]


def _plain(node) -> str:
    if len(node):
        raise ValueError(f"Rich text or media in {_local(node)} is not supported")
    return node.text or ""


def _attributes(node, allowed: set[str]):
    unknown = set(node.attrib) - allowed
    if unknown:
        raise ValueError(f"Unsupported {_local(node)} attributes: {', '.join(sorted(unknown))}")


def _extension(content: bytes) -> dict:
    # Inspect PIs only, never attempt to parse their payload as XML.
    matches = re.findall(rb"<\?chimera(?:\s+([^?]*))?\?>", content)
    if not matches:
        return {}
    if len(matches) != 1:
        raise ValueError("Multiple Chimera metadata processing instructions")
    root_position = re.search(rb"<(?:[A-Za-z_][\w.-]*:)?assessmentItem\b", content)
    if root_position is None or content.find(b"<?chimera") > root_position.start():
        raise ValueError("Chimera metadata must precede the assessment item")
    try:
        result = json.loads(base64.b64decode(matches[0].strip(), validate=True).decode("utf-8"))
    except (ValueError, UnicodeError) as exc:
        raise ValueError("Invalid Chimera metadata encoding") from exc
    if not isinstance(result, dict) or set(result) - {"metadata", "type", "settings", "dragSources"}:
        raise ValueError("Unsupported Chimera metadata fields")
    if "metadata" in result:
        if not isinstance(result["metadata"], dict):
            raise ValueError("Chimera metadata must be an object")
        result["metadata"] = QuestionMetadata.model_validate(result["metadata"]).model_dump(mode="json")
    if result.get("settings") is not None:
        if not isinstance(result["settings"], dict):
            raise ValueError("Chimera settings must be an object")
        if set(result["settings"]) - set(AssessmentSettings.model_fields):
            raise ValueError("Unknown Chimera assessment settings")
        result["settings"] = AssessmentSettings.model_validate(result["settings"]).model_dump(mode="json")
    if "dragSources" in result and (
        not isinstance(result["dragSources"], list)
        or any(not isinstance(value, str) for value in result["dragSources"])
        or len(result["dragSources"]) != len(set(result["dragSources"]))
    ):
        raise ValueError("dragSources must be an array of unique identifiers")
    return result


def _signature(node):
    return (node.tag, tuple(sorted(node.attrib.items())), (node.text or "").strip(),
            tuple(_signature(child) for child in node))


def _processing(root, response_id: str, kind: str) -> tuple[list[str] | None, bool]:
    processing = root.findall(_q("responseProcessing"))
    if len(processing) != 1:
        raise ValueError("Exactly one exact-match responseProcessing is required")
    processing = processing[0]
    template = processing.get("template")
    if template:
        if (template not in {
            "http://www.imsglobal.org/question/qti_v2p1/rptemplates/match_correct",
            "https://www.imsglobal.org/question/qti_v2p1/rptemplates/match_correct",
        } or len(processing) or set(processing.attrib) != {"template"}):
            raise ValueError("Unsupported response processing template or custom scoring")
        if kind == "short-answer":
            raise ValueError("Short-answer match_correct string semantics cannot be preserved")
        return None, False
    expected = ET.Element(_q("responseProcessing"))
    condition = ET.SubElement(expected, _q("responseCondition"))
    response_if = ET.SubElement(condition, _q("responseIf"))
    answers = None
    case_sensitive = False
    if kind == "short-answer":
        conditions = processing.findall(f"./{_q('responseCondition')}/{_q('responseIf')}")
        if len(conditions) != 1 or not len(conditions[0]):
            raise ValueError("Unsupported short-answer scoring")
        expression = conditions[0][0]
        if expression.tag != _q("or") or not len(expression):
            raise ValueError("Short answers require canonical exact stringMatch processing")
        answers = []
        cases = set()
        match_or = ET.SubElement(response_if, _q("or"))
        for candidate in expression:
            if candidate.tag != _q("stringMatch") or candidate.get("caseSensitive") not in {"true", "false"}:
                raise ValueError("Unsupported short-answer comparison")
            cases.add(candidate.get("caseSensitive"))
            base = candidate.find(_q("baseValue"))
            if base is None:
                raise ValueError("Missing accepted answer")
            answer = _plain(base)
            answers.append(answer)
            match = ET.SubElement(match_or, _q("stringMatch"), {"caseSensitive": candidate.get("caseSensitive")})
            ET.SubElement(match, _q("variable"), {"identifier": response_id})
            ET.SubElement(match, _q("baseValue"), {"baseType": "string"}).text = answer
        if len(cases) != 1:
            raise ValueError("Mixed short-answer case sensitivity is not supported")
        case_sensitive = cases == {"true"}
    else:
        match = ET.SubElement(response_if, _q("match"))
        ET.SubElement(match, _q("variable"), {"identifier": response_id})
        ET.SubElement(match, _q("correct"), {"identifier": response_id})
    for branch, score, feedback in (
        (response_if, "1", "correct"),
        (ET.SubElement(condition, _q("responseElse")), "0", "incorrect"),
    ):
        for name, base_type, text in (("SCORE", "float", score), ("FEEDBACK", "identifier", feedback)):
            outcome = ET.SubElement(branch, _q("setOutcomeValue"), {"identifier": name})
            ET.SubElement(outcome, _q("baseValue"), {"baseType": base_type}).text = text
    if _signature(processing) != _signature(expected):
        raise ValueError("Unsupported custom scoring; only canonical exact-match processing is supported")
    return answers, case_sensitive


def _parse_qti(content: bytes, source: str) -> dict:
    warnings = []
    title = ""
    try:
        root = _xml(content)
        if root.tag != _q("assessmentItem"):
            raise ValueError("Expected a QTI 2.1 assessmentItem; assessmentTest requires a ZIP package")
        title = root.get("title", "")
        _attributes(root, {"identifier", "title", "adaptive", "timeDependent", XML_LANG, "label",
                           "{http://www.w3.org/2001/XMLSchema-instance}schemaLocation"})
        for flag in ("adaptive", "timeDependent"):
            if root.get(flag, "false") not in {"false", "0"}:
                raise ValueError(f"{flag} items are not supported")
        if root.get("label"):
            warnings.append("The QTI item label is not preserved.")
        extension = _extension(content)
        allowed_root = {"responseDeclaration", "outcomeDeclaration", "itemBody", "responseProcessing", "modalFeedback"}
        for child in root:
            if child.tag != _q(_local(child)) or _local(child) not in allowed_root:
                if _local(child) in {"metadata", "qtiMetadata"}:
                    warnings.append("Standard QTI metadata fields are not preserved; only Chimera metadata is imported.")
                else:
                    raise ValueError(f"Unsupported QTI item element: {_local(child)}")
        for node in root.iter():
            if node.text and node.text.strip() and len(node) and node.tag not in {_q("itemBody")}:
                raise ValueError(f"Unsupported mixed text in {_local(node)}")
            if node.tail and node.tail.strip():
                raise ValueError("Unsupported mixed text outside a plain-text element")
        declarations = root.findall(_q("responseDeclaration"))
        bodies = root.findall(_q("itemBody"))
        if len(declarations) != 1 or len(bodies) != 1:
            raise ValueError("Exactly one responseDeclaration and itemBody are required")
        declaration = declarations[0]
        _attributes(declaration, {"identifier", "cardinality", "baseType"})
        if any(child.tag != _q("correctResponse") for child in declaration):
            raise ValueError("Response mappings, defaults, and custom scoring are not supported")
        correct_nodes = declaration.findall(_q("correctResponse"))
        if len(correct_nodes) > 1:
            raise ValueError("Duplicate correctResponse")
        correct = []
        for container in correct_nodes:
            _attributes(container, set())
            for value in container:
                if value.tag != _q("value") or value.attrib:
                    raise ValueError("Unsupported answer-key representation")
                correct.append(_plain(value))
        body = bodies[0]
        _attributes(body, set())
        interactions = [node for node in body.iter() if _local(node).endswith("Interaction")]
        if len(interactions) != 1:
            raise ValueError("Exactly one supported interaction (one fill-in-the-blank blank) is required")
        interaction = interactions[0]
        interaction_name = _local(interaction)
        if interaction.tag != _q(interaction_name):
            raise ValueError("Unsupported interaction namespace")
        response_id = declaration.get("identifier")
        if not response_id or interaction.get("responseIdentifier") != response_id:
            raise ValueError("Interaction responseIdentifier does not match the answer key")
        interaction_attributes = {
            "choiceInteraction": {"shuffle", "maxChoices", "minChoices"},
            "orderInteraction": {"shuffle"},
            "associateInteraction": {"shuffle", "maxAssociations", "minAssociations"},
            "matchInteraction": {"shuffle", "maxAssociations", "minAssociations"},
            "textEntryInteraction": {"expectedLength"},
        }
        _attributes(interaction, {"responseIdentifier"} | interaction_attributes.get(interaction_name, set()))
        if interaction.get("shuffle", "false") not in {"true", "1", "false", "0"}:
            raise ValueError("Invalid shuffle boolean")
        if interaction.get("shuffle", "false") not in {"false", "0"}:
            warnings.append("Interaction shuffle is not preserved; choices retain their document order.")
        if interaction.get("minChoices", "0") != "0" or interaction.get("minAssociations", "0") != "0":
            raise ValueError("Minimum response constraints cannot be preserved")
        prompt_nodes = interaction.findall(_q("prompt"))
        if interaction_name == "textEntryInteraction":
            if len(interaction) or (interaction.text or "").strip():
                raise ValueError("Unsupported text-entry content")
            paragraphs = [node for node in body if node.tag == _q("p")]
            if len(paragraphs) != 1 or (body.text or "").strip():
                raise ValueError("Fill-in-the-blank requires one plain-text prompt paragraph")
            paragraph = paragraphs[0]
            _attributes(paragraph, set())
            if len(body) == 1 and len(paragraph) == 2 and paragraph[1] is interaction:
                span = paragraph[0]
                if span.tag != _q("span") or (paragraph.text or "").strip():
                    raise ValueError("Inline fill-in-the-blank requires a plain-text span followed by one blank")
                _attributes(span, set())
                prompt = _plain(span)
            elif len(body) == 2 and body[0] is paragraph and body[1] is interaction:
                # Compatibility with the original exporter, which placed the
                # text-entry interaction directly after the prompt paragraph.
                prompt = _plain(paragraph)
            else:
                raise ValueError("Unsupported fill-in-the-blank layout; only one trailing blank is supported")
        else:
            if len(body) != 1 or (body.text or "").strip() or len(prompt_nodes) != 1:
                raise ValueError("Interaction requires one plain-text prompt and no other body content")
            _attributes(prompt_nodes[0], set())
            prompt = _plain(prompt_nodes[0])

        def choices(container, tag: str) -> list[dict]:
            result = []
            for node in container:
                if node.tag == _q("prompt"):
                    continue
                if node.tag != _q(tag):
                    raise ValueError(f"Unsupported interaction content: {_local(node)}")
                _attributes(node, {"identifier", "fixed", "matchMax", "matchMin"})
                if node.get("fixed", "false") not in {"false", "0"}:
                    raise ValueError("Fixed choices are not supported")
                if node.get("matchMax", "1") != "1" or node.get("matchMin", "0") != "0":
                    raise ValueError("Only one-to-one matching choices are supported")
                result.append({"identifier": node.get("identifier"), "text": _plain(node)})
            return result

        item = {"identifier": root.get("identifier"), "prompt": prompt}
        cardinality = declaration.get("cardinality")
        if interaction_name == "choiceInteraction":
            options = choices(interaction, "simpleChoice")
            maximum = interaction.get("maxChoices", "1")
            if cardinality == "single" and maximum != "1":
                raise ValueError("Single-choice maxChoices must be 1")
            if cardinality == "multiple" and maximum not in {"0", str(len(options))}:
                raise ValueError("Restricted multiple-response selection counts cannot be preserved")
            if cardinality not in {"single", "multiple"}:
                raise ValueError("Unsupported choice cardinality")
            kind = "multiple-response" if cardinality == "multiple" else "single-choice"
            if extension.get("type") == "true-false":
                if cardinality != "single" or options != [
                    {"identifier": "true", "text": "True"}, {"identifier": "false", "text": "False"},
                ]:
                    raise ValueError("True/false choices must be the canonical True and False options")
                kind = "true-false"
            if kind != "true-false":
                item["choices"] = options
            if kind == "multiple-response":
                item["correctResponses"] = correct
            else:
                if len(correct) != 1:
                    raise ValueError("Single choice requires exactly one correct answer")
                item["correctResponse"] = correct[0]
            expected = (cardinality, "identifier")
        elif interaction_name == "orderInteraction":
            kind = "ordering"
            item.update(choices=choices(interaction, "simpleChoice"), correctOrder=correct)
            expected = ("ordered", "identifier")
        elif interaction_name in {"associateInteraction", "matchInteraction"}:
            pairs = []
            for answer in correct:
                identifiers = answer.split()
                if len(identifiers) != 2:
                    raise ValueError("directedPair answers must contain two identifiers")
                pairs.append(dict(zip(("sourceIdentifier", "targetIdentifier"), identifiers)))
            if interaction.get("maxAssociations", "0") not in {"0", str(len(pairs))}:
                raise ValueError("Association limits cannot be preserved")
            item["correctMatches"] = pairs
            expected = ("multiple", "directedPair")
            if interaction_name == "associateInteraction":
                kind = "drag-drop"
                options = choices(interaction, "simpleAssociableChoice")
                source_ids = extension.get("dragSources")
                if source_ids is None:
                    source_ids = [pair["sourceIdentifier"] for pair in pairs]
                    target_ids = {pair["targetIdentifier"] for pair in pairs}
                    if (set(source_ids) & target_ids or set(source_ids) | target_ids != {option["identifier"] for option in options}):
                        raise ValueError("Ambiguous associateInteraction: unassigned choices require Chimera dragSources metadata")
                if not set(source_ids).issubset({option["identifier"] for option in options}):
                    raise ValueError("dragSources references unknown choices")
                item["draggables"] = [option for option in options if option["identifier"] in source_ids]
                item["targets"] = [option for option in options if option["identifier"] not in source_ids]
            else:
                kind = "matching"
                sets = interaction.findall(_q("simpleMatchSet"))
                if len(sets) != 2 or len(interaction) != 3:
                    raise ValueError("Matching requires exactly two simpleMatchSets")
                for match_set in sets:
                    _attributes(match_set, set())
                item["leftItems"] = choices(sets[0], "simpleAssociableChoice")
                item["rightItems"] = choices(sets[1], "simpleAssociableChoice")
        elif interaction_name == "textEntryInteraction":
            kind = "short-answer"
            expected = ("single", "string")
        else:
            raise ValueError(f"Unsupported QTI interaction: {interaction_name}")
        if (cardinality, declaration.get("baseType")) != expected:
            raise ValueError("Response cardinality/baseType does not match the interaction")
        if extension.get("type", kind) != kind:
            raise ValueError("Chimera type conflicts with the QTI interaction")
        if "dragSources" in extension and kind != "drag-drop":
            raise ValueError("dragSources is only valid for drag-drop")
        item["type"] = kind
        answers, case_sensitive = _processing(root, response_id, kind)
        if kind == "short-answer":
            if correct and correct != answers and correct != answers[:1]:
                raise ValueError("Short-answer answer key conflicts with string-match scoring")
            item.update(acceptedAnswers=answers, caseSensitive=case_sensitive)
        _validate_outcomes(root)
        seen_feedback = set()
        for feedback in root.findall(_q("modalFeedback")):
            _attributes(feedback, {"outcomeIdentifier", "identifier", "showHide"})
            identifier = feedback.get("identifier")
            if (feedback.get("outcomeIdentifier") != "FEEDBACK"
                    or feedback.get("showHide") != "show" or identifier not in {"correct", "incorrect"}
                    or identifier in seen_feedback):
                raise ValueError("Unsupported conditional feedback")
            seen_feedback.add(identifier)
            item["correctFeedback" if identifier == "correct" else "incorrectFeedback"] = _plain(feedback)
        item["metadata"] = extension.get("metadata", {
            "title": title, "language": root.get(XML_LANG, "en"),
        })
        validated = _validate_item(item)
        row = _row(source, title=validated["metadata"]["title"] or title, item=validated, warnings=warnings)
        if extension.get("settings") is not None:
            row["settings"] = extension["settings"]
            if row["settings"].get("remediationSceneId"):
                row["settings"]["remediationSceneId"] = None
                row["warnings"].append("Remediation scene links are not portable and have been dropped.")
        else:
            row["warnings"].append(SETTINGS_WARNING)
        return row
    except Exception as exc:
        return _row(source, title=title, errors=[str(exc)], warnings=warnings)


def _validate_outcomes(root):
    seen = set()
    for outcome in root.findall(_q("outcomeDeclaration")):
        _attributes(outcome, {"identifier", "cardinality", "baseType", "normalMaximum"})
        identifier = outcome.get("identifier")
        if identifier in seen or identifier not in {"SCORE", "FEEDBACK"}:
            raise ValueError("Unsupported or duplicate outcome declaration")
        seen.add(identifier)
        expected_type = "float" if identifier == "SCORE" else "identifier"
        if outcome.get("cardinality") != "single" or outcome.get("baseType") != expected_type:
            raise ValueError("Unsupported outcome type")
        if "normalMaximum" in outcome.attrib and (identifier != "SCORE" or outcome.get("normalMaximum") != "1"):
            raise ValueError("Non-unit score weights cannot be preserved")
        expected = ET.Element(_q("outcomeDeclaration"), dict(outcome.attrib))
        if len(outcome):
            default = ET.SubElement(expected, _q("defaultValue"))
            ET.SubElement(default, _q("value")).text = "0" if identifier == "SCORE" else "incorrect"
        if _signature(outcome) != _signature(expected):
            raise ValueError("Unsupported outcome default or custom scoring")
    if "SCORE" not in seen:
        raise ValueError("A SCORE outcome declaration is required")


def _resolve(base: str, href: str) -> str:
    parsed = urlsplit(href)
    path = unquote(parsed.path)
    if (not path or parsed.scheme or parsed.netloc or parsed.query or parsed.fragment
            or "\\" in path or "\x00" in path or ":" in path or path.startswith("/")):
        raise ValueError(f"Unsafe package reference: {href}")
    path = posixpath.normpath(posixpath.join(posixpath.dirname(base), path))
    if path == ".." or path.startswith("../") or path.startswith("/"):
        raise ValueError(f"Package reference escapes the archive: {href}")
    return path


def _package(filename: str, archive: ZipFile) -> list[dict]:
    if "imsmanifest.xml" not in archive.namelist():
        raise ValueError("A QTI ZIP must contain imsmanifest.xml identifying its question resources")
    manifest = _xml(archive.read("imsmanifest.xml"))
    if manifest.tag != f"{{{CP}}}manifest":
        raise ValueError("Unsupported package manifest namespace")
    for node in manifest.iter():
        if "{http://www.w3.org/XML/1998/namespace}base" in node.attrib:
            raise ValueError("xml:base package references are not supported")
    resources = manifest.findall(f"./{{{CP}}}resources/{{{CP}}}resource")
    items = []
    tests = []
    warnings = []
    if manifest.find(f".//{{{CP}}}metadata") is not None:
        warnings.append("Package metadata fields are not preserved; item metadata is imported separately.")
    resource_ids = set()
    for resource in resources:
        identifier = resource.get("identifier")
        if not identifier or identifier in resource_ids:
            raise ValueError("Missing or duplicate manifest resource identifier")
        resource_ids.add(identifier)
        kind = resource.get("type")
        if kind not in {"imsqti_item_xmlv2p1", "imsqti_test_xmlv2p1"}:
            continue
        path = _resolve("imsmanifest.xml", resource.get("href", ""))
        if path not in archive.namelist():
            raise ValueError(f"Missing package resource: {path}")
        (items if kind == "imsqti_item_xmlv2p1" else tests).append(path)
    ordered = []
    for path in tests:
        root = _xml(archive.read(path))
        if root.tag != _q("assessmentTest"):
            raise ValueError(f"Expected assessmentTest in {path}")
        for node in root.iter():
            if "{http://www.w3.org/XML/1998/namespace}base" in node.attrib:
                raise ValueError("xml:base test references are not supported")
        warnings.append(TEST_WARNING)
        for reference in root.iter(_q("assessmentItemRef")):
            target = _resolve(path, reference.get("href", ""))
            if target not in items:
                raise ValueError(f"Test reference is not a declared QTI item resource: {target}")
            ordered.append(target)
    ordered.extend(path for path in items if path not in ordered)
    if not ordered:
        raise ValueError("The manifest contains no QTI 2.1 question resources")
    if len(ordered) > MAX_QUESTIONS:
        raise ValueError(f"At most {MAX_QUESTIONS} questions may be imported")
    rows = []
    for path in ordered:
        row = _parse_qti(archive.read(path), f"{filename}:{path}")
        row["warnings"].extend(dict.fromkeys(warnings))
        rows.append(row)
    return rows


def parse_questions(filename: str, content: bytes) -> list[dict]:
    """Return preview rows, reporting invalid input rather than partially dropping it."""
    try:
        if not isinstance(content, bytes):
            raise ValueError("Question content must be bytes")
        if len(content) > MAX_UPLOAD:
            raise ValueError("Upload exceeds the 10 MB limit")
        extension = filename.rsplit(".", 1)[-1].lower()
        if extension == "xml":
            rows = [_parse_qti(content, filename)]
        elif extension == "csv":
            text = content.decode("utf-8-sig")
            rows = _table_rows(filename, iter(csv.reader(StringIO(text, newline=""), strict=True)))
        elif extension in {"zip", "xlsx"}:
            with _checked_zip(content) as archive:
                if extension == "zip":
                    rows = _package(filename, archive)
                else:
                    _check_xlsx_xml(archive)
                    with BytesIO(content) as stream:
                        workbook = load_workbook(stream, read_only=True, data_only=False, keep_links=False)
                        try:
                            if "Questions" in workbook.sheetnames:
                                worksheet = workbook["Questions"]
                            elif len(workbook.sheetnames) == 1:
                                worksheet = workbook.worksheets[0]
                            else:
                                raise ValueError("XLSX must have a Questions worksheet or a sole worksheet")
                            if (worksheet.max_row or 0) > MAX_QUESTIONS + 1 or (worksheet.max_column or 0) > 100:
                                raise ValueError("XLSX worksheet dimensions exceed the row/column limits")
                            worksheet.reset_dimensions()
                            rows = _table_rows(filename, iter(worksheet.iter_rows()), xlsx=True)
                        finally:
                            workbook.close()
        else:
            raise ValueError("Supported question files are .xml, .zip, .csv, and .xlsx")
        if len(rows) > MAX_QUESTIONS:
            raise ValueError(f"At most {MAX_QUESTIONS} questions may be imported")
        seen = set()
        for index, row in enumerate(rows):
            row["index"] = index
            if row["item"]:
                identifier = row["item"]["identifier"]
                if identifier in seen:
                    row["errors"].append(f"Duplicate question identifier: {identifier}")
                    row["item"] = None
                seen.add(identifier)
        return rows
    except Exception as exc:
        return [_row(filename, errors=[str(exc)])]


def _check_xlsx_xml(archive: ZipFile):
    # Do not trust the worksheet's advertised dimensions: a crafted file can
    # omit them or under-report cells and make a read-only iterator unbounded.
    for name in archive.namelist():
        if not name.endswith((".xml", ".rels")):
            continue
        is_worksheet = name.startswith("xl/worksheets/") and name.endswith(".xml")
        with archive.open(name) as stream:
            last_row = 0
            current_row = 0
            last_column = 0
            for event, node in SafeET.iterparse(
                stream, events=("start", "end"), forbid_dtd=True, forbid_entities=True, forbid_external=True,
            ):
                if event == "end":
                    node.clear()
                    continue
                if not is_worksheet:
                    continue
                tag = _local(node)
                if tag == "row":
                    try:
                        number = int(node.get("r", "0"))
                    except ValueError as exc:
                        raise ValueError("Invalid XLSX row reference") from exc
                    if number < 1 or number > MAX_QUESTIONS + 1:
                        raise ValueError("XLSX row limit exceeded")
                    if number <= last_row:
                        raise ValueError("Duplicate or unordered XLSX rows")
                    last_row = current_row = number
                    last_column = 0
                elif tag == "c":
                    reference = re.fullmatch(r"([A-Z]+)([1-9][0-9]*)", node.get("r", ""))
                    if reference is None:
                        raise ValueError("Invalid XLSX cell reference")
                    column = 0
                    for letter in reference[1]:
                        column = column * 26 + ord(letter) - ord("A") + 1
                    if column > 100 or int(reference[2]) > MAX_QUESTIONS + 1:
                        raise ValueError("XLSX row/column limits exceeded")
                    if column <= last_column or int(reference[2]) != current_row:
                        raise ValueError("Duplicate, unordered, or misplaced XLSX cells")
                    last_column = column
                elif tag == "f":
                    raise ValueError("Formulas are not allowed, including formulas with cached results")


def _export_values(item: dict, title: str) -> list[str]:
    metadata = item.get("metadata") or {}
    values = {key: metadata.get(key, "") for key in (
        "title", "description", "author", "source", "language", "difficulty",
    )}
    values["title"] = metadata.get("title", title)
    values["language"] = metadata.get("language", "en")
    values.update({key: item[key] for key in ("identifier", "type", "prompt", "correctFeedback", "incorrectFeedback")})
    for key in ("tags", "learningObjectiveRefs"):
        values[key] = json.dumps(metadata.get(key, []), ensure_ascii=False)
    kind = item["type"]
    for key in ("choices", "secondaryChoices"):
        values[key] = ""
    primary = {"drag-drop": "draggables", "matching": "leftItems"}.get(kind, "choices")
    if primary in item:
        values["choices"] = json.dumps(item[primary], ensure_ascii=False)
    secondary = {"drag-drop": "targets", "matching": "rightItems"}.get(kind)
    if secondary:
        values["secondaryChoices"] = json.dumps(item[secondary], ensure_ascii=False)
    answers = item.get("correctResponses", item.get("correctOrder", item.get("correctMatches", item.get("acceptedAnswers"))))
    if kind in {"single-choice", "true-false"}:
        answers = [item["correctResponse"]]
    values["answers"] = json.dumps(answers, ensure_ascii=False)
    values["caseSensitive"] = str(item["caseSensitive"]).lower() if kind == "short-answer" else ""
    return [str(values.get(column, "")) for column in COLUMNS]


def _write_spreadsheet(rows: list[list[str]], format: str) -> bytes:
    format = format.lower().lstrip(".")
    if format == "csv":
        stream = StringIO(newline="")
        writer = csv.writer(stream)
        writer.writerow(COLUMNS)
        writer.writerows([_csv_escape(value) for value in row] for row in rows)
        return stream.getvalue().encode("utf-8-sig")
    if format == "xlsx":
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Questions"
        for number, row in enumerate([COLUMNS, *rows], start=1):
            for column, value in enumerate(row, start=1):
                if len(value) > 32767:
                    raise ValueError("XLSX cell text exceeds Excel's 32767-character limit")
                cell = worksheet.cell(number, column, value)
                cell.data_type = "s"
        worksheet.freeze_panes = "A2"
        stream = BytesIO()
        workbook.save(stream)
        workbook.close()
        return stream.getvalue()
    raise ValueError("Question spreadsheet format must be csv or xlsx")


def question_spreadsheet(scenes: list[Scene], format: str) -> bytes:
    """Export question data and metadata, not scene-level settings."""
    rows = [
        _export_values(scene.assessmentItem.model_dump(mode="json"), scene.title)
        for scene in scenes if scene.assessmentItem is not None
    ]
    return _write_spreadsheet(rows, format)


def question_template(format: str) -> bytes:
    """Return a literal-string template demonstrating all seven question types."""
    choices = [{"identifier": "a", "text": "First"}, {"identifier": "b", "text": "Second"}]
    secondary = [{"identifier": "x", "text": "First target"}, {"identifier": "y", "text": "Second target"}]
    pairs = [{"sourceIdentifier": "a", "targetIdentifier": "x"},
             {"sourceIdentifier": "b", "targetIdentifier": "y"}]
    examples = [
        {"type": "single-choice", "choices": choices, "correctResponse": "a"},
        {"type": "multiple-response", "choices": choices, "correctResponses": ["a", "b"]},
        {"type": "true-false", "correctResponse": "true"},
        {"type": "ordering", "choices": choices, "correctOrder": ["b", "a"]},
        {"type": "drag-drop", "draggables": choices, "targets": secondary, "correctMatches": pairs},
        {"type": "matching", "leftItems": choices, "rightItems": secondary, "correctMatches": pairs},
        {"type": "short-answer", "acceptedAnswers": ["Example", "Another answer"], "caseSensitive": False},
    ]
    rows = []
    for number, example in enumerate(examples, start=1):
        label = "Fill-in-the-blank" if example["type"] == "short-answer" else example["type"]
        title = f"Example {label}"
        item = _validate_item({
            "identifier": f"example_{number}", "prompt": f"Answer this {label} question.",
            "metadata": {"title": title}, **example,
        })
        rows.append(_export_values(item, title))
    return _write_spreadsheet(rows, format)
