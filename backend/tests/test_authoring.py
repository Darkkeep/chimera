import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from pptx import Presentation
from pptx.util import Inches

from backend.app import main, storage
from backend.app.media import resolved_project
from backend.app.models import ProjectCreate, Scene, SceneCreate, SceneElement, SceneLayout, SceneUpdate
from backend.app.scene_print import render_scene_images
from backend.app.word_export import build_storyboard_docx


class AuthoringTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        root = Path(temporary.name)
        self.enterContext(patch.object(storage, "ROOT", root))
        self.enterContext(patch.object(storage, "PROJECTS", root / "projects"))
        self.project = storage.create_project(ProjectCreate(title="One authored scene"))

    def legacy(self):
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = "Source title"
        source = storage.project_path(self.project.id) / "original" / "source-deck.pptx"
        deck.save(source)
        self.project = storage.update_from_import(self.project.id, [Scene(id="retained", sourceSlide=1,
            title="Edited title", extractedText=["Keep the edited content"], notes="Keep notes", required=False)], [])
        return source

    def test_prepare_is_one_revision_idempotent_and_preserves_source_and_scene_identity(self):
        source = self.legacy()
        source_bytes = source.read_bytes()
        before = self.project
        result = storage.prepare_authored_project(before.id, before.storyboard.revision)
        self.assertEqual(result.storyboard.revision, before.storyboard.revision + 1)
        scene = result.storyboard.scenes[0]
        for field in ["id", "title", "extractedText", "notes", "required", "sourceSlide"]:
            self.assertEqual(getattr(scene, field), getattr(before.storyboard.scenes[0], field))
        self.assertTrue(scene.visualLayout.enabled)
        self.assertIn("Keep the edited content", "\n".join(element.text for element in scene.visualLayout.elements))
        self.assertEqual(source.read_bytes(), source_bytes)
        self.assertEqual(storage.prepare_authored_project(result.id, result.storyboard.revision).storyboard.revision, result.storyboard.revision)
        self.assertTrue(list((storage.project_path(result.id) / "history").glob("*")))
        with self.assertRaises(storage.StoryboardRevisionConflict):
            storage.prepare_authored_project(result.id, before.storyboard.revision)

    def test_missing_source_fails_without_partial_scene_or_revision_changes(self):
        source = self.legacy()
        source.unlink()
        before = storage.read_project(self.project.id).model_dump()
        with self.assertRaisesRegex(ValueError, "retained PowerPoint source"):
            storage.prepare_authored_project(self.project.id, self.project.storyboard.revision)
        self.assertEqual(storage.read_project(self.project.id).model_dump(), before)

    def test_unmapped_source_image_is_not_replaced_with_a_generic_layout(self):
        self.legacy()
        scene = self.project.storyboard.scenes[0].model_copy(update={"sourceSlide": None, "previewImage": "renders/old-slide.png"})
        storage.write_storyboard(storage.project_path(self.project.id), self.project.storyboard.model_copy(update={"scenes": [scene]}), self.project.storyboard.revision)
        before = storage.read_project(self.project.id)
        with self.assertRaisesRegex(ValueError, "Source layout is unavailable"):
            storage.prepare_authored_project(before.id, before.storyboard.revision)
        self.assertEqual(storage.read_project(before.id), before)

    def test_existing_composition_is_authoritative_and_not_reextracted(self):
        self.legacy()
        layout = SceneLayout(enabled=False, backgroundMode="none", elements=[
            SceneElement(id="authored", kind="text", text="Actual authored title", x=150, y=200, width=700, height=100),
        ])
        self.project = storage.update_scene(self.project.id, "retained", SceneUpdate(visualLayout=layout))
        self.assertEqual(storage.prepare_authored_project(self.project.id).storyboard.scenes[0].visualLayout, layout)
        resolved = resolved_project(self.project)
        self.assertTrue(resolved.storyboard.scenes[0].visualLayout.enabled)
        self.assertEqual(resolved.storyboard.scenes[0].visualLayout.elements[0].x, 150)

    def test_document_embeds_current_composition_instead_of_source_preview(self):
        self.project = storage.create_scene(self.project.id, SceneCreate(title="Current authored title", extractedText=["Current text"]))
        self.project = storage.prepare_authored_project(self.project.id)
        scene = self.project.storyboard.scenes[0]
        self.project = self.project.model_copy(update={"storyboard": self.project.storyboard.model_copy(update={
            "scenes": [scene.model_copy(update={"previewImage": "renders/obsolete-source.png"})],
        })})
        self.assertIsNone(resolved_project(self.project, published=True).storyboard.scenes[0].previewImage)
        images = render_scene_images(self.project)
        self.assertTrue(next(iter(images.values())).startswith(b"\x89PNG"))
        with patch("backend.app.scene_print.render_scene_images", return_value=images):
            document = build_storyboard_docx(self.project)
        with ZipFile(io.BytesIO(document)) as archive:
            embedded = [archive.read(name) for name in archive.namelist() if name.startswith("word/media/")]
            self.assertTrue(embedded)
            # Word normalizes the PNG container; verify pixel identity, not encoding.
            from PIL import Image
            actual = Image.open(io.BytesIO(embedded[0])).convert("RGB")
            expected = Image.open(io.BytesIO(next(iter(images.values())))).convert("RGB")
            self.assertEqual(actual.size, expected.size)
            self.assertEqual(actual.tobytes(), expected.tobytes())

    def test_preparation_requires_revision_and_uses_scene_authorization(self):
        from fastapi import HTTPException
        self.assertEqual(main.ROUTE_ACTIONS["prepare_project_authoring"], "scene.manage")
        with self.assertRaises(HTTPException) as result:
            main.prepare_project_authoring(self.project.id)
        self.assertEqual(result.exception.status_code, 428)
