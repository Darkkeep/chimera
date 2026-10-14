import unittest
import io
import tempfile
import zipfile
from copy import deepcopy
from pathlib import Path
from xml.etree import ElementTree

from pydantic import ValidationError
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Inches
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement

from backend.app.models import Scene, SceneElement, SceneLayout
from backend.app.ingest import extract_presentation
from backend.app.pptx_layout import extract_master_layout
from backend.app import main, storage
from backend.app.models import ProjectCreate, SceneUpdate
from unittest.mock import patch
from backend.app.scorm import PLAYER_JS, _scorm_12_player


class VisualLayoutTests(unittest.TestCase):
    def test_import_deduplicates_master_backdrops_and_includes_unused_masters(self):
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        master = deck.slide_masters[0]
        master.background.fill.solid()
        master.background.fill.fore_color.rgb = RGBColor(12, 34, 56)
        artwork = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
        artwork.text = "Brand & <identity>"
        master.shapes._spTree.append(deepcopy(artwork._element))
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "masters.pptx"
            deck.save(source)
            with zipfile.ZipFile(source) as archive:
                entries = {name: archive.read(name) for name in archive.namelist()}
            presentation_xml = ElementTree.fromstring(entries["ppt/presentation.xml"])
            relationships = ElementTree.fromstring(entries["ppt/_rels/presentation.xml.rels"])
            content_types = ElementTree.fromstring(entries["[Content_Types].xml"])
            for number in (2, 3):
                master_xml = ElementTree.fromstring(entries["ppt/slideMasters/slideMaster1.xml"])
                master_xml.find(qn("p:cSld")).set("name", f"Copied master {number}")
                if number == 3:
                    master_xml.find(".//" + qn("p:bg") + "//" + qn("a:srgbClr")).set("val", "AA5500")
                part_name = f"ppt/slideMasters/slideMaster{number}.xml"
                entries[part_name] = ElementTree.tostring(master_xml)
                entries[f"ppt/slideMasters/_rels/slideMaster{number}.xml.rels"] = entries["ppt/slideMasters/_rels/slideMaster1.xml.rels"]
                ElementTree.SubElement(relationships, "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship", {
                    "Id": f"rIdMaster{number}", "Type": RT.SLIDE_MASTER, "Target": f"slideMasters/slideMaster{number}.xml",
                })
                ElementTree.SubElement(presentation_xml.find(qn("p:sldMasterIdLst")), qn("p:sldMasterId"), {
                    "id": str(2147483648 + number), qn("r:id"): f"rIdMaster{number}",
                })
                ElementTree.SubElement(content_types, "{http://schemas.openxmlformats.org/package/2006/content-types}Override", {
                    "PartName": "/" + part_name, "ContentType": "application/vnd.openxmlformats-officedocument.presentationml.slideMaster+xml",
                })
            entries["ppt/presentation.xml"] = ElementTree.tostring(presentation_xml)
            entries["ppt/_rels/presentation.xml.rels"] = ElementTree.tostring(relationships)
            entries["[Content_Types].xml"] = ElementTree.tostring(content_types)
            with zipfile.ZipFile(source, "w") as archive:
                for name, content in entries.items():
                    archive.writestr(name, content)
            self.assertEqual(len(Presentation(source).slide_masters), 3)
            scenes, assets = extract_presentation(source, Path(folder))
            backgrounds = [asset for asset in assets if asset.role == "background"]
            self.assertEqual(len(backgrounds), 2)
            self.assertEqual([asset.filename for asset in backgrounds], ["Master background 1.svg", "Master background 3.svg"])
            self.assertTrue(all(asset.id not in scenes[0].imageAssets for asset in backgrounds))
            for asset in backgrounds:
                content = (Path(folder) / asset.path).read_bytes()
                self.assertIn(b"Brand &amp; &lt;identity&gt;", content)
                self.assertNotIn(b"Click to edit", content)
                ElementTree.fromstring(content)
            self.assertIn(b"#AA5500", (Path(folder) / backgrounds[1].path).read_bytes())
            _, repeated_assets = extract_presentation(source, Path(folder))
            self.assertEqual([asset.id for asset in backgrounds], [asset.id for asset in repeated_assets if asset.role == "background"])

    def test_master_backdrop_embeds_picture_bytes_and_vector_artwork(self):
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        image = io.BytesIO()
        Image.new("RGB", (20, 20), "blue").save(image, format="PNG")
        picture = slide.shapes.add_picture(io.BytesIO(image.getvalue()), Inches(1), Inches(1), Inches(2))
        master = deck.slide_masters[0]
        image_xml = deepcopy(picture._element)
        blip = image_xml.find(".//" + qn("a:blip"))
        image_part = slide.part.rels[blip.get(qn("r:embed"))].target_part
        blip.set(qn("r:embed"), master.part.relate_to(image_part, RT.IMAGE))
        master.shapes._spTree.append(image_xml)
        shape = slide.shapes.add_shape(MSO_SHAPE.PARALLELOGRAM, Inches(4), Inches(1), Inches(2), Inches(2))
        master.shapes._spTree.append(deepcopy(shape._element))
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "artwork.pptx"
            deck.save(source)
            _, assets = extract_presentation(source, Path(folder))
            background = next(asset for asset in assets if asset.role == "background")
            content = (Path(folder) / background.path).read_bytes()
        self.assertIn(b"data:image/png;base64,", content)
        self.assertIn(b"polygon", content)
        self.assertNotIn(b"/api/", content)

    def test_master_only_layout_excludes_placeholders_layout_and_slide_content(self):
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        master = slide.slide_layout.slide_master
        master.background.fill.solid()
        master.background.fill.fore_color.rgb = RGBColor(12, 34, 56)
        artwork = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
        artwork.text = "Master artwork"
        master.shapes._spTree.append(deepcopy(artwork._element))
        artwork.text = "Slide content"
        slide.slide_layout.shapes._spTree.append(deepcopy(artwork._element))
        layout = extract_master_layout(deck, master, lambda relationship: self.fail("No images expected"))
        self.assertEqual(layout.backgroundColor, "#0C2238")
        self.assertEqual([element.text for element in layout.elements], ["Master artwork"])
        self.assertTrue(all(element.source == "master" for element in layout.elements))

    def test_preset_shapes_and_zero_extent_connectors_are_extracted_as_vectors(self):
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        parallelogram = slide.shapes.add_shape(MSO_SHAPE.PARALLELOGRAM, Inches(1), Inches(1), Inches(3), Inches(1))
        parallelogram.fill.solid()
        parallelogram.fill.fore_color.rgb = RGBColor(30, 120, 180)
        parallelogram.line.color.rgb = RGBColor(10, 40, 70)
        slide.shapes.add_shape(MSO_SHAPE.CLOUD, Inches(5), Inches(5), Inches(2), Inches(1))
        custom = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(5), Inches(2), Inches(1))
        shape_properties = custom._element.find(qn("p:spPr"))
        preset = shape_properties.find(qn("a:prstGeom"))
        shape_properties.remove(preset)
        custom_geometry = OxmlElement("a:custGeom")
        for name in ("avLst", "gdLst", "ahLst", "cxnLst"):
            custom_geometry.append(OxmlElement(f"a:{name}"))
        rect = OxmlElement("a:rect")
        for name, value in (("l", "0"), ("t", "0"), ("r", "100"), ("b", "100")):
            rect.set(name, value)
        custom_geometry.append(rect)
        path_list = OxmlElement("a:pathLst")
        path_node = OxmlElement("a:path")
        path_node.set("w", "100")
        path_node.set("h", "100")
        move = OxmlElement("a:moveTo")
        point = OxmlElement("a:pt")
        point.set("x", "50")
        point.set("y", "0")
        move.append(point)
        path_node.append(move)
        for x, y in ((100, 100), (0, 100), (50, 0)):
            line_to = OxmlElement("a:lnTo")
            point = OxmlElement("a:pt")
            point.set("x", str(x))
            point.set("y", str(y))
            line_to.append(point)
            path_node.append(line_to)
        path_node.append(OxmlElement("a:close"))
        path_list.append(path_node)
        custom_geometry.append(path_list)
        shape_properties.append(custom_geometry)
        slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(1), Inches(3), Inches(5), Inches(3))
        slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(6), Inches(1), Inches(6), Inches(4))
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "shapes.pptx"
            deck.save(source)
            scenes, assets = extract_presentation(source, Path(folder))
        layout = scenes[0].visualLayout
        vector_shapes = [element for element in layout.elements if element.kind == "shape"]
        self.assertEqual(len([asset for asset in assets if asset.role == "content"]), 0)
        self.assertEqual(len([asset for asset in assets if asset.role == "background"]), 1)
        self.assertEqual(len(vector_shapes), 5)
        self.assertIn("polygon", vector_shapes[0].shapeSvg)
        self.assertIn("#1E78B4", vector_shapes[0].shapeSvg.upper())
        self.assertIn('preserveAspectRatio="none"', vector_shapes[0].shapeSvg)
        self.assertIn("<rect", vector_shapes[1].shapeSvg)
        self.assertTrue(any("editable approximation" in warning for warning in layout.warnings))
        self.assertIn("<path", vector_shapes[2].shapeSvg)
        self.assertTrue(all("<path" in element.shapeSvg for element in vector_shapes[3:]))
        self.assertTrue(all('stroke="none"' not in element.shapeSvg for element in vector_shapes[3:]))
        self.assertTrue(all(element.width > 0 and element.height > 0 for element in vector_shapes))
        self.assertFalse(any("Parallelogram" in warning or "Connector" in warning for warning in layout.warnings))

    def test_inherited_picture_background_and_slide_override(self):
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        image = io.BytesIO()
        Image.new("RGB", (20, 20), "blue").save(image, format="PNG")
        picture = slide.shapes.add_picture(io.BytesIO(image.getvalue()), 0, 0, Inches(1))
        image_part = slide.part.rels[picture._element.find(".//" + qn("a:blip")).get(qn("r:embed"))].target_part
        picture._element.getparent().remove(picture._element)
        master = slide.slide_layout.slide_master
        background = OxmlElement("p:bg")
        properties = OxmlElement("p:bgPr")
        fill = OxmlElement("a:blipFill")
        blip = OxmlElement("a:blip")
        blip.set(qn("r:embed"), master.part.relate_to(image_part, RT.IMAGE))
        fill.append(blip)
        properties.append(fill)
        background.append(properties)
        master._element.find(qn("p:cSld")).insert(0, background)
        override = deck.slides.add_slide(deck.slide_layouts[6])
        override.background.fill.solid()
        override.background.fill.fore_color.rgb = RGBColor(1, 2, 3)
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "deck.pptx"
            deck.save(source)
            scenes, assets = extract_presentation(source, Path(folder))
        self.assertIsNotNone(scenes[0].backgroundAsset)
        self.assertEqual(scenes[0].backgroundAsset, assets[0].id)
        self.assertIsNone(scenes[1].backgroundAsset)
        self.assertEqual(scenes[1].visualLayout.backgroundColor, "#010203")

    def test_import_preserves_geometry_and_master_images(self):
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.slide_layout.slide_master.background.fill.solid()
        slide.slide_layout.slide_master.background.fill.fore_color.rgb = RGBColor(12, 34, 56)
        image = io.BytesIO()
        Image.new("RGB", (30, 20), "red").save(image, format="PNG")
        master_picture = slide.shapes.add_picture(io.BytesIO(image.getvalue()), Inches(0), Inches(0), Inches(1))
        picture_xml = deepcopy(master_picture._element)
        blip = picture_xml.find(".//" + qn("a:blip"))
        image_part = slide.part.rels[blip.get(qn("r:embed"))].target_part
        blip.set(qn("r:embed"), slide.slide_layout.slide_master.part.relate_to(image_part, RT.IMAGE))
        slide.slide_layout.slide_master.shapes._spTree.append(picture_xml)
        master_picture._element.getparent().remove(master_picture._element)
        slide.shapes.add_picture(io.BytesIO(image.getvalue()), Inches(1), Inches(2), Inches(2))
        slide.shapes.add_picture(io.BytesIO(image.getvalue()), Inches(4), Inches(2), Inches(2))
        textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
        textbox.text = "Editable text"
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "deck.pptx"
            deck.save(source)
            scenes, assets = extract_presentation(source, Path(folder))
        layout = scenes[0].visualLayout
        self.assertEqual(layout.backgroundColor, "#0C2238")
        self.assertEqual(len([asset for asset in assets if asset.role == "content"]), 1)
        self.assertEqual(len([element for element in layout.elements if element.kind == "image"]), 3)
        self.assertTrue(layout.elements[0].locked)
        text = next(element for element in layout.elements if element.kind == "text")
        self.assertAlmostEqual(text.x, 128)
        self.assertEqual(text.text, "Editable text")
        self.assertFalse(layout.enabled)

    def test_identical_image_assets_are_shared_across_extracted_slides(self):
        deck = Presentation()
        image = io.BytesIO()
        Image.new("RGB", (30, 20), "red").save(image, format="PNG")
        image_bytes = image.getvalue()
        for _ in range(2):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_picture(io.BytesIO(image_bytes), Inches(1), Inches(1), Inches(2))
            slide.shapes.add_picture(io.BytesIO(image_bytes), Inches(4), Inches(1), Inches(2))
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "deck.pptx"
            deck.save(source)
            scenes, assets = extract_presentation(source, Path(folder))
        self.assertEqual(len([asset for asset in assets if asset.role == "content"]), 1)
        self.assertEqual(scenes[0].imageAssets, [assets[0].id])
        self.assertEqual(scenes[1].imageAssets, [assets[0].id])
        self.assertEqual([element.assetId for element in scenes[0].visualLayout.elements if element.kind == "image"], [assets[0].id, assets[0].id])
        self.assertEqual([element.assetId for element in scenes[1].visualLayout.elements if element.kind == "image"], [assets[0].id, assets[0].id])

    def test_legacy_scene_remains_unchanged(self):
        scene = Scene(id="old", title="Old scene", previewImage="renders/slide.png")
        self.assertIsNone(scene.visualLayout)
        self.assertEqual(scene.previewImage, "renders/slide.png")

    def test_legacy_extraction_preserves_edits_and_is_idempotent(self):
        with tempfile.TemporaryDirectory() as folder, patch.object(storage, "ROOT", Path(folder)), patch.object(storage, "PROJECTS", Path(folder) / "projects"):
            project = storage.create_project(ProjectCreate(title="Migration"))
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "Original"
            source = storage.project_path(project.id) / "original" / "source-deck.pptx"
            deck.save(source)
            storage.update_from_import(project.id, [Scene(id="retained", sourceSlide=1, title="Edited title", extractedText=["Edited body"], notes="Keep notes")], [])
            current = storage.read_project(project.id)
            updated = main.extract_scene_layout(project.id, "retained", str(current.storyboard.revision))
            scene = updated.storyboard.scenes[0]
            self.assertEqual(scene.id, "retained")
            self.assertEqual(scene.notes, "Keep notes")
            self.assertEqual(scene.extractedText, ["Edited body"])
            self.assertIn("Edited body", scene.visualLayout.elements[0].text)
            self.assertFalse(scene.visualLayout.enabled)
            again = main.extract_scene_layout(project.id, "retained", str(updated.storyboard.revision))
            self.assertEqual(again.storyboard.revision, updated.storyboard.revision)
            invalid = SceneLayout(elements=[SceneElement(id="bad", kind="image", assetId="missing", x=0, y=0, width=20, height=20)])
            with self.assertRaises(ValueError):
                storage.update_scene(project.id, "retained", SceneUpdate(visualLayout=invalid))

    def test_repeated_assets_have_independent_elements(self):
        elements = [SceneElement(id=identifier, kind="image", assetId="same-image", x=10, y=20, width=100, height=80) for identifier in ["first", "second"]]
        layout = SceneLayout(elements=elements)
        self.assertEqual(len(SceneLayout.model_validate_json(layout.model_dump_json()).elements), 2)
        with self.assertRaises(ValidationError):
            SceneLayout(elements=[elements[0], elements[0]])

    def test_both_export_runtimes_render_visual_layout(self):
        for runtime in [PLAYER_JS, _scorm_12_player()]:
            self.assertIn("if (visual) renderVisualScene(scene, body)", runtime)
            self.assertIn("if (!visual && scene.previewImage)", runtime)
            self.assertIn("node.textContent = element.text", runtime)
            self.assertIn("element.crop", runtime)
            self.assertIn('element.kind === "shape"', runtime)
            self.assertIn("element.shapeSvg", runtime)
            self.assertIn('document.querySelector("#contents")', runtime)

    def test_modal_text_changes_update_enabled_layout(self):
        with tempfile.TemporaryDirectory() as folder, patch.object(storage, "ROOT", Path(folder)), patch.object(storage, "PROJECTS", Path(folder) / "projects"):
            project = storage.create_project(ProjectCreate(title="Text update"))
            layout = SceneLayout(enabled=True, elements=[SceneElement(id="text", kind="text", text="Before", x=80, y=90, width=300, height=200)])
            storage.update_from_import(project.id, [Scene(id="visual", title="Title", visualLayout=layout)], [])
            updated = storage.update_scene(project.id, "visual", SceneUpdate(extractedText=["After"]))
            self.assertEqual(updated.storyboard.scenes[0].visualLayout.elements[0].text, "After")
            self.assertEqual(updated.storyboard.scenes[0].visualLayout.elements[0].x, 80)

    def test_modal_edits_are_preserved_before_visual_activation(self):
        with tempfile.TemporaryDirectory() as folder, patch.object(storage, "ROOT", Path(folder)), patch.object(storage, "PROJECTS", Path(folder) / "projects"):
            project = storage.create_project(ProjectCreate(title="Inactive text"))
            layout = SceneLayout(elements=[SceneElement(id="title", source="slide", kind="text", text="Before title", x=10, y=10, width=100, height=50), SceneElement(id="body", source="slide", kind="text", text="Before body", x=10, y=80, width=100, height=50)])
            storage.update_from_import(project.id, [Scene(id="inactive", title="Before title", extractedText=["Before body"], visualLayout=layout)], [])
            updated = storage.update_scene(project.id, "inactive", SceneUpdate(title="After title", extractedText=["After body"]))
            self.assertEqual([element.text for element in updated.storyboard.scenes[0].visualLayout.elements], ["After title", "After body"])
            self.assertFalse(updated.storyboard.scenes[0].visualLayout.enabled)

    def test_detached_image_is_removed_from_visual_elements(self):
        from backend.app.models import ExtractedAsset

        with tempfile.TemporaryDirectory() as folder, patch.object(storage, "ROOT", Path(folder)), patch.object(storage, "PROJECTS", Path(folder) / "projects"):
            project = storage.create_project(ProjectCreate(title="Detach image"))
            asset = ExtractedAsset(id="image", kind="image", filename="image.png", contentType="image/png", sizeBytes=0, sha256="test", path="assets/images/image.png")
            layout = SceneLayout(enabled=True, elements=[SceneElement(id="placed", kind="image", assetId="image", x=10, y=10, width=100, height=100)])
            storage.update_from_import(project.id, [Scene(id="visual", title="Scene", imageAssets=["image"], visualLayout=layout)], [asset])
            updated = storage.update_scene(project.id, "visual", SceneUpdate(imageAssets=[]))
            self.assertEqual(updated.storyboard.scenes[0].visualLayout.elements, [])
            self.assertEqual(len(updated.storyboard.assets), 1)

    def test_invalid_geometry_and_asset_are_rejected(self):
        for changes in [{"width": 0}, {"x": float("nan")}, {"kind": "image"}, {"crop": [0.6, 0, 0.6, 0]}]:
            with self.assertRaises(ValidationError):
                SceneElement.model_validate({"id": "text", "kind": "text", "x": 0, "y": 0, "width": 100, "height": 100, **changes})