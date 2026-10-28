import base64
import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from fastapi.testclient import TestClient
from pydantic import ValidationError

from backend.app import ai_budget, backup, content_generation, main, media, ocr, scorm, storage, word_export
from backend.app.models import AssetMetadataUpdate, MediaPlacement, MediaPlacementsUpdate, ProjectCreate, ReferencePassage, ReviewCommentCreate, SceneCreate


class MediaCompletionTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        root = Path(temporary.name)
        for module in (storage, main):
            self.enterContext(patch.object(module, "ROOT", root))
        self.enterContext(patch.object(storage, "PROJECTS", root / "projects"))
        self.client = TestClient(main.app)
        self.project = storage.create_project(ProjectCreate(title="Media qualification"))
        self.project = storage.create_scene(self.project.id, SceneCreate(title="One"))
        self.project = storage.create_scene(self.project.id, SceneCreate(title="Two"))

    def video(self):
        self.project = storage.add_project_asset(self.project.id, "lesson.mp4", b"\x00\x00\x00\x18ftypisom\x00\x00\x00\x00", "video")
        return self.project.storyboard.assets[-1]

    def test_shared_copy_is_pinned_and_never_placed_and_duplicate_copy_preserves_edits(self):
        asset = self.video()
        source = media.publish_shared(self.project.id, asset.id)
        target = storage.create_project(ProjectCreate(title="Shared consumer"))
        copied = media.copy_shared(target.id, source["id"], target.storyboard.revision)
        self.assertEqual(copied.storyboard.revision, target.storyboard.revision + 1)
        self.assertEqual(copied.storyboard.assets[0].sharedSourceId, source["id"])
        self.assertEqual(copied.storyboard.courseMedia, [])
        copied = media.update_metadata(target.id, asset.id, AssetMetadataUpdate(transcript="Local edits"), copied.storyboard.revision)
        unchanged = media.copy_shared(target.id, source["id"], copied.storyboard.revision)
        self.assertEqual(unchanged.storyboard.assets[0].transcript, "Local edits")
        path, _ = media.shared_record(source["id"])
        (path / "record.json").rename(path / "withdrawn.json")
        self.assertEqual(media.list_shared(), [])
        self.assertEqual(storage.read_project(target.id).storyboard.assets[0].sha256, asset.sha256)
        with self.assertRaises(FileNotFoundError):
            media.copy_shared(target.id, source["id"], copied.storyboard.revision)

    def test_corrupt_shared_media_cannot_modify_target(self):
        asset = self.video()
        record = media.publish_shared(self.project.id, asset.id)
        path, _ = media.shared_record(record["id"])
        (path / "content").write_bytes(b"corrupt")
        before = storage.read_project(self.project.id)
        with self.assertRaisesRegex(ValueError, "checksum"):
            media.copy_shared(self.project.id, record["id"], before.storyboard.revision)
        self.assertEqual(storage.read_project(self.project.id), before)

    def test_placement_precedence_exclusion_and_stale_write(self):
        asset = self.video()
        self.project = media.write_placements(self.project.id, MediaPlacementsUpdate(placements=[MediaPlacement(assetId=asset.id, position="before", order=3)]), self.project.storyboard.revision)
        first, second = self.project.storyboard.scenes
        self.project = media.write_placements(self.project.id, MediaPlacementsUpdate(excludedAssetIds=[asset.id]), self.project.storyboard.revision, first.id)
        self.assertEqual(media.effective_placements(self.project, self.project.storyboard.scenes[0]), [])
        self.assertEqual(media.effective_placements(self.project, second)[0].position, "before")
        old = self.project.storyboard.revision
        self.project = media.write_placements(self.project.id, MediaPlacementsUpdate(placements=[MediaPlacement(assetId=asset.id)], excludedAssetIds=[asset.id]), old, first.id)
        self.assertEqual(media.effective_placements(self.project, self.project.storyboard.scenes[0])[0].position, "after")
        with self.assertRaises(storage.StoryboardRevisionConflict):
            media.write_placements(self.project.id, MediaPlacementsUpdate(), old)
        with self.assertRaisesRegex(ValueError, "placements"):
            media.delete_asset(self.project.id, asset.id, self.project.storyboard.revision)

    def test_invalid_caption_and_placement_leave_revision_unchanged(self):
        asset = self.video()
        version = self.project.storyboard.revision
        for captions in ("text", "WEBVTT\n\n00:00:04.000 --> 00:00:01.000\nInvalid", "WEBVTT\n\n00:61.000 --> 01:00.000\nInvalid"):
            with self.subTest(captions=captions), self.assertRaises(ValueError):
                media.update_metadata(self.project.id, asset.id, AssetMetadataUpdate(captionsVtt=captions), version)
        with self.assertRaises(ValueError):
            media.write_placements(self.project.id, MediaPlacementsUpdate(placements=[MediaPlacement(assetId="unknown")]), version)
        self.assertEqual(storage.read_project(self.project.id).storyboard.revision, version)

    def test_required_viewing_is_video_only_non_looping_and_threshold_bounded(self):
        asset = self.video()
        self.project = storage.add_project_asset(self.project.id, "narration.mp3", b"audio", "audio")
        audio = self.project.storyboard.assets[-1]
        revision = self.project.storyboard.revision
        with self.assertRaisesRegex(ValueError, "only be configured for video"):
            media.write_placements(self.project.id, MediaPlacementsUpdate(placements=[MediaPlacement(assetId=audio.id, requiredViewing=True)]), revision)
        with self.assertRaisesRegex(ValueError, "cannot loop"):
            media.write_placements(self.project.id, MediaPlacementsUpdate(placements=[MediaPlacement(assetId=asset.id, requiredViewing=True, loop=True)]), revision)
        with self.assertRaises(ValidationError):
            MediaPlacement(assetId=asset.id, viewingThresholdPercent=101)
        self.assertEqual(storage.read_project(self.project.id).storyboard.revision, revision)

    def test_required_video_settings_are_packaged_and_included_in_resume_estimate(self):
        asset = self.video()
        from backend.app.models import AssetAttributionUpdate
        self.project = media.update_metadata(self.project.id, asset.id, AssetMetadataUpdate(
            transcript="Inspect the valve.",
            captionsVtt="WEBVTT\n\n00:00:00.000 --> 00:00:02.000\nInspect the valve.",
        ), self.project.storyboard.revision)
        self.project = storage.update_project_asset_attribution(self.project.id, asset.id, AssetAttributionUpdate(
            attributionStatus="complete", attributionText="Synthetic test video",
        ))
        self.project = media.write_placements(self.project.id, MediaPlacementsUpdate(placements=[
            MediaPlacement(assetId=asset.id, requiredViewing=True, viewingThresholdPercent=85),
        ]), self.project.storyboard.revision)
        for builder, budget in ((scorm.build_scorm_12, 4096), (scorm.build_scorm_2004, 64000)):
            package, report = builder(self.project)
            with ZipFile(io.BytesIO(package)) as archive:
                story = json.loads(archive.read("storyboard.json"))
                self.assertTrue(all(scene["mediaPlacements"][0]["requiredViewing"] for scene in story["scenes"]))
                self.assertTrue(all(scene["mediaPlacements"][0]["viewingThresholdPercent"] == 85 for scene in story["scenes"]))
                self.assertIn("videoProgress", archive.read("player.js").decode())
            self.assertEqual(report["resumeData"]["budgetBytes"], budget)
            self.assertEqual(report["resumeData"]["requiredVideoCount"], 1)
            self.assertIn("Estimate includes one range", report["resumeData"]["videoRangeFragmentationNote"])

    def test_packages_publish_only_used_media_with_captions_and_transcripts(self):
        asset = self.video()
        unused = storage.add_project_asset(self.project.id, "unused.mp3", b"unused", "audio")
        self.project = unused
        self.project = media.update_metadata(self.project.id, asset.id, AssetMetadataUpdate(transcript="Inspect the valve.", captionsVtt="WEBVTT\n\n00:00:00.000 --> 00:00:02.000\nInspect the valve."), self.project.storyboard.revision)
        from backend.app.models import AssetAttributionUpdate
        self.project = storage.update_project_asset_attribution(self.project.id, asset.id, AssetAttributionUpdate(attributionStatus="complete", attributionText="Synthetic test video"))
        self.project = media.write_placements(self.project.id, MediaPlacementsUpdate(placements=[MediaPlacement(assetId=asset.id)]), self.project.storyboard.revision)
        for builder, budget in ((scorm.build_scorm_12, 4096), (scorm.build_scorm_2004, 64000)):
            package, report = builder(self.project)
            with ZipFile(io.BytesIO(package)) as archive:
                story = json.loads(archive.read("storyboard.json"))
                self.assertEqual(len(story["assets"]), 1)
                self.assertEqual(story["assets"][0]["transcript"], "Inspect the valve.")
                self.assertTrue(story["assets"][0]["captionsVtt"].startswith("WEBVTT"))
                self.assertTrue(all(s["videoAssets"] == [asset.id] for s in story["scenes"]))
                self.assertFalse(any("unused" in name for name in archive.namelist()))
                self.assertIn('media.addTextTrack("captions"', archive.read("player.js").decode())
                self.assertEqual(report["resumeData"]["budgetBytes"], budget)

    def test_video_production_requires_captions_transcript_and_credit(self):
        asset = self.video()
        self.project = media.write_placements(self.project.id, MediaPlacementsUpdate(placements=[MediaPlacement(assetId=asset.id)]), self.project.storyboard.revision)
        with self.assertRaisesRegex(ValueError, "captions"):
            scorm.build_scorm_12(self.project)

    def test_word_discussion_is_explicit_and_stays_out_of_published_content(self):
        comment = storage.create_review_comment(self.project.id, ReviewCommentCreate(sceneId=self.project.storyboard.scenes[0].id, authorLabel="Reviewer", body="Private production discussion"))
        def content(data):
            doc = Document(io.BytesIO(data))
            return "\n".join(p.text for p in doc.paragraphs)
        self.assertNotIn(comment.body, content(word_export.build_storyboard_docx(self.project)))
        self.assertIn(comment.body, content(word_export.build_storyboard_docx(self.project, include_review_comments=True)))
        package, _ = scorm.build_scorm_12(self.project)
        with ZipFile(io.BytesIO(package)) as archive:
            self.assertNotIn(comment.body, archive.read("storyboard.json").decode())

    def test_new_routes_require_revision_and_validate_video_container(self):
        root = f"/api/projects/{self.project.id}"
        self.assertEqual(self.client.put(root + "/course-media", json={"placements": []}).status_code, 428)
        rejected = self.client.post(root + "/assets", files={"file": ("bad.mp4", b"not a video", "video/mp4")})
        self.assertEqual(rejected.status_code, 422)
        shared = self.client.get(root + "/shared-media")
        self.assertEqual(shared.status_code, 200)
        self.assertEqual(shared.json(), [])

    def test_ocr_creates_version_and_requires_inspection_before_grounding(self):
        resource = storage.save_reference_resource(self.project.id, "scan.pdf", "Scanned guide", "study-guide", True, [], b"synthetic pdf", "application/pdf", [], "failed", "No text")
        passages = [ReferencePassage(locator="page:1", text="Inspect before operation.")]
        with patch.object(ocr, "extract_pdf", return_value=(passages, "partial", "Local OCR (eng); inspect text")):
            updated = ocr.run_reference_ocr(self.project.id, resource.id, resource.currentVersion, "eng")
        self.assertEqual(len(updated.versions), 2)
        self.assertEqual(updated.versions[0].extractionStatus, "failed")
        with self.assertRaisesRegex(ValueError, "readable"):
            content_generation._current_reference_context(self.project.id, [resource.id])
        accepted = ocr.accept_reference_ocr(self.project.id, resource.id, updated.currentVersion)
        self.assertEqual(accepted.versions[-1].extractionStatus, "complete")
        self.assertEqual(content_generation._current_reference_context(self.project.id, [resource.id])[0]["passages"][0]["locator"], "page:1")
        with self.assertRaisesRegex(ValueError, "changed"):
            ocr.accept_reference_ocr(self.project.id, resource.id, resource.currentVersion)

    def test_ai_budget_and_reference_budget_fail_before_provider_submission(self):
        with self.assertRaisesRegex(ValueError, "No provider request"):
            ai_budget.checked_json({"source": "漢" * 100000})
        resource = storage.save_reference_resource(self.project.id, "long.txt", "Long", "study-guide", True, [], b"text", "text/plain", [ReferencePassage(locator="line:1", text="x" * 8001)], "complete", None)
        with self.assertRaisesRegex(ValueError, "silently omitted"):
            content_generation._current_reference_context(self.project.id, [resource.id])
        self.assertEqual(ai_budget.provider_usage({"usage": {"prompt_tokens": 42, "completion_tokens": 5, "secret": "not retained"}}), {"prompt_tokens": 42, "completion_tokens": 5})

    def test_media_archive_round_trip_preserves_placements_metadata_and_independent_copy(self):
        asset = self.video()
        self.project = media.update_metadata(self.project.id, asset.id, AssetMetadataUpdate(transcript="Archive transcript", captionsVtt="WEBVTT\n\n00:00:00.000 --> 00:00:01.000\nInspect."), self.project.storyboard.revision)
        self.project = media.write_placements(self.project.id, MediaPlacementsUpdate(placements=[MediaPlacement(assetId=asset.id)]), self.project.storyboard.revision)
        self.project = media.write_placements(self.project.id, MediaPlacementsUpdate(excludedAssetIds=[asset.id]), self.project.storyboard.revision, self.project.storyboard.scenes[0].id)
        archive = backup.build_project_archive(self.project.id)
        before = self.project.storyboard.model_dump()
        media.write_placements(self.project.id, MediaPlacementsUpdate(), self.project.storyboard.revision)
        restored = backup.restore_project_archive(self.project.id, archive)
        self.assertEqual(restored.storyboard.model_dump(), before)
        self.assertEqual(media.effective_placements(restored, restored.storyboard.scenes[0]), [])
        self.assertEqual(media.effective_placements(restored, restored.storyboard.scenes[1])[0].assetId, asset.id)

    def test_word_citations_remain_pinned_when_reference_is_replaced(self):
        from backend.app.models import EnhancementProposalCreate
        resource = storage.save_reference_resource(self.project.id, "guide.txt", "Operating guide", "study-guide", True, [], b"Inspect first.", "text/plain", [ReferencePassage(locator="line:1", text="Inspect first.")], "complete", None)
        original_hash = resource.versions[0].sha256
        proposal = storage.create_enhancement_proposal(self.project.id, EnhancementProposalCreate(targetSceneId=self.project.storyboard.scenes[0].id, proposedValue="Inspect first.", rationale="Grounded clarification", referenceCitations=[{"resourceId": resource.id, "version": "1", "locator": "line:1", "excerpt": "Inspect first."}]))
        storage.decide_enhancement_proposal(self.project.id, proposal.id, "accept", self.project.storyboard.revision)
        storage.save_reference_resource(self.project.id, "guide.txt", "Operating guide", "study-guide", True, [], b"Different new guidance.", "text/plain", [ReferencePassage(locator="line:1", text="Different new guidance.")], "complete", None, resource.id)
        doc = Document(io.BytesIO(word_export.build_storyboard_docx(storage.read_project(self.project.id))))
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("version 1 | line:1", text)
        self.assertIn(original_hash, text)
        self.assertNotIn("Different new guidance.", text)

    def test_api_layout_warnings_are_advisory_and_available_without_browser_fitting(self):
        from backend.app.models import SceneUpdate
        from backend.app.validation import validate_project
        scene = self.project.storyboard.scenes[0]
        self.project = storage.update_scene(self.project.id, scene.id, SceneUpdate(visualLayout={"enabled": True, "backgroundColor": "#ffffff", "elements": [
            {"id": "a", "kind": "text", "x": 10, "y": 10, "width": 200, "height": 100, "color": "#eeeeee", "text": "Low contrast"},
            {"id": "b", "kind": "text", "x": 20, "y": 20, "width": 200, "height": 100, "text": "Overlapping text"},
        ]}))
        report = validate_project(self.project)
        warnings = [i for i in report.issues if i.code == "visual-layout-review"]
        self.assertEqual(len(warnings), 2)
        self.assertTrue(all(i.severity == "warning" and i.sceneId == scene.id for i in warnings))
        self.assertTrue(report.passed)
