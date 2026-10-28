import base64
import csv
import json
import re
import unittest
import warnings
from io import BytesIO, StringIO
from unittest.mock import patch
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from openpyxl import Workbook, load_workbook

from backend.app import question_import as importer
from backend.app.models import AssessmentSettings, Scene
from backend.app.qti import assessment_item_xml, assessment_test_package, qti_tag


def scenes():
    rows = importer.parse_questions("questions.csv", importer.question_template("csv"))
    return [
        Scene(id=f"scene_{index}", title=row["title"], sceneType="question",
              assessmentItem=row["item"],
              assessmentSettings=AssessmentSettings(maxAttempts=3, passingScore=75, scoreWeight=2))
        for index, row in enumerate(rows)
    ]


def xml_bytes(root):
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def without_extension(content):
    return b"\n".join(line for line in content.split(b"\n") if not line.startswith(b"<?chimera "))


def with_extension(content, extension):
    content = without_extension(content)
    declaration, body = content.split(b"\n", 1)
    encoded = base64.b64encode(json.dumps(extension, ensure_ascii=False).encode("utf-8"))
    return declaration + b"\n<?chimera " + encoded + b"?>\n" + body


def zip_bytes(entries):
    stream = BytesIO()
    with ZipFile(stream, "w", ZIP_DEFLATED) as archive:
        for name, value in entries:
            archive.writestr(name, value)
    return stream.getvalue()


def csv_bytes(headers, rows):
    stream = StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow(headers)
    writer.writerows(rows)
    return stream.getvalue().encode("utf-8-sig")


class QuestionImportParserTests(unittest.TestCase):
    def assert_valid(self, rows, count=1):
        self.assertEqual(len(rows), count, rows)
        self.assertEqual([row["index"] for row in rows], list(range(count)))
        for row in rows:
            self.assertEqual(row["errors"], [], row)
            self.assertIsNotNone(row["item"], row)
        return rows

    def assert_invalid(self, content, filename="item.xml", contains=None):
        rows = importer.parse_questions(filename, content)
        self.assertTrue(any(row["errors"] for row in rows), rows)
        self.assertTrue(all(row["item"] is None for row in rows if row["errors"]), rows)
        if contains:
            self.assertIn(contains.lower(), " ".join(error for row in rows for error in row["errors"]).lower())
        return rows

    def test_templates_cover_all_seven_types(self):
        for format in ("csv", "xlsx"):
            with self.subTest(format=format):
                rows = self.assert_valid(
                    importer.parse_questions(f"template.{format}", importer.question_template(format)), 7,
                )
                self.assertEqual([row["item"]["type"] for row in rows], [
                    "single-choice", "multiple-response", "true-false", "ordering",
                    "drag-drop", "matching", "short-answer",
                ])
                self.assertTrue(all(any("settings" in value for value in row["warnings"]) for row in rows))
        with self.assertRaises(ValueError):
            importer.question_template("ods")

    def test_spreadsheet_roundtrip_all_data_metadata_and_literal_formula_strings(self):
        questions = scenes()
        for index, scene in enumerate(questions):
            item = scene.assessmentItem.model_dump(mode="json")
            item["metadata"] = {
                **scene.assessmentItem.metadata.model_dump(), "title": f"=Title {index}",
                "description": "'description", "tags": ["one", "two"],
                "author": "+Author", "source": "@Source", "language": "fr-CA",
                "difficulty": "hard", "learningObjectiveRefs": ["objective-1"],
            }
            scene.assessmentItem = importer.ITEM_ADAPTER.validate_python(item)
            scene.assessmentItem.prompt = "\t=SUM(1,2)\nA second line, with \"quotes\" and café"
            scene.assessmentItem.correctFeedback = ""
            scene.assessmentItem.incorrectFeedback = "-Not yet"
        for format in ("csv", "xlsx"):
            with self.subTest(format=format):
                data = importer.question_spreadsheet(questions, format)
                rows = self.assert_valid(importer.parse_questions(f"roundtrip.{format}", data), 7)
                self.assertEqual([row["item"] for row in rows],
                                 [scene.assessmentItem.model_dump(mode="json") for scene in questions])
                if format == "csv":
                    self.assertIn(b"'=Title", data)
                else:
                    workbook = load_workbook(BytesIO(data), data_only=False)
                    self.assertTrue(all(cell.data_type != "f" for row in workbook.active for cell in row))
                    workbook.close()

    def test_empty_metadata_title_and_apostrophe_escape_are_lossless(self):
        scene = scenes()[0]
        scene.assessmentItem.metadata.title = ""
        for text in ("'", "''", "'plain", "=formula", " +formula", "plain", "\n@formula"):
            scene.assessmentItem.prompt = text
            for format in ("csv", "xlsx"):
                with self.subTest(text=text, format=format):
                    rows = self.assert_valid(importer.parse_questions(
                        f"escape.{format}", importer.question_spreadsheet([scene], format),
                    ))
                    self.assertEqual(rows[0]["item"], scene.assessmentItem.model_dump(mode="json"))

    def test_csv_bom_multiline_minimal_headers_and_default_feedback(self):
        headers = ["identifier", "type", "title", "prompt", "answers"]
        rows = self.assert_valid(importer.parse_questions("minimal.csv", csv_bytes(headers, [
            ["question", "short-answer", "Title", "A prompt\nwith a newline", '["answer"]'],
        ])))
        self.assertEqual(rows[0]["item"]["prompt"], "A prompt\nwith a newline")
        self.assertEqual(rows[0]["item"]["correctFeedback"], "Correct.")
        self.assertEqual(rows[0]["item"]["metadata"]["language"], "en")
        self.assertEqual(rows[0]["source"], "minimal.csv:row 2")

    def test_malformed_csv_and_unknown_nonempty_columns(self):
        headers = ["identifier", "type", "title", "prompt", "answers"]
        row = ["question", "short-answer", "Title", "Prompt", '["answer"]']
        cases = [
            (csv_bytes(headers + ["title"], [row + ["duplicate"]]), "duplicate"),
            (csv_bytes(headers[:-1], [row[:-1]]), "missing"),
            (csv_bytes(headers, [row[:-1]]), "column count"),
            (csv_bytes(headers + ["custom"], [row + ["not preserved"]]), "unsupported"),
            (csv_bytes(headers + ["", ""], [row + ["not preserved", ""]]), "blank header"),
            (csv_bytes(headers, [row[:-1] + ["not JSON"]]), "JSON"),
            (b"\xff\xfe", None),
            (b'"unterminated', None),
        ]
        for data, message in cases:
            with self.subTest(message=message):
                self.assert_invalid(data, "bad.csv", message)
        self.assert_valid(importer.parse_questions("empty-column.csv", csv_bytes(
            headers + ["custom"], [row + [""]],
        )))

    def test_csv_validation_errors_are_per_row_and_indices_sequential(self):
        data = importer.question_template("csv").decode("utf-8-sig")
        rows = list(csv.reader(StringIO(data)))
        rows[2][4] = "not json"
        rows[3][0] = rows[1][0]
        parsed = importer.parse_questions("bad.csv", csv_bytes(rows[0], rows[1:]))
        self.assertEqual([row["index"] for row in parsed], list(range(7)))
        self.assertTrue(parsed[1]["errors"])
        self.assertTrue(parsed[2]["errors"])
        self.assertIsNotNone(parsed[0]["item"])
        self.assertIsNotNone(parsed[6]["item"])

    def test_spreadsheet_rejects_ignored_fields_and_invalid_metadata(self):
        data = list(csv.reader(StringIO(importer.question_template("csv").decode("utf-8-sig"))))
        for column, value, row_index in (
            ("tags", '["duplicate","duplicate"]', 1),
            ("tags", '[2]', 1),
            ("language", "not a language", 1),
            ("difficulty", "impossible", 1),
            ("secondaryChoices", "[]", 1),
            ("choices", "[]", 3),
            ("caseSensitive", "perhaps", 7),
            ("choices", '[{"identifier":"a","text":"A","extra":1}]', 1),
        ):
            with self.subTest(column=column, value=value):
                copy = [row[:] for row in data]
                copy[row_index][copy[0].index(column)] = value
                self.assert_invalid(csv_bytes(copy[0], copy[1:]), "bad.csv")

    def test_xlsx_sheet_selection_formulas_and_dimensions(self):
        def book_bytes(workbook):
            stream = BytesIO()
            workbook.save(stream)
            workbook.close()
            return stream.getvalue()

        workbook = load_workbook(BytesIO(importer.question_template("xlsx")))
        workbook.create_sheet("Guidance")
        self.assert_valid(importer.parse_questions("sheet.xlsx", book_bytes(workbook)), 7)
        workbook = Workbook()
        workbook.create_sheet("Other")
        self.assert_invalid(book_bytes(workbook), "sheets.xlsx", "Questions")
        workbook = load_workbook(BytesIO(importer.question_template("xlsx")))
        workbook.active["D2"] = "=1+1"
        self.assert_invalid(book_bytes(workbook), "formula.xlsx", "formula")
        workbook = load_workbook(BytesIO(importer.question_template("xlsx")))
        workbook.active.cell(502, 1, "too far")
        self.assert_invalid(book_bytes(workbook), "long.xlsx", "row")
        workbook = load_workbook(BytesIO(importer.question_template("xlsx")))
        workbook.active.cell(2, 101, "too far")
        self.assert_invalid(book_bytes(workbook), "wide.xlsx", "column")

    def test_xlsx_cached_formulas_and_forged_dimensions_are_rejected(self):
        data = importer.question_template("xlsx")
        with ZipFile(BytesIO(data)) as archive:
            original = [(name, archive.read(name)) for name in archive.namelist()]
        for cell, message in (
            (b'<c r="D2"><f>1+1</f><v>2</v></c>', "formula"),
            (b'<c r="ZZ2" t="inlineStr"><is><t>hidden</t></is></c>', "column"),
        ):
            entries = []
            for name, content in original:
                if name == "xl/worksheets/sheet1.xml":
                    content = re.sub(rb"<dimension[^>]*/>", b'<dimension ref="A1:A1"/>', content)
                    content = re.sub(rb'<c r="D2"[^>]*>.*?</c>', lambda _: cell, content)
                entries.append((name, content))
            self.assert_invalid(zip_bytes(entries), "forged.xlsx", message)
        malicious = []
        for name, content in original:
            if name == "xl/sharedStrings.xml":
                continue
            if name == "docProps/core.xml":
                content = b'<!DOCTYPE core [<!ENTITY x "unsafe">]><core>&x;</core>'
            malicious.append((name, content))
        self.assert_invalid(zip_bytes(malicious), "dtd.xlsx", "DTD")

    def test_qti_export_roundtrip_all_types_metadata_settings_and_feedback(self):
        for scene in scenes():
            scene.assessmentItem.metadata.description = "A description"
            scene.assessmentItem.metadata.tags = ["topic"]
            scene.assessmentItem.metadata.learningObjectiveRefs = ["objective"]
            scene.assessmentItem.metadata.language = "sv-SE"
            with self.subTest(kind=scene.assessmentItem.type):
                rows = self.assert_valid(importer.parse_questions("item.xml", assessment_item_xml(scene)))
                self.assertEqual(rows[0]["item"], scene.assessmentItem.model_dump(mode="json"))
                self.assertEqual(rows[0]["settings"], scene.assessmentSettings.model_dump(mode="json"))

    def test_qti_standard_title_language_and_missing_settings(self):
        scene = scenes()[0]
        root = ET.fromstring(assessment_item_xml(scene))
        root.set("title", "Standard title")
        root.set(importer.XML_LANG, "de-DE")
        rows = self.assert_valid(importer.parse_questions("item.xml", xml_bytes(root)))
        self.assertEqual(rows[0]["item"]["metadata"]["title"], "Standard title")
        self.assertEqual(rows[0]["item"]["metadata"]["language"], "de-DE")
        self.assertTrue(any("settings" in warning for warning in rows[0]["warnings"]))

    def test_qti_true_false_unmarked_retains_choice_semantics(self):
        scene = scenes()[2]
        rows = self.assert_valid(importer.parse_questions("item.xml", without_extension(assessment_item_xml(scene))))
        self.assertEqual(rows[0]["item"]["type"], "single-choice")
        self.assertEqual(rows[0]["item"]["correctResponse"], "true")

    def test_qti_drag_distractors_require_explicit_partition(self):
        scene = scenes()[4]
        item = scene.assessmentItem.model_dump(mode="json")
        item["draggables"].append({"identifier": "unused", "text": "Distractor"})
        scene.assessmentItem = importer.ITEM_ADAPTER.validate_python(item)
        data = assessment_item_xml(scene)
        self.assertEqual(self.assert_valid(importer.parse_questions("item.xml", data))[0]["item"],
                         scene.assessmentItem.model_dump(mode="json"))
        self.assert_invalid(without_extension(data), contains="ambiguous")
        data = assessment_item_xml(scenes()[4])
        self.assert_valid(importer.parse_questions("item.xml", without_extension(data)))

    def test_qti_short_answer_case_sensitive_and_multiple_accepted_answers(self):
        scene = scenes()[-1]
        for sensitive in (True, False):
            scene.assessmentItem.caseSensitive = sensitive
            scene.assessmentItem.acceptedAnswers = ["café", "New York", " leading space "]
            row = self.assert_valid(importer.parse_questions("item.xml", assessment_item_xml(scene)))[0]
            self.assertEqual(row["item"]["caseSensitive"], sensitive)
            self.assertEqual(row["item"]["acceptedAnswers"], scene.assessmentItem.acceptedAnswers)

    def test_fill_in_the_blank_template_labels_preserve_short_answer_type(self):
        for format in ("csv", "xlsx"):
            with self.subTest(format=format):
                row = self.assert_valid(importer.parse_questions(
                    f"template.{format}", importer.question_template(format),
                ), 7)[-1]
                self.assertEqual(row["item"]["type"], "short-answer")
                self.assertEqual(row["title"], "Example Fill-in-the-blank")
                self.assertIn("Fill-in-the-blank", row["item"]["prompt"])

    def test_inline_fill_in_the_blank_and_legacy_layout_preserve_exact_answers(self):
        scene = scenes()[-1]
        scene.assessmentItem.prompt = "Café,\ncomplete this sentence:"
        scene.assessmentItem.acceptedAnswers = [" leading space ", "New York", "café", "CAFÉ"]
        for sensitive in (True, False):
            scene.assessmentItem.caseSensitive = sensitive
            for inline in (True, False):
                for answer_key in ("primary", "all", "absent"):
                    with self.subTest(sensitive=sensitive, inline=inline, answer_key=answer_key):
                        root = ET.fromstring(assessment_item_xml(scene))
                        body = root.find(qti_tag("itemBody"))
                        body.clear()
                        paragraph = ET.SubElement(body, qti_tag("p"))
                        if inline:
                            ET.SubElement(paragraph, qti_tag("span")).text = scene.assessmentItem.prompt
                            parent = paragraph
                        else:
                            paragraph.text = scene.assessmentItem.prompt
                            parent = body
                        ET.SubElement(parent, qti_tag("textEntryInteraction"), {
                            "responseIdentifier": "RESPONSE", "expectedLength": "40",
                        })
                        declaration = root.find(qti_tag("responseDeclaration"))
                        for correct in declaration.findall(qti_tag("correctResponse")):
                            declaration.remove(correct)
                        if answer_key != "absent":
                            correct = ET.SubElement(declaration, qti_tag("correctResponse"))
                            answers = scene.assessmentItem.acceptedAnswers
                            for answer in (answers[:1] if answer_key == "primary" else answers):
                                ET.SubElement(correct, qti_tag("value")).text = answer
                        row = self.assert_valid(importer.parse_questions("blank.xml", xml_bytes(root)))[0]
                        self.assertEqual(row["item"], scene.assessmentItem.model_dump(mode="json"))

    def test_fill_in_the_blank_rejects_multiple_blanks_rich_text_and_conflicting_key(self):
        def inline_root():
            root = ET.fromstring(assessment_item_xml(scenes()[-1]))
            body = root.find(qti_tag("itemBody"))
            body.clear()
            paragraph = ET.SubElement(body, qti_tag("p"))
            ET.SubElement(paragraph, qti_tag("span")).text = "Complete the sentence"
            ET.SubElement(paragraph, qti_tag("textEntryInteraction"), {
                "responseIdentifier": "RESPONSE", "expectedLength": "40",
            })
            return root, paragraph

        root, paragraph = inline_root()
        ET.SubElement(paragraph, qti_tag("textEntryInteraction"), {"responseIdentifier": "RESPONSE"})
        self.assert_invalid(xml_bytes(root), contains="one")
        root, paragraph = inline_root()
        ET.SubElement(paragraph[0], qti_tag("b")).text = "rich"
        self.assert_invalid(xml_bytes(root))
        root, paragraph = inline_root()
        paragraph[1].tail = " Text after the blank cannot be dropped"
        self.assert_invalid(xml_bytes(root), contains="text")
        root, paragraph = inline_root()
        declaration = root.find(qti_tag("responseDeclaration"))
        for correct in declaration.findall(qti_tag("correctResponse")):
            declaration.remove(correct)
        correct = ET.SubElement(declaration, qti_tag("correctResponse"))
        ET.SubElement(correct, qti_tag("value")).text = "not accepted by scoring"
        self.assert_invalid(xml_bytes(root), contains="conflicts")

    def test_qti_match_correct_template_only(self):
        root = ET.fromstring(assessment_item_xml(scenes()[0]))
        processing = root.find(qti_tag("responseProcessing"))
        processing.clear()
        processing.set("template", "http://www.imsglobal.org/question/qti_v2p1/rptemplates/match_correct")
        self.assert_valid(importer.parse_questions("item.xml", xml_bytes(root)))
        processing.set("template", "https://evil.example/custom")
        self.assert_invalid(xml_bytes(root), contains="template")

    def test_qti_rejects_custom_scoring_and_unsupported_content(self):
        original = assessment_item_xml(scenes()[0])

        def alter(path, attribute=None, value=None, child=None):
            root = ET.fromstring(original)
            node = root.find(path)
            if attribute:
                node.set(attribute, value)
            if child:
                ET.SubElement(node, qti_tag(child)).text = "unsupported"
            return xml_bytes(root)

        variants = [
            alter(".", "adaptive", "true"),
            alter(".", "timeDependent", "true"),
            alter(qti_tag("itemBody"), child="img"),
            alter(f".//{qti_tag('simpleChoice')}", child="b"),
            alter(qti_tag("responseDeclaration"), child="mapping"),
            alter(qti_tag("responseProcessing"), child="customOperator"),
            alter(qti_tag("outcomeDeclaration"), "normalMaximum", "2"),
            alter(f".//{qti_tag('choiceInteraction')}", "maxChoices", "2"),
            alter(f".//{qti_tag('choiceInteraction')}", "minChoices", "1"),
            alter(f".//{qti_tag('simpleChoice')}", "fixed", "true"),
            alter(f".//{qti_tag('modalFeedback')}", child="p"),
        ]
        for content in variants:
            with self.subTest(content=content):
                self.assert_invalid(content)
        root = ET.fromstring(original)
        root.find(f".//{qti_tag('choiceInteraction')}").tag = qti_tag("hotspotInteraction")
        self.assert_invalid(xml_bytes(root), contains="unsupported")
        root = ET.fromstring(original)
        root.find(f".//{qti_tag('setOutcomeValue')}/{qti_tag('baseValue')}").text = "0.5"
        self.assert_invalid(xml_bytes(root), contains="scoring")

    def test_shuffle_and_standard_metadata_warn_and_remediation_is_dropped(self):
        scene = scenes()[0]
        scene.assessmentSettings.remediationSceneId = "nonportable-scene"
        data = assessment_item_xml(scene)
        encoded_extension = data.split(b"<?chimera ", 1)[1].split(b"?>", 1)[0]
        root = ET.fromstring(data)
        root.find(f".//{qti_tag('choiceInteraction')}").set("shuffle", "true")
        metadata = ET.SubElement(root, qti_tag("qtiMetadata"))
        ET.SubElement(metadata, qti_tag("field")).text = "not preserved"
        data = with_extension(xml_bytes(root), json.loads(base64.b64decode(encoded_extension)))
        row = self.assert_valid(importer.parse_questions("item.xml", data))[0]
        self.assertIsNone(row["settings"]["remediationSceneId"])
        for text in ("shuffle", "metadata", "Remediation"):
            self.assertTrue(any(text in warning for warning in row["warnings"]), row)

    def test_qti_invalid_extension_and_type_conflicts(self):
        original = assessment_item_xml(scenes()[0])
        for extension in (
            {"type": "matching"},
            {"metadata": {"unknown": "value"}},
            {"metadata": []},
            {"settings": {"maxAttempts": 0}},
            {"settings": {"unknown": 2}},
            {"dragSources": ["a", "a"]},
            {"dragSources": ["a"]},
            {"unknown": "value"},
        ):
            with self.subTest(extension=extension):
                self.assert_invalid(with_extension(original, extension))
        self.assert_invalid(original.replace(b"<?chimera ", b"<?chimera @"), contains="encoding")
        duplicate = original.replace(b"<?chimera ", b"<?chimera e30=?>\n<?chimera ")
        self.assert_invalid(duplicate, contains="multiple")

    def test_xml_entities_dtd_and_non_item_roots_are_rejected(self):
        for content in (
            b'<!DOCTYPE assessmentItem [<!ENTITY x "expansion">]><assessmentItem>&x;</assessmentItem>',
            b'<!DOCTYPE assessmentItem SYSTEM "file:///C:/sensitive"><assessmentItem/>',
            b'<!DOCTYPE assessmentItem><assessmentItem/>',
            b"<malformed",
            b'<assessmentItem xmlns="http://www.imsglobal.org/xsd/imsqti_v2p2"/>',
            f'<assessmentTest xmlns="{importer.QTI}"/>'.encode(),
        ):
            self.assert_invalid(content)

    def test_package_uses_test_order_then_declared_unreferenced_items(self):
        questions = scenes()[:3]
        data = assessment_test_package(questions, "Test")
        with ZipFile(BytesIO(data)) as archive:
            entries = [(name, archive.read(name)) for name in archive.namelist()]
        changed = []
        for name, content in entries:
            if name == "assessmentTest.xml":
                root = ET.fromstring(content)
                section = root.find(f".//{qti_tag('assessmentSection')}")
                refs = list(section)
                section.remove(refs[0])
                section.remove(refs[1])
                section.remove(refs[2])
                section.extend([refs[2], refs[0]])
                content = xml_bytes(root)
            changed.append((name, content))
        changed.append(("assets/not-a-question.xml", b"<asset/>"))
        rows = self.assert_valid(importer.parse_questions("package.zip", zip_bytes(changed)), 3)
        self.assertEqual([row["item"]["identifier"] for row in rows],
                         [questions[i].assessmentItem.identifier for i in (2, 0, 1)])
        self.assertTrue(all(any("test-level" in warning for warning in row["warnings"]) for row in rows))

    def test_manifest_resource_order_without_test_and_safe_relative_test_refs(self):
        q = scenes()
        cp = importer.CP
        manifest = f'''<manifest xmlns="{cp}" identifier="m"><resources>
          <resource identifier="b" type="imsqti_item_xmlv2p1" href="items/b.xml"/>
          <resource identifier="a" type="imsqti_item_xmlv2p1" href="items/a.xml"/>
        </resources></manifest>'''.encode()
        entries = [("imsmanifest.xml", manifest), ("items/a.xml", assessment_item_xml(q[0])),
                   ("items/b.xml", assessment_item_xml(q[1])), ("asset.xml", b"<anything/>")]
        rows = self.assert_valid(importer.parse_questions("package.zip", zip_bytes(entries)), 2)
        self.assertEqual([row["item"]["identifier"] for row in rows],
                         [q[1].assessmentItem.identifier, q[0].assessmentItem.identifier])
        manifest = manifest.replace(b"<resources>", b'<resources><resource identifier="t" type="imsqti_test_xmlv2p1" href="tests/test.xml"/>')
        test = f'''<assessmentTest xmlns="{importer.QTI}"><testPart>
            <assessmentSection><assessmentItemRef href="../items/a.xml"/></assessmentSection>
            </testPart></assessmentTest>'''.encode()
        entries[0] = ("imsmanifest.xml", manifest)
        entries.append(("tests/test.xml", test))
        rows = self.assert_valid(importer.parse_questions("package.zip", zip_bytes(entries)), 2)
        self.assertEqual(rows[0]["item"]["identifier"], q[0].assessmentItem.identifier)

    def test_zip_unsafe_paths_duplicates_encryption_and_symbolic_links(self):
        for path in ("../escape.xml", "/absolute.xml", "C:/absolute.xml", "a\\b.xml", "./item.xml", "a//b.xml"):
            with self.subTest(path=path):
                data = zip_bytes([(path, b"x")])
                if "\\" in path:
                    data = data.replace(b"a/b.xml", b"a\\b.xml")
                self.assert_invalid(data, "bad.zip", "unsafe")
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            data = zip_bytes([("same.xml", b"x"), ("same.xml", b"y")])
        self.assert_invalid(data, "duplicate.zip", "duplicate")
        data = bytearray(zip_bytes([("item.xml", b"x")]))
        local = data.find(b"PK\x03\x04")
        central = data.find(b"PK\x01\x02")
        data[local + 6] |= 1
        data[central + 8] |= 1
        self.assert_invalid(bytes(data), "encrypted.zip", "encrypted")
        info = ZipInfo("link.xml")
        info.external_attr = (0o120777 << 16)
        self.assert_invalid(zip_bytes([(info, b"target")]), "link.zip", "symbolic")

    def test_package_rejects_missing_manifest_broken_and_external_references(self):
        self.assert_invalid(zip_bytes([("item.xml", assessment_item_xml(scenes()[0]))]), "missing.zip", "manifest")
        for reference in ("https://example.com/item.xml", "../../escape.xml", "%2e%2e/escape.xml", "missing.xml", "item.xml?x=1"):
            manifest = f'''<manifest xmlns="{importer.CP}"><resources>
              <resource identifier="a" type="imsqti_item_xmlv2p1" href="{reference}"/>
            </resources></manifest>'''.encode()
            with self.subTest(reference=reference):
                self.assert_invalid(zip_bytes([("imsmanifest.xml", manifest)]), "bad.zip")
        self.assert_invalid(b"not a zip", "bad.zip")
        self.assert_invalid(b"not xlsx", "bad.xlsx")
        self.assert_invalid(b"no", "bad.doc")

    def test_limits_upload_questions_zip_members_and_uncompressed_size(self):
        self.assert_invalid(b"x" * (importer.MAX_UPLOAD + 1), "large.csv", "10 MB")
        headers = ["identifier", "type", "title", "prompt", "answers"]
        data = csv_bytes(headers, [
            [f"q_{index}", "short-answer", "Title", "Prompt", '["answer"]']
            for index in range(501)
        ])
        self.assert_invalid(data, "many.csv", "500")
        data = zip_bytes([(f"entry_{index}", b"") for index in range(2001)])
        self.assert_invalid(data, "entries.zip", "2000")
        # A low patched limit exercises the same cumulative central-directory
        # accounting without allocating a large ZIP bomb fixture.
        with patch.object(importer, "MAX_UNCOMPRESSED", 100):
            self.assert_invalid(zip_bytes([("a", b"x" * 51), ("b", b"x" * 51)]),
                                "large.zip", "uncompressed")


if __name__ == "__main__":
    unittest.main()
