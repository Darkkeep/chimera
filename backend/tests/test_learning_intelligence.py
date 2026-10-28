import os
import json
import tempfile
import time
import unittest
from contextlib import ExitStack
from io import BytesIO
from pathlib import Path
from unittest.mock import Mock, patch

import httpx
from jsonschema import Draft202012Validator
from fastapi import HTTPException
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Inches
from starlette.requests import Request

from backend.app import main, storage
from backend.app.learning_intelligence import PythiaProvider, PythiaUnavailableError, validate_assessment_blueprint, validate_course_review, validate_storyboard_analysis, validate_storyboard_transformation
from backend.app.models import ProjectCreate, PythiaAnalysisProposalCreate, PythiaAssessmentProposalCreate, SceneCreate


def transformation_envelope(project_id: str, scene_ids: list[str]) -> dict[str, object]:
    source_refs = [f"chimera://{project_id}/scenes/{scene_id}" for scene_id in scene_ids]
    return {
        "schema_version": "1.0",
        "artifact": {"type": "storyboard-transformation", "id": "artifact-transform-1", "version": 1},
        "producer": {"skill": "storyboard-transformer", "version": "1.0"},
        "sources": source_refs,
        "status": "draft",
        "validation": {"state": "passed", "validators": ["schema"]},
        "payload": {
            "project_id": project_id,
            "source_scene_ids": scene_ids,
            "source_refs": source_refs,
            "rationale": "Restructure the selected content into clear steps.",
            "proposed_scenes": [{
                "proposal_scene_id": "proposed-scene-1",
                "source_scene_ids": [scene_ids[0]],
                "purpose": "State the immediate action.",
                "title": "Take the immediate action",
                "learner_content": ["Use the nearest marked exit."],
                "interaction_recommendation": "Ask the learner to select an exit.",
                "production_notes": "Emphasize the assembly location.",
                "media_requirements": [],
                "assessment_opportunities": [],
                "rationale": "Keep the first action prominent.",
            }, {
                "proposal_scene_id": "proposed-scene-2",
                "source_scene_ids": [scene_ids[-1]],
                "purpose": "Explain the follow-up action.",
                "title": "Continue to the assembly point",
                "learner_content": ["Leave personal items behind."],
                "interaction_recommendation": "Ask the learner what to do after exiting.",
                "production_notes": "Name the local assembly point.",
                "media_requirements": ["A site-specific assembly point map."],
                "assessment_opportunities": ["Identify the correct follow-up action."],
                "rationale": "Separate follow-up guidance from evacuation steps.",
            }],
        },
    }


def assessment_blueprint_envelope(project_id: str, scene_ids: list[str], objective_ids: list[str]) -> dict[str, object]:
    return {
        "schema_version": "1.0",
        "artifact": {"type": "assessment-blueprint", "id": "artifact-assessment-1", "version": 1},
        "producer": {"skill": "assessment-designer", "version": "1.0"},
        "sources": [f"chimera://{project_id}/scenes/{scene_id}" for scene_id in scene_ids],
        "status": "draft",
        "validation": {"state": "passed", "validators": ["assessment-blueprint-1.0"]},
        "payload": {
            "project_id": project_id,
            "candidate_interpretations": [],
            "unresolved_findings": [],
            "assumptions": ["Review the distractors with an SME."],
            "alignment": [{
                "objective_id": objective_ids[0],
                "evidence": "The learner selects a safe response.",
                "method": "Single choice",
                "scoring": "One correct answer",
                "fairness": "Use plain language.",
            }],
            "ai_use_boundaries": "AI drafts items; SMEs approve scoring.",
            "limitations": ["No audience research was supplied."],
            "review_questions": ["Confirm the local procedure."],
            "proposed_items": [{
                "proposal_item_id": "proposal-item-1",
                "target_scene_id": scene_ids[0],
                "objective_ids": [objective_ids[0]],
                "item": {
                    "type": "single-choice",
                    "identifier": "evacuation-response",
                    "prompt": "What should you do when the alarm sounds?",
                    "choices": [
                        {"identifier": "CHOICE_A", "text": "Use the nearest marked exit."},
                        {"identifier": "CHOICE_B", "text": "Collect personal items first."},
                    ],
                    "correctResponse": "CHOICE_A",
                    "correctFeedback": "Correct.",
                    "incorrectFeedback": "Leave personal items behind.",
                },
                "rationale": "Checks the stated evacuation objective.",
                "evidence_refs": [f"chimera://{project_id}/scenes/{scene_ids[0]}"],
            }],
        },
    }


def course_review_envelope(project_id: str, inventory: dict[str, object]) -> dict[str, object]:
    course_ref = inventory["course_ref"]
    source_refs = [course_ref, *(scene["source_ref"] for scene in inventory["scenes"])]
    domains = ["objectives", "instructional_alignment", "assessment", "accessibility", "source_currency", "standards", "tracking", "packaging", "evaluation"]
    return {
        "schema_version": "1.0",
        "artifact": {"type": "course-review", "id": "artifact-course-review-1", "version": 1},
        "producer": {"skill": "kernel-orchestrator", "version": "1.0"},
        "sources": source_refs,
        "status": "draft",
        "validation": {"state": "passed", "validators": ["course-review-1.0"]},
        "payload": {
            "project_id": project_id,
            "course_ref": course_ref,
            "domain_status": [{
                "domain": domain,
                "status": "review_required" if domain == "standards" else "insufficient_evidence" if domain == "evaluation" else "ready",
                "finding_count": 1 if domain == "standards" else 0,
            } for domain in domains],
            "findings": [{
                "finding_id": "standards-1",
                "category": "standards",
                "severity": "review",
                "statement": "Standards selection is unavailable.",
                "object_refs": [{"object_type": "course", "object_id": course_ref}],
                "evidence_refs": [course_ref],
                "recommendation": "Select or document the delivery standard.",
                "limitations": ["No delivery target was submitted."],
            }],
            "specialist_routes": [{"domain": "standards", "capability": "standards.validate", "provider": "standards-architect"}],
            "limitations": ["Only the submitted project inventory was reviewed."],
            "review_questions": ["Which delivery standard is required?"],
        },
    }


def assert_contract_schema(test_case: unittest.TestCase, schema_name: str, document: dict[str, object]) -> None:
    schema_path = Path(__file__).resolve().parents[2] / "vendor" / "pythia-contracts" / "schemas" / schema_name
    test_case.assertTrue(schema_path.is_file(), f"Shared contract schema is required: {schema_path}")
    schema = json.loads(schema_path.read_text(encoding="utf-8"))
    errors = sorted(
        Draft202012Validator(schema).iter_errors(document),
        key=lambda error: list(error.absolute_path),
    )
    test_case.assertEqual([], [error.message for error in errors], schema_name)


class LearningIntelligenceTests(unittest.TestCase):
    def test_provider_defaults_to_local_engine_api(self):
        with patch.dict("os.environ", {}, clear=True):
            provider = PythiaProvider()

        self.assertEqual(provider.base_url, "http://localhost:8000")

    def test_provider_rejects_unsafe_base_urls(self):
        for base_url in (
            "ftp://pythia.example",
            "https://user:secret@pythia.example",
            "https://pythia.example?token=secret",
            "https://pythia.example#fragment",
        ):
            with self.subTest(base_url=base_url), self.assertRaises(ValueError):
                PythiaProvider(base_url)

    def test_submit_task_sends_contract_to_engine_api(self):
        response = Mock()
        response.json.return_value = {"task_id": "chimera-test", "status": "submitted"}
        with patch("backend.app.learning_intelligence.httpx.request", return_value=response) as request:
            result = PythiaProvider("http://localhost:8000/").submit_task({"task_id": "chimera-test"})

        self.assertEqual(result["status"], "submitted")
        request.assert_called_once_with(
            "POST", "http://localhost:8000/tasks", json={"task_id": "chimera-test"}, timeout=10.0,
        )

    def test_provider_sends_server_configured_authentication_to_engine(self):
        response = Mock()
        response.json.return_value = {"status": "ok"}
        with patch.dict("os.environ", {"CHIMERA_PYTHIA_API_KEY": "chimera-client-token"}), patch(
            "backend.app.learning_intelligence.httpx.request", return_value=response,
        ) as request:
            self.assertEqual({"status": "ok"}, PythiaProvider("http://localhost:8001").health())

        request.assert_called_once_with(
            "GET", "http://localhost:8001/health", json=None, timeout=10.0,
            headers={"Authorization": "Bearer chimera-client-token"},
        )

    def test_provider_diagnostics_read_engine_registry_and_telemetry(self):
        responses = [Mock(), Mock()]
        responses[0].json.return_value = {"providers": []}
        responses[1].json.return_value = {"events": []}
        with patch("backend.app.learning_intelligence.httpx.request", side_effect=responses) as request:
            provider = PythiaProvider("http://localhost:8000")
            self.assertEqual({"providers": []}, provider.providers())
            self.assertEqual({"events": []}, provider.telemetry())

        self.assertEqual("GET", request.call_args_list[0].args[0])
        self.assertEqual("http://localhost:8000/providers", request.call_args_list[0].args[1])
        self.assertEqual("http://localhost:8000/telemetry", request.call_args_list[1].args[1])

    def test_intelligence_engine_status_exposes_read_only_engine_diagnostics(self):
        provider = Mock()
        provider.health.return_value = {"status": "ok", "version": "1.0"}
        provider.capabilities.return_value = {"capabilities": {"storyboard.enhance": {"provider": "storyboard-producer"}}}
        provider.providers.return_value = {"budget": {"max_concurrent_tasks": 4}, "providers": []}
        provider.telemetry.return_value = {"events": []}
        policy = Mock()
        policy.allows.return_value = True
        with patch.object(main, "_pythia_provider", return_value=provider), patch.object(
            main, "get_authorization_policy", return_value=policy,
        ):
            status = main.get_intelligence_engine()

        self.assertTrue(status["available"])
        self.assertIn("storyboard.enhance", status["capabilities"])
        self.assertEqual(4, status["providerRegistry"]["budget"]["max_concurrent_tasks"])
        self.assertEqual([], status["telemetry"]["events"])
        self.assertTrue(status["canViewUsage"])

    def test_intelligence_engine_hides_global_usage_from_non_admins(self):
        provider = Mock()
        provider.health.return_value = {"status": "ok", "version": "1.0"}
        provider.capabilities.return_value = {"capabilities": {"storyboard.enhance": {}}}
        provider.providers.return_value = {"providers": []}
        policy = Mock()
        policy.allows.return_value = False
        with patch.object(main, "_pythia_provider", return_value=provider), patch.object(
            main, "get_authorization_policy", return_value=policy,
        ):
            status = main.get_intelligence_engine()

        self.assertFalse(status["canViewUsage"])
        self.assertEqual([], status["telemetry"]["events"])
        provider.telemetry.assert_not_called()

    def test_http_errors_preserve_status_for_route_translation(self):
        response = Mock()
        response.status_code = 409
        response.text = "duplicate task ID"
        response.headers = {"X-Pythia-Request-Id": "engine-request-123"}
        response.raise_for_status.side_effect = httpx.HTTPStatusError("conflict", request=Mock(), response=response)
        with patch("backend.app.learning_intelligence.httpx.request", return_value=response):
            with self.assertRaises(PythiaUnavailableError) as raised:
                PythiaProvider().health()

        self.assertEqual(raised.exception.status_code, 409)
        self.assertEqual(raised.exception.request_id, "engine-request-123")
        self.assertIn("Engine request ID: engine-request-123", str(raised.exception))
        translated = main._pythia_error(raised.exception)
        self.assertEqual(translated.status_code, 409)
        self.assertEqual(translated.headers["X-Pythia-Request-Id"], "engine-request-123")

    def test_engine_auth_limit_and_capacity_statuses_remain_recoverable(self):
        for status_code in (401, 403, 413, 422, 429, 503):
            with self.subTest(status_code=status_code):
                translated = main._pythia_error(PythiaUnavailableError("Engine request failed", status_code))
                self.assertEqual(status_code, translated.status_code)

    def test_storyboard_analysis_envelope_must_match_scene_and_schema(self):
        envelope = {
            "artifact": {"type": "storyboard-analysis"},
            "status": "draft",
            "validation": {"state": "passed"},
            "producer": {"skill": "storyboard-producer"},
            "payload": {
                "project_id": "course-123",
                "scene_id": "scene_a81f73",
                "source_refs": [],
                "findings": [],
                "recommendations": [],
                "uncertainties": [],
                "review_questions": [],
            },
        }

        self.assertIs(validate_storyboard_analysis(envelope, "course-123", "scene_a81f73"), envelope)
        with self.assertRaisesRegex(ValueError, "does not match"):
            validate_storyboard_analysis(envelope, "other-course", "scene_a81f73")

    def test_storyboard_analysis_rejects_unstructured_findings(self):
        envelope = {
            "artifact": {"type": "storyboard-analysis"},
            "status": "draft",
            "validation": {"state": "passed"},
            "producer": {"skill": "storyboard-producer"},
            "payload": {
                "project_id": "course-123",
                "scene_id": "scene_a81f73",
                "source_refs": [],
                "findings": ["Identifier-only result"],
                "recommendations": [],
                "uncertainties": [],
                "review_questions": [],
            },
        }

        with self.assertRaisesRegex(ValueError, "findings must contain objects"):
            validate_storyboard_analysis(envelope, "course-123", "scene_a81f73")

    def test_storyboard_transformation_requires_a_valid_selected_scene_proposal(self):
        source_ref = "chimera://course-123/scenes/scene_a81f73"
        envelope = {
            "schema_version": "1.0",
            "artifact": {"type": "storyboard-transformation", "id": "artifact-1", "version": 1},
            "producer": {"skill": "storyboard-transformer", "version": "1.0"},
            "sources": [source_ref],
            "status": "draft",
            "validation": {"state": "passed", "validators": ["schema"]},
            "payload": {
                "project_id": "course-123",
                "source_scene_ids": ["scene_a81f73"],
                "source_refs": [source_ref],
                "rationale": "Split the instruction into a sequence.",
                "proposed_scenes": [{
                    "proposal_scene_id": "proposed-scene-1",
                    "source_scene_ids": ["scene_a81f73"],
                    "purpose": "Explain the immediate response.",
                    "title": "Leave immediately",
                    "learner_content": ["Use the nearest marked exit."],
                    "interaction_recommendation": "Ask the learner to identify an exit.",
                    "production_notes": "Emphasize the assembly location.",
                    "media_requirements": [],
                    "assessment_opportunities": [],
                    "rationale": "Keep the immediate action prominent.",
                }],
            },
        }

        assert_contract_schema(self, "artifact-envelope.schema.json", envelope)
        assert_contract_schema(self, "storyboard-transformation.schema.json", envelope["payload"])
        self.assertIs(validate_storyboard_transformation(envelope, "course-123", ["scene_a81f73"]), envelope)
        with self.assertRaisesRegex(ValueError, "selected scenes"):
            validate_storyboard_transformation(envelope, "course-123", ["other-scene"])
        invalid_sources = {**envelope, "payload": {**envelope["payload"], "source_refs": ["chimera://other/scenes/other"]}}
        with self.assertRaisesRegex(ValueError, "source_refs"):
            validate_storyboard_transformation(invalid_sources, "course-123", ["scene_a81f73"])

    def test_assessment_blueprint_validates_native_item_and_scope_references(self):
        envelope = assessment_blueprint_envelope("course-123", ["scene_a81f73"], ["objective-evacuation"])

        assert_contract_schema(self, "artifact-envelope.schema.json", envelope)
        assert_contract_schema(self, "assessment-blueprint.schema.json", envelope["payload"])
        self.assertIs(validate_assessment_blueprint(envelope, "course-123", ["scene_a81f73"], ["objective-evacuation"]), envelope)
        invalid_item = {**envelope, "payload": {**envelope["payload"], "proposed_items": [{
            **envelope["payload"]["proposed_items"][0],
            "target_scene_id": "other-scene",
        }]}}
        with self.assertRaisesRegex(ValueError, "outside the selected scope"):
            validate_assessment_blueprint(invalid_item, "course-123", ["scene_a81f73"], ["objective-evacuation"])
        invalid_objective = {**envelope, "payload": {**envelope["payload"], "proposed_items": [{
            **envelope["payload"]["proposed_items"][0],
            "objective_ids": ["unsubmitted-objective"],
        }]}}
        with self.assertRaisesRegex(ValueError, "supplied objectives"):
            validate_assessment_blueprint(invalid_objective, "course-123", ["scene_a81f73"], ["objective-evacuation"])

    def test_course_review_validator_rejects_object_references_outside_inventory(self):
        inventory = {
            "course_ref": "chimera://course-123/course",
            "objectives": [],
            "scenes": [{"scene_id": "scene-1", "source_ref": "chimera://course-123/scenes/scene-1"}],
            "assessment_items": [],
            "assets": [],
            "source_records": [],
        }
        envelope = course_review_envelope("course-123", inventory)

        assert_contract_schema(self, "artifact-envelope.schema.json", envelope)
        assert_contract_schema(self, "course-review.schema.json", envelope["payload"])
        self.assertIs(validate_course_review(envelope, "course-123", inventory), envelope)
        invalid = {**envelope, "payload": {**envelope["payload"], "findings": [{
            **envelope["payload"]["findings"][0],
            "object_refs": [{"object_type": "scene", "object_id": "unknown-scene"}],
        }]}}
        with self.assertRaisesRegex(ValueError, "outside the submitted scope"):
            validate_course_review(invalid, "course-123", inventory)
        invalid_evidence = {**envelope, "payload": {**envelope["payload"], "findings": [{
            **envelope["payload"]["findings"][0],
            "evidence_refs": [{"untrusted": "object"}],
        }]}}
        with self.assertRaisesRegex(ValueError, "evidence_refs"):
            validate_course_review(invalid_evidence, "course-123", inventory)


class PythiaRouteTests(unittest.TestCase):
    def setUp(self):
        self.stack = ExitStack()
        self.addCleanup(self.stack.close)
        self.root = Path(self.stack.enter_context(tempfile.TemporaryDirectory()))
        self.stack.enter_context(patch.multiple(storage, ROOT=self.root, PROJECTS=self.root / "projects"))
        self.stack.enter_context(patch.object(main, "ROOT", self.root))
        self.project = storage.create_project(ProjectCreate(title="Pythia fixture"))
        self.project = storage.create_scene(self.project.id, SceneCreate(
            title="Stable scene",
            extractedText=["First screen text", "Second screen text"],
            actionDescription="Choose the nearest exit.",
            notes="Emphasize the local assembly point.",
        ))
        self.scene = self.project.storyboard.scenes[0]

    def test_hosted_pythia_proposal_routes_require_review_permissions(self):
        def request_for(path):
            return Request({
                "type": "http",
                "method": "POST",
                "path": path,
                "headers": [],
                "query_string": b"",
                "scheme": "http",
                "server": ("testserver", 80),
            })

        self.assertEqual(
            main._hosted_request_action(request_for(f"/api/projects/{self.project.id}/pythia-analysis-proposals")),
            "enhancement.proposal.create",
        )
        self.assertEqual(
            main._hosted_request_action(request_for(f"/api/projects/{self.project.id}/pythia-analysis-proposals/proposal-1/decision")),
            "enhancement.proposal.decide",
        )

    def test_submit_uses_only_the_scene_scoped_task_contract(self):
        provider = Mock()
        provider.health.return_value = {"status": "ok", "version": "1.0"}
        provider.capabilities.return_value = {"capabilities": [{"id": "storyboard.analyze", "provider": "synthetic"}]}
        provider.submit_task.side_effect = lambda request: {"task_id": request["task_id"], "status": "submitted"}
        with patch.object(main, "PythiaProvider", return_value=provider):
            task = main.submit_pythia_analysis(self.project.id, self.scene.id)

        request = provider.submit_task.call_args.args[0]
        self.assertEqual(task["status"], "submitted")
        self.assertEqual(task["source_revision"], self.project.storyboard.revision)
        self.assertEqual(request["schema_version"], "1.0")
        self.assertEqual(request["capability"], "storyboard.analyze")
        self.assertEqual(request["consumer"], "chimera")
        self.assertEqual(request["project_id"], self.project.id)
        self.assertEqual(request["inputs"], [{
            "scene_id": self.scene.id,
            "source_ref": f"chimera://{self.project.id}/scenes/{self.scene.id}",
            "scene_content": {
                "title": "Stable scene",
                "body_text": "First screen text\nSecond screen text\nChoose the nearest exit.",
                "speaker_notes": "Emphasize the local assembly point.",
                "objective_ids": [],
            },
        }])
        self.assertEqual(request["constraints"], {})
        self.assertEqual(request["context"], {})
        assert_contract_schema(self, "task-request.schema.json", request)
        assert_contract_schema(self, "storyboard-analysis-input.schema.json", request["inputs"][0])

    @unittest.skipUnless(
        all(os.environ.get(name) for name in (
            "CHIMERA_PYTHIA_BASE_URL",
            "CHIMERA_PYTHIA_API_KEY",
            "CHIMERA_MANTEION_API_KEY",
            "CHIMERA_FOREIGN_API_KEY",
        )),
        "Set isolated qualification credentials to run the live Pythia interoperability check.",
    )
    def test_live_pythia_candidate_round_trip_and_cross_client_provenance(self):
        base_url = os.environ["CHIMERA_PYTHIA_BASE_URL"]
        provider = PythiaProvider(base_url)
        self.assertEqual(provider.health().get("contract_version"), "1.1.0")

        submitted = main.submit_pythia_analysis(self.project.id, self.scene.id)
        self.assertEqual(submitted["consumer"], "chimera")
        self.assertEqual(submitted["status"], "submitted")
        task_id = submitted["task_id"]

        task = provider.task_status(task_id)
        for _ in range(100):
            if task.get("status") == "awaiting_review":
                break
            self.assertNotEqual(task.get("status"), "failed", task)
            time.sleep(0.025)
            task = provider.task_status(task_id)
        self.assertEqual(task.get("status"), "awaiting_review")
        artifact_id = task["artifact_id"]
        artifact = provider.artifact(artifact_id)
        validate_storyboard_analysis(artifact, self.project.id, self.scene.id)
        self.assertEqual(
            artifact["sources"],
            [f"chimera://{self.project.id}/scenes/{self.scene.id}"],
        )

        with patch.dict(os.environ, {"CHIMERA_PYTHIA_API_KEY": os.environ["CHIMERA_MANTEION_API_KEY"]}):
            manteion = PythiaProvider(base_url)
            self.assertEqual(manteion.task_status(task_id), task)
            self.assertEqual(manteion.artifact(artifact_id), artifact)

        with patch.dict(os.environ, {"CHIMERA_PYTHIA_API_KEY": os.environ["CHIMERA_FOREIGN_API_KEY"]}):
            foreign = PythiaProvider(base_url)
            with self.assertRaises(PythiaUnavailableError) as raised:
                foreign.artifact(artifact_id)
        self.assertEqual(raised.exception.status_code, 404)

    def test_transform_submit_preserves_selected_scene_order_and_content(self):
        self.project = storage.create_scene(self.project.id, SceneCreate(
            title="Second scene",
            extractedText=["Keep the route clear."],
            notes="Connect this instruction to the alarm.",
        ))
        second_scene = self.project.storyboard.scenes[-1]
        provider = Mock()
        provider.health.return_value = {"status": "ok", "version": "1.0"}
        provider.capabilities.return_value = {"capabilities": [{"id": "storyboard.transform"}]}
        provider.submit_task.side_effect = lambda request: {"task_id": request["task_id"], "status": "submitted"}
        with patch.object(main, "PythiaProvider", return_value=provider):
            status = main.get_pythia_status(self.project.id)
            task = main.submit_pythia_transform(self.project.id, main.PythiaTransformRequest(
                sceneIds=[second_scene.id, self.scene.id],
                constraints={"audience": "new staff"},
            ))

        request = provider.submit_task.call_args.args[0]
        self.assertTrue(status["supportsStoryboardTransform"])
        self.assertEqual(task["source_revision"], self.project.storyboard.revision)
        self.assertEqual(task["source_scene_ids"], [self.scene.id, second_scene.id])
        self.assertEqual(request["capability"], "storyboard.transform")
        self.assertEqual([item["scene_id"] for item in request["inputs"]], [self.scene.id, second_scene.id])
        self.assertEqual(request["inputs"][1]["scene_content"]["body_text"], "Keep the route clear.")
        self.assertEqual(request["inputs"][1]["scene_content"]["speaker_notes"], "Connect this instruction to the alarm.")
        self.assertEqual(request["constraints"], {"audience": "new staff"})
        assert_contract_schema(self, "task-request.schema.json", request)
        for scene_input in request["inputs"]:
            assert_contract_schema(self, "storyboard-transform-input.schema.json", scene_input)

    def test_transform_submit_refuses_unadvertised_capability(self):
        provider = Mock()
        provider.health.return_value = {"status": "ok"}
        provider.capabilities.return_value = {"capabilities": [{"id": "storyboard.analyze"}]}
        with patch.object(main, "PythiaProvider", return_value=provider):
            with self.assertRaises(HTTPException) as raised:
                main.submit_pythia_transform(self.project.id, main.PythiaTransformRequest(sceneIds=[self.scene.id]))

        self.assertEqual(raised.exception.status_code, 422)
        provider.submit_task.assert_not_called()

    def test_assessment_design_submits_canonical_capability_and_context(self):
        provider = Mock()
        provider.health.return_value = {"status": "ok", "version": "1.0"}
        provider.capabilities.return_value = {"capabilities": [{"id": "assessment.design"}]}
        provider.submit_task.side_effect = lambda request: {"task_id": request["task_id"], "status": "submitted"}
        request_data = main.PythiaAssessmentDesignRequest(
            sceneIds=[self.scene.id],
            objectives=[main.PythiaAssessmentObjective(id="objective-evacuation", statement="Choose a safe exit.", sceneIds=[self.scene.id])],
            audience="New staff",
            desiredDifficulty="application",
            completionRules="One correct response is required.",
        )
        with patch.object(main, "PythiaProvider", return_value=provider):
            status = main.get_pythia_status(self.project.id)
            task = main.submit_pythia_assessment_design(self.project.id, request_data)

        submitted = provider.submit_task.call_args.args[0]
        self.assertTrue(status["supportsAssessmentDesign"])
        self.assertEqual(status["assessmentCapability"], "assessment.design")
        self.assertEqual(task["capability"], "assessment.design")
        self.assertEqual(submitted["capability"], "assessment.design")
        self.assertEqual(submitted["inputs"][0]["scene_id"], self.scene.id)
        self.assertEqual(submitted["context"]["objectives"], [{"id": "objective-evacuation", "statement": "Choose a safe exit."}])
        self.assertEqual(submitted["context"]["audience"], "New staff")
        self.assertEqual(submitted["context"]["existing_items"], [{"scene_id": self.scene.id, "item": None}])
        assert_contract_schema(self, "task-request.schema.json", submitted)
        assert_contract_schema(self, "storyboard-transform-input.schema.json", submitted["inputs"][0])
        assert_contract_schema(self, "assessment-design-context.schema.json", submitted["context"])

    def test_assessment_design_refuses_unadvertised_capability(self):
        provider = Mock()
        provider.health.return_value = {"status": "ok"}
        provider.capabilities.return_value = {"capabilities": [{"id": "storyboard.transform"}]}
        with patch.object(main, "PythiaProvider", return_value=provider):
            with self.assertRaises(HTTPException) as raised:
                main.submit_pythia_assessment_design(self.project.id, main.PythiaAssessmentDesignRequest(
                    sceneIds=[self.scene.id],
                    objectives=[main.PythiaAssessmentObjective(id="objective-1", statement="Complete the action.", sceneIds=[self.scene.id])],
                    audience="New staff",
                    desiredDifficulty="application",
                    completionRules="One correct response is required.",
                ))

        self.assertEqual(raised.exception.status_code, 422)
        provider.submit_task.assert_not_called()

    def test_course_review_submits_only_inventory_fields_chimera_owns(self):
        provider = Mock()
        provider.health.return_value = {"status": "ok", "version": "1.0"}
        provider.capabilities.return_value = {"capabilities": [{"id": "course.review"}]}
        provider.submit_task.side_effect = lambda request: {"task_id": request["task_id"], "status": "submitted"}
        with patch.object(main, "PythiaProvider", return_value=provider):
            status = main.get_pythia_status(self.project.id)
            task = main.submit_pythia_course_review(self.project.id)

        request = provider.submit_task.call_args.args[0]
        inventory = request["inputs"][0]
        self.assertTrue(status["supportsCourseReview"])
        self.assertEqual(request["capability"], "course.review")
        self.assertEqual(task["source_revision"], self.project.storyboard.revision)
        self.assertEqual(inventory["course_ref"], f"chimera://{self.project.id}/course")
        self.assertEqual(inventory["objectives"], [])
        self.assertEqual(inventory["scenes"][0]["scene_id"], self.scene.id)
        self.assertEqual(inventory["scenes"][0]["objective_ids"], [])
        self.assertEqual(inventory["standards_status"], "unavailable")
        self.assertEqual(inventory["tracking_status"], "unavailable")
        assert_contract_schema(self, "task-request.schema.json", request)
        assert_contract_schema(self, "course-review-input.schema.json", inventory)

    def test_course_review_persists_findings_and_resolution_without_storyboard_mutation(self):
        inventory = main._pythia_course_inventory(self.project)
        provider = Mock()
        provider.task_status.return_value = {
            "task_id": "task-course-review-1",
            "project_id": self.project.id,
            "consumer": "chimera",
            "capability": "course.review",
            "status": "awaiting_review",
            "artifact_id": "artifact-course-review-1",
        }
        provider.artifact.return_value = course_review_envelope(self.project.id, inventory)
        revision = self.project.storyboard.revision
        with patch.object(main, "PythiaProvider", return_value=provider), TestClient(main.app) as client:
            created = client.post(
                f"/api/projects/{self.project.id}/pythia-course-reviews",
                json={
                    "taskId": "task-course-review-1",
                    "artifactId": "artifact-course-review-1",
                    "sourceRevision": revision,
                },
            )
            listed = client.get(f"/api/projects/{self.project.id}/pythia-course-reviews")
            review_id = created.json()["id"]
            resolved = client.patch(
                f"/api/projects/{self.project.id}/pythia-course-reviews/{review_id}/findings/standards-1",
                json={"status": "resolved", "resolutionNote": "Standard selected in the LMS plan."},
            )

        self.assertEqual(created.status_code, 201, created.text)
        self.assertEqual(listed.status_code, 200, listed.text)
        self.assertEqual(listed.json()[0]["artifact"]["artifact"]["type"], "course-review")
        self.assertEqual(listed.json()[0]["findings"][0]["status"], "open")
        self.assertEqual(resolved.status_code, 200, resolved.text)
        self.assertEqual(resolved.json()["findings"][0]["status"], "resolved")
        self.assertEqual(resolved.json()["findings"][0]["decisionHistory"][0]["resolutionNote"], "Standard selected in the LMS plan.")
        self.assertEqual(storage.read_project(self.project.id).storyboard.revision, revision)

    def test_assessment_artifact_becomes_editable_proposal_and_mutates_only_on_accept(self):
        objective_ids = ["objective-evacuation"]
        source_revision = self.project.storyboard.revision
        provider = Mock()
        provider.task_status.return_value = {
            "task_id": "task-assessment-1",
            "project_id": self.project.id,
            "consumer": "chimera",
            "capability": "assessment.design",
            "status": "awaiting_review",
            "artifact_id": "artifact-assessment-1",
        }
        provider.artifact.return_value = assessment_blueprint_envelope(self.project.id, [self.scene.id], objective_ids)
        with patch.object(main, "PythiaProvider", return_value=provider), TestClient(main.app) as client:
            saved = client.post(
                f"/api/projects/{self.project.id}/pythia-assessment-proposals",
                json={
                    "taskId": "task-assessment-1",
                    "artifactId": "artifact-assessment-1",
                    "sourceSceneIds": [self.scene.id],
                    "objectiveIds": objective_ids,
                    "sourceRevision": source_revision,
                },
            )
            proposal = saved.json()[0]
            before_accept = storage.read_project(self.project.id)
            accepted = client.post(
                f"/api/projects/{self.project.id}/enhancement-proposals/{proposal['id']}/decision",
                headers={"If-Match": f'"{source_revision}"'},
                json={"decision": "accept"},
            )

        self.assertEqual(saved.status_code, 201, saved.text)
        self.assertEqual(len(saved.json()), 1)
        self.assertEqual(proposal["field"], "assessmentItem")
        self.assertEqual(proposal["sourceTaskId"], "task-assessment-1")
        self.assertEqual(proposal["sourceArtifactId"], "artifact-assessment-1")
        self.assertEqual(proposal["sourceCapability"], "assessment.design")
        self.assertEqual(proposal["proposedValue"]["identifier"], "evacuation-response")
        self.assertIsNone(before_accept.storyboard.scenes[0].assessmentItem)
        self.assertEqual(accepted.status_code, 200, accepted.text)
        accepted_project = storage.read_project(self.project.id)
        self.assertEqual(accepted_project.storyboard.revision, source_revision + 1)
        self.assertEqual(accepted_project.storyboard.scenes[0].assessmentItem.identifier, "evacuation-response")

    def test_transform_proposal_accept_replaces_only_selected_scenes_and_writes_history(self):
        self.project = storage.create_scene(self.project.id, SceneCreate(title="Unselected scene"))
        unselected = self.project.storyboard.scenes[-1]
        self.project = storage.create_scene(self.project.id, SceneCreate(title="Second selected scene"))
        second_selected = self.project.storyboard.scenes[-1]
        source_scene_ids = [self.scene.id, second_selected.id]
        source_revision = self.project.storyboard.revision
        provider = Mock()
        provider.task_status.return_value = {
            "task_id": "task-transform-1",
            "project_id": self.project.id,
            "consumer": "chimera",
            "status": "awaiting_review",
            "artifact_id": "artifact-transform-1",
        }
        artifact = transformation_envelope(self.project.id, source_scene_ids)
        edited_scenes = artifact["payload"]["proposed_scenes"]
        edited_scenes[0]["title"] = "Edited immediate action"
        provider.artifact.return_value = artifact
        with patch.object(main, "PythiaProvider", return_value=provider), TestClient(main.app) as client:
            created = client.post(
                f"/api/projects/{self.project.id}/pythia-transform-proposals",
                json={
                    "taskId": "task-transform-1",
                    "artifactId": "artifact-transform-1",
                    "sourceSceneIds": source_scene_ids,
                    "sourceRevision": source_revision,
                },
            )
            listed = client.get(f"/api/projects/{self.project.id}/pythia-transform-proposals")
            readiness = client.get(f"/api/projects/{self.project.id}/readiness")
            before_accept = storage.read_project(self.project.id)
            accepted = client.post(
                f"/api/projects/{self.project.id}/pythia-transform-proposals/{created.json()['id']}/decision",
                headers={"If-Match": f'"{source_revision}"'},
                json={"decision": "accept", "proposedScenes": edited_scenes},
            )

        self.assertEqual(created.status_code, 201, created.text)
        self.assertEqual(listed.status_code, 200, listed.text)
        self.assertEqual(listed.json()[0]["status"], "pending")
        self.assertEqual(readiness.json()["pendingProposals"], 1)
        self.assertEqual(before_accept.storyboard.revision, source_revision)
        self.assertEqual(accepted.status_code, 200, accepted.text)
        self.assertEqual(accepted.json()["proposal"]["status"], "accepted")
        project = storage.read_project(self.project.id)
        self.assertEqual(project.storyboard.revision, source_revision + 1)
        self.assertEqual(
            [scene.title for scene in project.storyboard.scenes],
            ["Edited immediate action", "Continue to the assembly point", unselected.title],
        )
        self.assertEqual(accepted.json()["proposal"]["decisionHistory"][0]["acceptedScenes"][0]["title"], "Edited immediate action")
        self.assertTrue(list((storage.project_path(self.project.id) / "history").glob("storyboard-*.json")))

    def test_transform_proposal_reject_leaves_storyboard_unchanged(self):
        source_scene_ids = [self.scene.id]
        source_revision = self.project.storyboard.revision
        provider = Mock()
        provider.task_status.return_value = {
            "task_id": "task-transform-1",
            "project_id": self.project.id,
            "consumer": "chimera",
            "status": "awaiting_review",
            "artifact_id": "artifact-transform-1",
        }
        provider.artifact.return_value = transformation_envelope(self.project.id, source_scene_ids)
        with patch.object(main, "PythiaProvider", return_value=provider), TestClient(main.app) as client:
            created = client.post(
                f"/api/projects/{self.project.id}/pythia-transform-proposals",
                json={
                    "taskId": "task-transform-1",
                    "artifactId": "artifact-transform-1",
                    "sourceSceneIds": source_scene_ids,
                    "sourceRevision": source_revision,
                },
            )
            rejected = client.post(
                f"/api/projects/{self.project.id}/pythia-transform-proposals/{created.json()['id']}/decision",
                headers={"If-Match": f'"{source_revision}"'},
                json={"decision": "reject"},
            )

        self.assertEqual(rejected.status_code, 200, rejected.text)
        self.assertEqual(rejected.json()["proposal"]["status"], "rejected")
        self.assertEqual(storage.read_project(self.project.id).storyboard.revision, source_revision)
        self.assertEqual(storage.read_project(self.project.id).storyboard.scenes[0].title, self.scene.title)

    def test_transform_proposal_cannot_be_accepted_after_source_revision_changes(self):
        source_scene_ids = [self.scene.id]
        source_revision = self.project.storyboard.revision
        provider = Mock()
        provider.task_status.return_value = {
            "task_id": "task-transform-1",
            "project_id": self.project.id,
            "consumer": "chimera",
            "status": "awaiting_review",
            "artifact_id": "artifact-transform-1",
        }
        provider.artifact.return_value = transformation_envelope(self.project.id, source_scene_ids)
        with patch.object(main, "PythiaProvider", return_value=provider), TestClient(main.app) as client:
            created = client.post(
                f"/api/projects/{self.project.id}/pythia-transform-proposals",
                json={
                    "taskId": "task-transform-1",
                    "artifactId": "artifact-transform-1",
                    "sourceSceneIds": source_scene_ids,
                    "sourceRevision": source_revision,
                },
            )
            self.project = storage.create_scene(self.project.id, SceneCreate(title="Author's newer scene"))
            current_revision = self.project.storyboard.revision
            accepted = client.post(
                f"/api/projects/{self.project.id}/pythia-transform-proposals/{created.json()['id']}/decision",
                headers={"If-Match": f'"{current_revision}"'},
                json={"decision": "accept"},
            )

        self.assertEqual(accepted.status_code, 409, accepted.text)
        self.assertEqual(storage.list_pythia_transform_proposals(self.project.id)[0].status, "stale")
        self.assertEqual(storage.read_project(self.project.id).storyboard.revision, current_revision)
        self.assertEqual(storage.read_project(self.project.id).storyboard.scenes[0].title, self.scene.title)

    def test_submit_refuses_unadvertised_capability(self):
        provider = Mock()
        provider.health.return_value = {"status": "ok"}
        provider.capabilities.return_value = {"capabilities": [{"id": "other.capability"}]}
        with patch.object(main, "PythiaProvider", return_value=provider):
            with self.assertRaises(HTTPException) as raised:
                main.submit_pythia_analysis(self.project.id, self.scene.id)

        self.assertEqual(raised.exception.status_code, 422)
        provider.submit_task.assert_not_called()

    def test_resubmission_uses_a_new_task_id(self):
        provider = Mock()
        provider.health.return_value = {"status": "ok"}
        provider.capabilities.return_value = {"capabilities": [{"id": "storyboard.analyze"}]}
        provider.submit_task.side_effect = lambda request: {"task_id": request["task_id"], "status": "submitted"}
        with patch.object(main, "PythiaProvider", return_value=provider):
            first = main.submit_pythia_analysis(self.project.id, self.scene.id)
            second = main.submit_pythia_analysis(self.project.id, self.scene.id)

        self.assertNotEqual(first["task_id"], second["task_id"])

    def test_invalid_pythia_configuration_returns_recoverable_unavailable(self):
        with patch.dict("os.environ", {"CHIMERA_PYTHIA_BASE_URL": "ftp://invalid.example"}):
            with self.assertRaises(HTTPException) as raised:
                main.get_pythia_status(self.project.id)

        self.assertEqual(raised.exception.status_code, 503)

    def test_task_status_cannot_cross_project_boundaries(self):
        provider = Mock()
        provider.task_status.return_value = {"project_id": "other-project", "consumer": "chimera"}
        with patch.object(main, "PythiaProvider", return_value=provider):
            with self.assertRaises(HTTPException) as raised:
                main.get_pythia_task(self.project.id, "foreign-task")

        self.assertEqual(raised.exception.status_code, 404)

    def test_pythia_analysis_proposal_is_persisted_and_review_does_not_mutate_storyboard(self):
        source_ref = f"chimera://{self.project.id}/scenes/{self.scene.id}"
        provider = Mock()
        provider.task_status.return_value = {
            "task_id": "task-analysis-1",
            "project_id": self.project.id,
            "consumer": "chimera",
            "status": "awaiting_review",
            "artifact_id": "artifact-analysis-1",
        }
        provider.artifact.return_value = {
            "schema_version": "1.0",
            "artifact": {"type": "storyboard-analysis", "id": "artifact-analysis-1", "version": 1},
            "status": "draft",
            "validation": {"state": "passed", "validators": ["storyboard-analysis-1.0"]},
            "producer": {"skill": "storyboard-producer", "version": "1.0"},
            "sources": [source_ref],
            "payload": {
                "project_id": self.project.id,
                "scene_id": self.scene.id,
                "source_refs": [source_ref],
                "findings": [{
                    "id": "finding-1",
                    "category": "content_clarity",
                    "statement": "The evacuation action is clear.",
                    "evidence_refs": [source_ref],
                    "uncertainty": "The assembly location is not named.",
                }],
                "recommendations": [{
                    "id": "recommendation-1",
                    "statement": "Name the assembly location.",
                    "finding_refs": ["finding-1"],
                }],
                "uncertainties": [],
                "review_questions": [],
            },
        }
        assert_contract_schema(self, "artifact-envelope.schema.json", provider.artifact.return_value)
        assert_contract_schema(self, "storyboard-analysis.schema.json", provider.artifact.return_value["payload"])
        source_revision = self.project.storyboard.revision
        with patch.object(main, "PythiaProvider", return_value=provider), TestClient(main.app) as client:
            created = client.post(
                f"/api/projects/{self.project.id}/pythia-analysis-proposals",
                json={
                    "taskId": "task-analysis-1",
                    "artifactId": "artifact-analysis-1",
                    "sceneId": self.scene.id,
                    "sourceRevision": source_revision,
                },
            )
            listed = client.get(f"/api/projects/{self.project.id}/pythia-analysis-proposals")
            readiness = client.get(f"/api/projects/{self.project.id}/readiness")
            repeated = client.post(
                f"/api/projects/{self.project.id}/pythia-analysis-proposals",
                json={
                    "taskId": "task-analysis-1",
                    "artifactId": "artifact-analysis-1",
                    "sceneId": self.scene.id,
                    "sourceRevision": source_revision,
                },
            )
            proposal_id = created.json()["id"]
            accepted = client.post(
                f"/api/projects/{self.project.id}/pythia-analysis-proposals/{proposal_id}/decision",
                json={"decision": "accept"},
            )

        self.assertEqual(created.status_code, 201, created.text)
        self.assertEqual(listed.status_code, 200, listed.text)
        self.assertEqual(len(listed.json()), 1)
        self.assertEqual(readiness.json()["pendingProposals"], 1)
        self.assertEqual(repeated.json()["id"], proposal_id)
        self.assertEqual(accepted.json()["status"], "accepted")
        self.assertEqual(accepted.json()["artifact"]["payload"]["findings"][0]["id"], "finding-1")
        self.assertEqual(accepted.json()["decisionHistory"][0]["storyboardRevision"], source_revision)
        self.assertEqual(storage.read_project(self.project.id).storyboard.revision, source_revision)

    def test_pythia_proposal_requires_task_artifact_match(self):
        provider = Mock()
        provider.task_status.return_value = {
            "task_id": "task-analysis-1",
            "project_id": self.project.id,
            "consumer": "chimera",
            "status": "awaiting_review",
            "artifact_id": "actual-artifact",
        }
        with patch.object(main, "PythiaProvider", return_value=provider), TestClient(main.app) as client:
            response = client.post(
                f"/api/projects/{self.project.id}/pythia-analysis-proposals",
                json={
                    "taskId": "task-analysis-1",
                    "artifactId": "different-artifact",
                    "sceneId": self.scene.id,
                    "sourceRevision": self.project.storyboard.revision,
                },
            )

        self.assertEqual(response.status_code, 409, response.text)
        provider.artifact.assert_not_called()

    def test_valid_artifact_is_returned_without_storyboard_mutation(self):
        provider = Mock()
        provider.artifact.return_value = {
            "artifact": {"type": "storyboard-analysis"},
            "status": "draft",
            "validation": {"state": "passed"},
            "producer": {"skill": "storyboard-producer"},
            "payload": {
                "project_id": self.project.id,
                "scene_id": self.scene.id,
                "source_refs": [f"chimera://{self.project.id}/scenes/{self.scene.id}"],
                "findings": [{
                    "id": "evacuation-route",
                    "category": "content_clarity",
                    "statement": "The scene identifies the required evacuation action.",
                    "evidence_refs": [f"chimera://{self.project.id}/scenes/{self.scene.id}"],
                    "uncertainty": "The destination is not named.",
                }],
                "recommendations": [{
                    "id": "name-assembly-point",
                    "statement": "Name the designated assembly point.",
                    "finding_refs": ["evacuation-route"],
                }],
                "uncertainties": [],
                "review_questions": [],
            },
        }
        with patch.object(main, "PythiaProvider", return_value=provider):
            result = main.get_pythia_artifact(self.project.id, self.scene.id, "fixture.artifact")

        self.assertEqual(result["payload"]["scene_id"], self.scene.id)
        self.assertEqual(storage.read_project(self.project.id).storyboard.revision, self.project.storyboard.revision)

    def test_http_round_trip_discovers_submits_polls_and_retrieves_draft(self):
        provider = Mock()
        provider.health.return_value = {"status": "ok", "version": "1.0"}
        provider.capabilities.return_value = {"capabilities": [{"id": "storyboard.analyze", "provider": "synthetic"}]}
        provider.submit_task.side_effect = lambda request: {
            "task_id": request["task_id"], "project_id": self.project.id, "consumer": "chimera", "status": "submitted",
        }
        provider.task_status.side_effect = lambda task_id: {
            "task_id": task_id, "project_id": self.project.id, "consumer": "chimera", "status": "awaiting_review",
            "artifact_id": "fixture.artifact",
        }
        provider.artifact.return_value = {
            "artifact": {"type": "storyboard-analysis"},
            "status": "draft",
            "validation": {"state": "passed"},
            "producer": {"skill": "storyboard-producer"},
            "payload": {
                "project_id": self.project.id,
                "scene_id": self.scene.id,
                "source_refs": [f"chimera://{self.project.id}/scenes/{self.scene.id}"],
                "findings": [{
                    "id": "synthetic-finding",
                    "category": "content_clarity",
                    "statement": "Synthetic fixture finding.",
                    "evidence_refs": [f"chimera://{self.project.id}/scenes/{self.scene.id}"],
                    "uncertainty": "The synthetic fixture has no additional context.",
                }],
                "recommendations": [{
                    "id": "synthetic-recommendation",
                    "statement": "Review the synthetic finding.",
                    "finding_refs": ["synthetic-finding"],
                }],
                "uncertainties": [],
                "review_questions": [],
            },
        }
        with patch.object(main, "PythiaProvider", return_value=provider), TestClient(main.app) as client:
            status = client.get(f"/api/projects/{self.project.id}/pythia-status")
            submitted = client.post(f"/api/projects/{self.project.id}/scenes/{self.scene.id}/pythia-analysis")
            task_id = submitted.json()["task_id"]
            polled = client.get(f"/api/projects/{self.project.id}/pythia-tasks/{task_id}")
            artifact = client.get(
                f"/api/projects/{self.project.id}/scenes/{self.scene.id}/pythia-artifacts/fixture.artifact",
            )

        self.assertEqual(status.status_code, 200, status.text)
        self.assertTrue(status.json()["supportsStoryboardAnalyze"])
        self.assertEqual(submitted.status_code, 202, submitted.text)
        self.assertEqual(submitted.json()["status"], "submitted")
        self.assertEqual(polled.json()["status"], "awaiting_review")
        self.assertEqual(artifact.json()["payload"]["findings"][0]["id"], "synthetic-finding")
        self.assertEqual(artifact.json()["payload"]["recommendations"][0]["finding_refs"], ["synthetic-finding"])
        self.assertEqual(storage.read_project(self.project.id).storyboard.revision, self.project.storyboard.revision)

    @unittest.skipUnless(os.environ.get("CHIMERA_PYTHIA_INTEGRATION") == "1", "set CHIMERA_PYTHIA_INTEGRATION=1 to call local Pythia")
    def test_live_pythia_round_trip_with_fixture_scene_content(self):
        self.project = storage.create_scene(self.project.id, SceneCreate(
            title="Respond to a fire alarm",
            extractedText=["Leave by the nearest marked exit and assemble at the designated area. Do not retrieve personal items."],
            notes="Emphasize the local assembly point and the reason personal items must be left behind.",
        ))
        scene = self.project.storyboard.scenes[-1]
        revision = self.project.storyboard.revision
        base_url = os.environ.get("CHIMERA_PYTHIA_BASE_URL", "http://127.0.0.1:8001")

        with patch.dict("os.environ", {"CHIMERA_PYTHIA_BASE_URL": base_url}):
            task = main.submit_pythia_analysis(self.project.id, scene.id)
            source_revision = task["source_revision"]
            deadline = time.monotonic() + 20
            while task.get("status") in {"submitted", "queued", "running"} and time.monotonic() < deadline:
                time.sleep(0.1)
                task = main.get_pythia_task(self.project.id, task["task_id"])

            self.assertEqual(task.get("status"), "awaiting_review", task)
            self.assertIsInstance(task.get("artifact_id"), str)
            artifact = main.get_pythia_artifact(self.project.id, scene.id, task["artifact_id"])
            proposal = main.post_pythia_analysis_proposal(self.project.id, PythiaAnalysisProposalCreate(
                taskId=task["task_id"],
                artifactId=task["artifact_id"],
                sceneId=scene.id,
                sourceRevision=source_revision,
            ))
            decided = main.post_pythia_analysis_proposal_decision(
                self.project.id,
                proposal.id,
                main.PythiaAnalysisProposalDecision(decision="accept"),
            )

        self.assertEqual(artifact["payload"]["project_id"], self.project.id)
        self.assertEqual(artifact["payload"]["scene_id"], scene.id)
        self.assertEqual(proposal.status, "pending")
        self.assertEqual(decided.status, "accepted")
        self.assertTrue(artifact["payload"]["findings"])
        self.assertTrue(all(
            isinstance(item, dict) and isinstance(item.get("id"), str)
            and isinstance(item.get("statement"), str) and item.get("evidence_refs")
            for item in artifact["payload"]["findings"]
        ))
        self.assertTrue(all(isinstance(item, dict) for item in artifact["payload"]["recommendations"]))
        self.assertEqual(storage.read_project(self.project.id).storyboard.revision, revision)

    @unittest.skipUnless(os.environ.get("CHIMERA_PYTHIA_TRANSFORM_INTEGRATION") == "1", "set CHIMERA_PYTHIA_TRANSFORM_INTEGRATION=1 to call local Pythia")
    def test_live_pythia_transform_round_trip_persists_and_accepts_proposal(self):
        self.project = storage.create_scene(self.project.id, SceneCreate(
            title="Respond to a fire alarm",
            extractedText=["Leave by the nearest marked exit."],
            notes="Emphasize the local assembly point.",
        ))
        self.project = storage.create_scene(self.project.id, SceneCreate(
            title="Support others safely",
            extractedText=["Report missing colleagues only after reaching the assembly area."],
            notes="Do not re-enter the building.",
        ))
        source_scenes = self.project.storyboard.scenes[-2:]
        source_scene_ids = [scene.id for scene in source_scenes]
        base_url = os.environ.get("CHIMERA_PYTHIA_BASE_URL", "http://127.0.0.1:8002")

        with patch.dict("os.environ", {"CHIMERA_PYTHIA_BASE_URL": base_url}):
            submitted = main.submit_pythia_transform(
                self.project.id,
                main.PythiaTransformRequest(sceneIds=source_scene_ids),
            )
            source_revision = submitted["source_revision"]
            task = submitted
            deadline = time.monotonic() + 20
            while task.get("status") in {"submitted", "queued", "running"} and time.monotonic() < deadline:
                time.sleep(0.1)
                task = main.get_pythia_task(self.project.id, submitted["task_id"])

            self.assertEqual(task.get("status"), "awaiting_review", task)
            proposal = main.post_pythia_transform_proposal(self.project.id, main.PythiaTransformProposalCreate(
                taskId=task["task_id"],
                artifactId=task["artifact_id"],
                sourceSceneIds=source_scene_ids,
                sourceRevision=source_revision,
            ))
            result = main.post_pythia_transform_proposal_decision(
                self.project.id,
                proposal.id,
                main.PythiaTransformProposalDecision(decision="accept"),
                if_match=f'"{source_revision}"',
            )

        self.assertEqual(proposal.status, "pending")
        self.assertEqual(proposal.artifact["artifact"]["type"], "storyboard-transformation")
        self.assertTrue(proposal.artifact["payload"]["proposed_scenes"])
        self.assertEqual(result.proposal.status, "accepted")
        self.assertEqual(result.project.storyboard.revision, source_revision + 1)
        self.assertEqual(len(result.project.storyboard.scenes), 3)
        self.assertEqual(result.project.storyboard.scenes[0].title, self.scene.title)
        self.assertEqual(len(result.project.storyboard.scenes[1:]), len(proposal.artifact["payload"]["proposed_scenes"]))

    @unittest.skipUnless(os.environ.get("CHIMERA_PYTHIA_ASSESSMENT_INTEGRATION") == "1", "set CHIMERA_PYTHIA_ASSESSMENT_INTEGRATION=1 to call local Pythia")
    def test_live_pythia_assessment_design_becomes_reviewable_native_proposal(self):
        self.project = storage.update_scene(
            self.project.id,
            self.scene.id,
            main.SceneUpdate(
                extractedText=["Leave by the nearest marked exit and report to the assembly area."],
                notes="Emphasize the local assembly point.",
            ),
        )
        source_revision = self.project.storyboard.revision
        base_url = os.environ.get("CHIMERA_PYTHIA_BASE_URL", "http://127.0.0.1:8003")
        with patch.dict("os.environ", {"CHIMERA_PYTHIA_BASE_URL": base_url}):
            submitted = main.submit_pythia_assessment_design(
                self.project.id,
                main.PythiaAssessmentDesignRequest(
                    sceneIds=[self.scene.id],
                    objectives=[main.PythiaAssessmentObjective(
                        id="objective-evacuation",
                        statement="Choose the safe evacuation response.",
                        sceneIds=[self.scene.id],
                    )],
                    audience="New staff",
                    desiredDifficulty="application",
                    completionRules="One correct response is required.",
                ),
            )
            task = submitted
            deadline = time.monotonic() + 20
            while task.get("status") in {"submitted", "queued", "running"} and time.monotonic() < deadline:
                time.sleep(0.1)
                task = main.get_pythia_task(self.project.id, submitted["task_id"])
            self.assertEqual(task.get("status"), "awaiting_review", task)
            proposals = main.post_pythia_assessment_proposal(self.project.id, PythiaAssessmentProposalCreate(
                taskId=task["task_id"],
                artifactId=task["artifact_id"],
                sourceSceneIds=submitted["source_scene_ids"],
                objectiveIds=submitted["objective_ids"],
                sourceRevision=submitted["source_revision"],
            ))

        self.assertTrue(proposals)
        proposal = proposals[0]
        self.assertEqual(proposal.sourceTaskId, task["task_id"])
        self.assertEqual(proposal.sourceArtifactId, task["artifact_id"])
        self.assertIsNone(storage.read_project(self.project.id).storyboard.scenes[0].assessmentItem)
        decided, updated_project = storage.decide_enhancement_proposal(
            self.project.id,
            proposal.id,
            "accept",
            source_revision,
        )
        self.assertEqual(decided.status, "accepted")
        self.assertEqual(updated_project.storyboard.revision, source_revision + 1)
        self.assertEqual(updated_project.storyboard.scenes[0].assessmentItem.identifier, proposal.proposedValue["identifier"])

    @unittest.skipUnless(os.environ.get("CHIMERA_PYTHIA_COURSE_REVIEW_INTEGRATION") == "1", "set CHIMERA_PYTHIA_COURSE_REVIEW_INTEGRATION=1 to call local Pythia")
    def test_live_pythia_course_review_persists_and_resolves_finding_without_mutation(self):
        self.project = storage.create_scene(self.project.id, SceneCreate(
            title="Respond to a fire alarm",
            extractedText=["Leave by the nearest marked exit."],
            notes="Emphasize the local assembly point.",
        ))
        source_revision = self.project.storyboard.revision
        base_url = os.environ.get("CHIMERA_PYTHIA_BASE_URL", "http://127.0.0.1:8003")
        with patch.dict("os.environ", {"CHIMERA_PYTHIA_BASE_URL": base_url}):
            submitted = main.submit_pythia_course_review(self.project.id)
            task = submitted
            deadline = time.monotonic() + 20
            while task.get("status") in {"submitted", "queued", "running"} and time.monotonic() < deadline:
                time.sleep(0.1)
                task = main.get_pythia_task(self.project.id, submitted["task_id"])
            self.assertEqual(task.get("status"), "awaiting_review", task)
            review = main.post_pythia_course_review(self.project.id, main.PythiaCourseReviewCreate(
                taskId=task["task_id"],
                artifactId=task["artifact_id"],
                sourceRevision=submitted["source_revision"],
            ))
            self.assertTrue(review.findings)
            finding = review.findings[0]
            decided = main.patch_pythia_course_finding(
                self.project.id,
                review.id,
                finding.findingId,
                main.PythiaCourseFindingDecision(status="resolved", resolutionNote="Reviewed against the current course plan."),
            )

        self.assertEqual(decided.findings[0].status, "resolved")
        self.assertEqual(decided.artifact["artifact"]["type"], "course-review")
        self.assertEqual(storage.read_project(self.project.id).storyboard.revision, source_revision)

    @unittest.skipUnless(os.environ.get("CHIMERA_PYTHIA_GOLDEN_PATH_INTEGRATION") == "1", "set CHIMERA_PYTHIA_GOLDEN_PATH_INTEGRATION=1 to run the complete local Pythia golden path")
    def test_live_pythia_golden_path_imports_reviews_transforms_assesses_and_exports(self):
        presentation = Presentation()
        for title, body, notes in [
            ("Respond to an evacuation alarm", "Use the nearest marked exit. Leave personal items behind.", "Emphasize the designated assembly area."),
            ("Continue to the assembly area", "Report to the designated assembly area and wait for instructions.", "Do not re-enter the building."),
        ]:
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            title_box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
            title_box.text_frame.text = title
            body_box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(2))
            body_box.text_frame.text = body
            slide.notes_slide.notes_text_frame.text = notes
        deck = BytesIO()
        presentation.save(deck)

        project = storage.create_project(ProjectCreate(title="Sprint 10 Golden Path"))
        base_url = os.environ.get("CHIMERA_PYTHIA_BASE_URL", "http://127.0.0.1:8001")

        with patch.dict("os.environ", {"CHIMERA_PYTHIA_BASE_URL": base_url}), TestClient(main.app) as client:
            def wait_for_task(task):
                deadline = time.monotonic() + 60
                while task.get("status") in {"submitted", "queued", "running"} and time.monotonic() < deadline:
                    time.sleep(0.1)
                    response = client.get(f"/api/projects/{project.id}/pythia-tasks/{task['task_id']}")
                    self.assertEqual(response.status_code, 200, response.text)
                    task = response.json()
                self.assertEqual(task.get("status"), "awaiting_review", task)
                self.assertIsInstance(task.get("artifact_id"), str)
                return task

            status = client.get(f"/api/projects/{project.id}/pythia-status")
            self.assertEqual(status.status_code, 200, status.text)
            self.assertEqual(status.json()["health"]["status"], "ok")
            self.assertTrue(status.json()["supportsStoryboardAnalyze"])
            self.assertTrue(status.json()["supportsStoryboardTransform"])
            self.assertTrue(status.json()["supportsAssessmentDesign"])
            self.assertTrue(status.json()["supportsCourseReview"])

            imported = client.post(
                f"/api/projects/{project.id}/source",
                files={"file": ("legacy-evacuation.pptx", deck.getvalue(), "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
                headers={"If-Match": f'"{project.storyboard.revision}"'},
            )
            self.assertEqual(imported.status_code, 200, imported.text)
            imported_project = imported.json()
            self.assertTrue(imported_project["hasSourceDeck"])
            source_scene_ids = [scene["id"] for scene in imported_project["storyboard"]["scenes"]]
            self.assertEqual(len(source_scene_ids), 2)
            source_revision = imported_project["storyboard"]["revision"]

            analysis = client.post(f"/api/projects/{project.id}/scenes/{source_scene_ids[0]}/pythia-analysis")
            self.assertEqual(analysis.status_code, 202, analysis.text)
            analysis_task = wait_for_task(analysis.json())
            analysis_artifact = client.get(
                f"/api/projects/{project.id}/scenes/{source_scene_ids[0]}/pythia-artifacts/{analysis_task['artifact_id']}"
            )
            self.assertEqual(analysis_artifact.status_code, 200, analysis_artifact.text)
            self.assertTrue(analysis_artifact.json()["payload"]["findings"])
            analysis_proposal = client.post(
                f"/api/projects/{project.id}/pythia-analysis-proposals",
                json={
                    "taskId": analysis_task["task_id"],
                    "artifactId": analysis_task["artifact_id"],
                    "sceneId": source_scene_ids[0],
                    "sourceRevision": source_revision,
                },
            )
            self.assertEqual(analysis_proposal.status_code, 201, analysis_proposal.text)
            analysis_decision = client.post(
                f"/api/projects/{project.id}/pythia-analysis-proposals/{analysis_proposal.json()['id']}/decision",
                json={"decision": "accept"},
            )
            self.assertEqual(analysis_decision.status_code, 200, analysis_decision.text)
            self.assertEqual(analysis_decision.json()["status"], "accepted")
            self.assertEqual(storage.read_project(project.id).storyboard.revision, source_revision)

            transform = client.post(
                f"/api/projects/{project.id}/pythia-transform",
                json={"sceneIds": source_scene_ids, "constraints": {}},
            )
            self.assertEqual(transform.status_code, 202, transform.text)
            transform_task = wait_for_task(transform.json())
            transform_proposal = client.post(
                f"/api/projects/{project.id}/pythia-transform-proposals",
                json={
                    "taskId": transform_task["task_id"],
                    "artifactId": transform_task["artifact_id"],
                    "sourceSceneIds": source_scene_ids,
                    "sourceRevision": source_revision,
                },
            )
            self.assertEqual(transform_proposal.status_code, 201, transform_proposal.text)
            transformed = client.post(
                f"/api/projects/{project.id}/pythia-transform-proposals/{transform_proposal.json()['id']}/decision",
                json={"decision": "accept"},
                headers={"If-Match": f'"{source_revision}"'},
            )
            self.assertEqual(transformed.status_code, 200, transformed.text)
            self.assertEqual(transformed.json()["proposal"]["status"], "accepted")
            transformed_project = transformed.json()["project"]
            target_scene_id = transformed_project["storyboard"]["scenes"][0]["id"]
            transformed_revision = transformed_project["storyboard"]["revision"]

            assessment = client.post(
                f"/api/projects/{project.id}/pythia-assessment-design",
                json={
                    "sceneIds": [target_scene_id],
                    "objectives": [{
                        "id": "objective-evacuation",
                        "statement": "Choose the safe response to an evacuation alarm.",
                        "sceneIds": [target_scene_id],
                    }],
                    "audience": "New staff",
                    "desiredDifficulty": "application",
                    "completionRules": "One correct response is required.",
                    "constraints": {},
                },
            )
            self.assertEqual(assessment.status_code, 202, assessment.text)
            assessment_task = wait_for_task(assessment.json())
            assessment_proposals = client.post(
                f"/api/projects/{project.id}/pythia-assessment-proposals",
                json={
                    "taskId": assessment_task["task_id"],
                    "artifactId": assessment_task["artifact_id"],
                    "sourceSceneIds": [target_scene_id],
                    "objectiveIds": ["objective-evacuation"],
                    "sourceRevision": transformed_revision,
                },
            )
            self.assertEqual(assessment_proposals.status_code, 201, assessment_proposals.text)
            self.assertTrue(assessment_proposals.json())
            assessment_proposal = assessment_proposals.json()[0]
            assessment_decision = client.post(
                f"/api/projects/{project.id}/enhancement-proposals/{assessment_proposal['id']}/decision",
                json={"decision": "accept"},
                headers={"If-Match": f'"{transformed_revision}"'},
            )
            self.assertEqual(assessment_decision.status_code, 200, assessment_decision.text)
            self.assertEqual(assessment_decision.json()["proposal"]["status"], "accepted")
            released_revision = assessment_decision.json()["project"]["storyboard"]["revision"]
            self.assertIsNotNone(next(
                scene for scene in assessment_decision.json()["project"]["storyboard"]["scenes"]
                if scene["id"] == target_scene_id
            )["assessmentItem"])

            course_review_task_response = client.post(f"/api/projects/{project.id}/pythia-course-review")
            self.assertEqual(course_review_task_response.status_code, 202, course_review_task_response.text)
            course_review_task = wait_for_task(course_review_task_response.json())
            course_review_response = client.post(
                f"/api/projects/{project.id}/pythia-course-reviews",
                json={
                    "taskId": course_review_task["task_id"],
                    "artifactId": course_review_task["artifact_id"],
                    "sourceRevision": released_revision,
                },
            )
            self.assertEqual(course_review_response.status_code, 201, course_review_response.text)
            review = course_review_response.json()
            findings = review["artifact"]["payload"]["findings"]
            self.assertTrue(findings)
            reviewed = review
            for finding in findings:
                decision = client.patch(
                    f"/api/projects/{project.id}/pythia-course-reviews/{review['id']}/findings/{finding['finding_id']}",
                    json={"status": "resolved", "resolutionNote": "Reviewed against the accepted course content."},
                )
                self.assertEqual(decision.status_code, 200, decision.text)
                reviewed = decision.json()
            self.assertTrue(all(item["status"] == "resolved" for item in reviewed["findings"]))
            self.assertEqual(storage.read_project(project.id).storyboard.revision, released_revision)

            release = client.patch(
                f"/api/projects/{project.id}/production-stage",
                json={"productionStage": "approved"},
            )
            self.assertEqual(release.status_code, 200, release.text)
            self.assertEqual(release.json()["productionStage"], "approved")

            package = client.get(f"/api/projects/{project.id}/scorm-2004.zip")
            self.assertEqual(package.status_code, 200, package.text[:500])
            from zipfile import ZipFile

            with ZipFile(BytesIO(package.content)) as archive:
                self.assertIn("imsmanifest.xml", archive.namelist())
                self.assertIn("index.html", archive.namelist())
            report = client.get(f"/api/projects/{project.id}/scorm-report/scorm2004")
            self.assertEqual(report.status_code, 200, report.text)
            self.assertEqual(report.json()["exportProfile"], "SCORM 2004 4th Edition")
            self.assertEqual(report.json()["buildStatus"], "passed")
            self.assertTrue(all(check["passed"] for check in report.json()["compatibilityChecks"]))


if __name__ == "__main__":
    unittest.main()