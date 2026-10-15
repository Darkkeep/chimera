"""Produce isolated editable Word acceptance artifacts, never user projects."""
import hashlib
import io
import json
from pathlib import Path
import subprocess
import sys
import tempfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import pymupdf
from docx import Document
from backend.app import media, storage, word_export
from backend.app.models import AssetAttributionUpdate, AssetMetadataUpdate, EnhancementProposalCreate, MediaPlacement, MediaPlacementsUpdate, ProjectCreate, ReferencePassage, ReviewCommentCreate, SceneCreate
from backend.app.render import find_libreoffice


def main():
    output = ROOT / "workspace" / "completion-1.1.0" / "word"
    output.mkdir(parents=True, exist_ok=True)
    results = []
    with tempfile.TemporaryDirectory(prefix="chimera-word-completion-") as directory:
        storage.ROOT = Path(directory) / "workspace"
        storage.PROJECTS = storage.ROOT / "projects"
        project = storage.create_project(ProjectCreate(title="Media handoff qualification"))
        project = storage.create_scene(project.id, SceneCreate(title="Inspect before operation", extractedText=["Inspect equipment before operation."], narration="Inspect equipment before operation.", notes="Synthetic authoring fixture."))
        project = storage.add_project_asset(project.id, "inspection.mp4", (ROOT / "tests/fixtures/media-qualification.mp4").read_bytes(), "video")
        asset = project.storyboard.assets[0]
        project = media.update_metadata(project.id, asset.id, AssetMetadataUpdate(transcript="Inspect equipment. No spoken audio.", captionsVtt="WEBVTT\n\n00:00:00.000 --> 00:00:02.000\nInspect equipment."), project.storyboard.revision)
        project = storage.update_project_asset_attribution(project.id, asset.id, AssetAttributionUpdate(attributionText="Locally generated synthetic qualification footage.", attributionStatus="complete"))
        project = media.write_placements(project.id, MediaPlacementsUpdate(placements=[MediaPlacement(assetId=asset.id)]), project.storyboard.revision)
        resource = storage.save_reference_resource(project.id, "guide.txt", "Inspection guide", "study-guide", True, [], b"Inspect equipment before operation.", "text/plain", [ReferencePassage(locator="line:1", text="Inspect equipment before operation.")], "complete", None)
        proposal = storage.create_enhancement_proposal(project.id, EnhancementProposalCreate(targetSceneId=project.storyboard.scenes[0].id, proposedValue="Inspect first.", rationale="Synthetic source-grounded handoff.", referenceCitations=[{"resourceId": resource.id, "version": "1", "locator": "line:1", "excerpt": "Inspect equipment before operation."}]))
        storage.decide_enhancement_proposal(project.id, proposal.id, "accept", project.storyboard.revision)
        storage.create_review_comment(project.id, ReviewCommentCreate(sceneId=project.storyboard.scenes[0].id, authorLabel="Synthetic reviewer", body="Confirm caption timing during human review."))
        project = storage.read_project(project.id)
        for name, builder in (("sme-master", word_export.build_sme_storyboard_docx), ("full-storyboard", word_export.build_full_storyboard_docx)):
            data = builder(project, include_review_comments=True)
            path = output / f"{name}.docx"
            path.write_bytes(data)
            document = Document(io.BytesIO(data))
            paragraphs = "\n".join(p.text for p in document.paragraphs)
            assert resource.versions[0].sha256 in paragraphs
            assert "Confirm caption timing" in paragraphs
            document.add_paragraph("Editable round-trip verification.")
            edited = output / f"{name}-edited.docx"
            document.save(edited)
            assert Document(str(edited)).paragraphs[-1].text == "Editable round-trip verification."
            result = {"format": name, "sourceRevision": project.storyboard.revision, "sha256": hashlib.sha256(data).hexdigest(), "editableRoundTrip": True, "interactiveEditing": "pending"}
            executable = find_libreoffice()
            if executable:
                profile = (Path(directory) / f"writer-{name}").as_uri()
                converted = subprocess.run([executable, "--headless", f"-env:UserInstallation={profile}", "--convert-to", "pdf", "--outdir", str(output), str(path)], capture_output=True, timeout=60)
                if converted.returncode:
                    raise RuntimeError("Word rendering failed")
                with pymupdf.open(output / f"{name}.pdf") as pdf:
                    result["pages"] = len(pdf)
                    result["orientations"] = ["landscape" if p.rect.width > p.rect.height else "portrait" for p in pdf]
                    for index in {0, len(pdf) - 1}:
                        pdf[index].get_pixmap(matrix=pymupdf.Matrix(1, 1)).save(output / f"{name}-page-{index + 1}.png")
            results.append(result)
    (output / "qualification.json").write_text(json.dumps(results, indent=2), encoding="utf-8")
    print(json.dumps(results, indent=2))


if __name__ == "__main__":
    main()
